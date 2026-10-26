"""
Module: services.report_render

The pluggable renderer interface for custom reports: a format is a `Renderer` registered for one or more
layout families, and a run's output is produced through `render` with no format named anywhere else.

Responsibilities:
- `Renderer`, `RenderInput`, `RenderedFile` and `register_renderer`.
- `formats_for(family)`: the formats a family can be produced in, in registration order (the catalogue and
  the route's format check read this).
- `render(fmt, family, input)`.
- The shipped renderers, registered at the foot of the module: `html`, `pdf`, `docx` (document family) and
  `csv`, `xlsx` (table family).

Design decisions:
- **A renderer sees only the neutral model.** A document renderer gets a `ReportDocument` (title, blocks,
  branding, orientation), a table renderer gets `Sheet`s; neither sees a layout, a record or a user, so
  formats cannot drift apart in content and a new format (ODT, ODS) is one registration.
- **Cells are text.** Every spreadsheet cell is written as a string, after the same formula guard the CSV
  export uses (`csv_safe`): a value from a record that starts with `=`, `+`, `-`, `@`, tab or CR gets a
  leading apostrophe, and the cell type is forced to text so a spreadsheet never evaluates it.
- **CSV of several sheets is a zip** of one CSV per sheet (unique, filename-safe names); one sheet is a
  plain CSV. XLSX holds every sheet as a tab (names cleaned of `[]:*?/\\`, at most 31 characters, unique
  case-insensitively); a cell is cut at Excel's 32,767-character limit.
- Existing reports keep their own path (`report_framework.render_pdf`/`render_docx`); this registry serves
  custom reports only.

Dependencies: `openpyxl` (XLSX), `services.report_pdf`, `services.report_docx`, `services.report_html`,
`services.csv_safety`, `services.report_model`.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet as _WriteOnlyWorksheet

from app.models.report_layout import FAMILY_DOCUMENT, FAMILY_TABLE
from app.services.csv_safety import csv_safe
from app.services.report_docx import render_docx_document
from app.services.report_html import render_html_document
from app.services.report_model import ReportDocument, Sheet
from app.services.report_pdf import render_pdf_document

#: Excel's longest cell text.
MAX_CELL_CHARS = 32_767
_SHEET_NAME_BAD = re.compile(r"[\[\]:*?/\\]")
_FILE_NAME_BAD = re.compile(r"[^A-Za-z0-9._ -]+")


@dataclass(frozen=True)
class RenderInput:
    """What a renderer is given.

    Attributes:
        title: The report title.
        document: The neutral document (document family), else `None`.
        sheets: The sheets (table family), else empty.
    """

    title: str
    document: ReportDocument | None = None
    sheets: Sequence[Sheet] = ()


@dataclass(frozen=True)
class RenderedFile:
    """A rendered file: its bytes and the extension (which also selects its content type)."""

    content: bytes
    extension: str


@dataclass(frozen=True)
class Renderer:
    """One output format.

    Attributes:
        key: The `format` value (`pdf`, `xlsx` ...).
        families: The layout families it can render.
        render: Produces the file from a `RenderInput` of one of those families.
    """

    key: str
    families: tuple[str, ...]
    render: Callable[[RenderInput], RenderedFile]


_RENDERERS: dict[str, Renderer] = {}


def register_renderer(renderer: Renderer) -> None:
    """Adds an output format.

    Raises:
        ValueError: The key is already registered.
    """
    if renderer.key in _RENDERERS:
        raise ValueError(f"format {renderer.key!r} already has a renderer")
    _RENDERERS[renderer.key] = renderer


def formats_for(family: str) -> list[str]:
    """The formats a layout family can be produced in, in registration order."""
    return [key for key, r in _RENDERERS.items() if family in r.families]


def render(fmt: str, family: str, source: RenderInput) -> RenderedFile:
    """Renders `source` as `fmt`.

    Raises:
        KeyError: `fmt` is not a registered format of `family`.
    """
    renderer = _RENDERERS.get(fmt)
    if renderer is None or family not in renderer.families:
        raise KeyError(fmt)
    return renderer.render(source)


# --- table family ----------------------------------------------------------------------------------------


def unique_names(names: Sequence[str], *, clean: Callable[[str], str], limit: int | None = None, suffix: str = "") -> list[str]:
    """`names` made distinct (case-insensitively) after `clean`, optionally cut to `limit` characters.

    A repeat gets ` (2)`, ` (3)` ... (inside the limit), and an empty name becomes `Sheet`.
    """
    taken: set[str] = set()
    out: list[str] = []
    for name in names:
        base = clean(name).strip() or "Sheet"
        candidate, n = base[:limit] if limit else base, 1
        while candidate.lower() in taken:
            n += 1
            tail = f" ({n})"
            candidate = (base[: limit - len(tail)] if limit else base) + tail
        taken.add(candidate.lower())
        out.append(candidate + suffix)
    return out


def _csv_bytes(sheet: Sheet) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([csv_safe(c) for c in sheet.columns])
    for row in sheet.rows:
        writer.writerow([csv_safe(c) for c in row])
    return buffer.getvalue().encode("utf-8")


def render_csv_sheets(source: RenderInput) -> RenderedFile:
    """One sheet as a CSV; several as a zip with one CSV per sheet. Every cell goes through `csv_safe`."""
    if len(source.sheets) == 1:
        return RenderedFile(_csv_bytes(source.sheets[0]), "csv")
    names = unique_names([s.name for s in source.sheets], clean=lambda n: _FILE_NAME_BAD.sub("_", n), suffix=".csv")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, sheet in zip(names, source.sheets, strict=True):
            archive.writestr(name, _csv_bytes(sheet))
    return RenderedFile(out.getvalue(), "zip")


def _text_cell(sheet: _WriteOnlyWorksheet, value: str, *, bold: bool = False) -> Cell:
    """A spreadsheet cell holding `value` as text: guarded against formulas and cut at Excel's limit."""
    cell = WriteOnlyCell(sheet, value=csv_safe(value)[:MAX_CELL_CHARS])
    cell.data_type = "s"  # never a formula, whatever the text looks like
    if bold:
        cell.font = Font(bold=True)
    return cell


def render_xlsx(source: RenderInput) -> RenderedFile:
    """An XLSX workbook with one tab per sheet: a bold header row (frozen), text cells, sized columns."""
    workbook = Workbook(write_only=True)
    names = unique_names([s.name for s in source.sheets], clean=lambda n: _SHEET_NAME_BAD.sub(" ", n), limit=31)
    for name, sheet in zip(names, source.sheets, strict=True):
        ws = workbook.create_sheet(title=name)
        ws.freeze_panes = "A2"
        for index, column in enumerate(sheet.columns, start=1):
            widest = max([len(column), *(len(row[index - 1]) for row in sheet.rows[:200] if index - 1 < len(row))])
            ws.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 8), 60)
        ws.append([_text_cell(ws, c, bold=True) for c in sheet.columns])
        for row in sheet.rows:
            ws.append([_text_cell(ws, c) for c in row])
    if not source.sheets:
        workbook.create_sheet(title="Sheet")
    out = io.BytesIO()
    workbook.save(out)
    return RenderedFile(out.getvalue(), "xlsx")


# --- shipped renderers ---------------------------------------------------------------------------------------


def _document(source: RenderInput) -> ReportDocument:
    assert source.document is not None
    return source.document


register_renderer(Renderer("pdf", (FAMILY_DOCUMENT,), lambda s: RenderedFile(render_pdf_document(_document(s)), "pdf")))
register_renderer(Renderer("docx", (FAMILY_DOCUMENT,), lambda s: RenderedFile(render_docx_document(_document(s)), "docx")))
register_renderer(Renderer("html", (FAMILY_DOCUMENT,), lambda s: RenderedFile(render_html_document(_document(s)), "html")))
register_renderer(Renderer("csv", (FAMILY_TABLE,), render_csv_sheets))
register_renderer(Renderer("xlsx", (FAMILY_TABLE,), render_xlsx))

__all__ = [
    "MAX_CELL_CHARS", "RenderInput", "RenderedFile", "Renderer", "formats_for", "register_renderer", "render",
    "render_csv_sheets", "render_xlsx", "unique_names",
]
