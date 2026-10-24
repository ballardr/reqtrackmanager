"""Tests for report cell tones: a cell the screen colours keeps its colour in PDF and Word.

Covers the tone-to-colour table (pinned to the stylesheet's own tints), the section's
shape/name validation, and both renderers applying a tone while leaving untoned cells,
the header and the CSV alone.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.rl_accel import fp_str

from app.services.report_document import ReportBranding
from app.services.report_docx import render_docx_document
from app.services.report_framework import ReportSection
from app.services.report_model import TONE_FILL, ReportDocument, Table, text_on
from app.services.report_pdf import render_pdf_document

THEME = Path(__file__).resolve().parents[2] / "frontend" / "src" / "styles" / "theme.css"


def _document(*, compact: bool = True) -> ReportDocument:
    table = Table(
        ["Name", "Rating"], [["A", "Critical"], ["B", "Low"], ["C", "None"]], compact=compact,
        tones=[[None, "danger"], [None, "accent"], [None, None]],
    )
    return ReportDocument(title="T", blocks=[table], landscape=True)


def _fills(cell) -> str | None:
    shading = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
    return shading.get(qn("w:fill")) if shading is not None else None


# --- Tone table -------------------------------------------------------------------------


def test_tone_fills_match_the_stylesheets_scoring_cell_tints():
    """The print colour of each tone is the on-screen `.scoring-cell--<tone>` tint on the light surface."""
    if not THEME.is_file():
        pytest.skip("the frontend stylesheet is not visible (the backend container mounts only backend/)")
    css = THEME.read_text()
    light = css[css.index(':root[data-theme="light"]'):]
    light = light[: light.index("}")]
    variable = {"muted": "text-muted", "info": "info", "accent": "accent", "warning": "warning", "danger": "danger"}
    assert set(TONE_FILL) == set(variable)
    for tone, name in variable.items():
        base = re.search(rf"--color-{name}:\s*(#[0-9a-fA-F]{{6}})", light).group(1)
        share = int(re.search(rf"\.scoring-cell--{tone}\s*{{[^}}]*?var\(--color-{name}\)\s+(\d+)%", css).group(1)) / 100
        channels = (int(base[i:i + 2], 16) for i in (1, 3, 5))
        expected = "".join(f"{round(c * share + 255 * (1 - share)):02X}" for c in channels)
        assert TONE_FILL[tone] == expected, tone


def test_every_scoring_band_tone_has_a_print_colour():
    from app.modules.registry import SCORING_BAND_TONES

    assert set(TONE_FILL) == set(SCORING_BAND_TONES)


# --- Section validation -----------------------------------------------------------------


def test_section_rejects_tones_not_shaped_like_rows():
    with pytest.raises(ValueError, match="same shape"):
        ReportSection("k", "T", ["a", "b"], [["1", "2"]], tones=[[None]])
    with pytest.raises(ValueError, match="same shape"):
        ReportSection("k", "T", ["a"], [["1"], ["2"]], tones=[[None]])


def test_section_rejects_an_unknown_tone_instead_of_dropping_it_silently():
    with pytest.raises(ValueError, match="unknown tone"):
        ReportSection("k", "T", ["a"], [["1"]], tones=[["magenta"]])
    assert ReportSection("k", "T", ["a"], [["1"]], tones=[["danger"]]).tones == [["danger"]]


def test_tone_at_ignores_missing_out_of_range_and_unknown_tones():
    table = Table(["a"], [["1"]], tones=[["nope"]])
    assert table.tone_at(0, 0) is None and table.tone_at(5, 0) is None and table.tone_at(0, 5) is None
    assert Table(["a"], [["1"]]).tone_at(0, 0) is None


# --- Word -------------------------------------------------------------------------------


@pytest.mark.parametrize("compact", [True, False])
def test_docx_fills_toned_cells_and_leaves_the_rest(compact):
    doc = Document(io.BytesIO(render_docx_document(_document(compact=compact))))
    rows = doc.tables[0].rows
    assert _fills(rows[1].cells[1]) == TONE_FILL["danger"]
    assert _fills(rows[2].cells[1]) == TONE_FILL["accent"]
    assert _fills(rows[3].cells[1]) in (None, "F5F5F5")  # untoned: no fill, or the zebra stripe
    assert _fills(rows[1].cells[0]) != TONE_FILL["danger"]
    assert rows[1].cells[1].text == "Critical"  # the band name stays in the text for a greyscale print
    run = rows[1].cells[1].paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == text_on(TONE_FILL["danger"])


def test_docx_tone_overrides_the_zebra_stripe():
    document = ReportDocument(title="T", blocks=[Table(
        ["a"], [["1"], ["2"]], compact=True, tones=[[None], ["warning"]],
    )])
    rows = Document(io.BytesIO(render_docx_document(document))).tables[0].rows
    shades = rows[2].cells[0]._tc.tcPr.findall(qn("w:shd"))
    assert [s.get(qn("w:fill")) for s in shades] == [TONE_FILL["warning"]]  # one fill, not the stripe as well


def test_docx_header_keeps_the_accent_even_when_every_body_cell_is_toned():
    document = ReportDocument(title="T", blocks=[Table(["a"], [["1"]], compact=True, tones=[["danger"]])])
    rows = Document(io.BytesIO(render_docx_document(document))).tables[0].rows
    assert _fills(rows[0].cells[0]) not in (None, TONE_FILL["danger"])


# --- PDF --------------------------------------------------------------------------------


@pytest.mark.parametrize("compact", [True, False])
def test_pdf_fills_toned_cells(monkeypatch, compact):
    monkeypatch.setattr(rl_config, "pageCompression", 0)  # readable content streams
    pdf = render_pdf_document(_document(compact=compact))

    def fill_operator(tone: str) -> bytes:
        rgb = colors.HexColor(f"#{TONE_FILL[tone]}")
        return f"{fp_str(rgb.red, rgb.green, rgb.blue)} rg".encode()

    assert pdf.count(fill_operator("danger")) >= 1
    assert pdf.count(fill_operator("accent")) >= 1
    assert fill_operator("warning") not in pdf  # a tone nobody used is not painted
    text = colors.HexColor(f"#{text_on(TONE_FILL['danger'])}")
    assert f"{fp_str(text.red, text.green, text.blue)} rg".encode() in pdf


def test_pdf_without_tones_paints_no_tone_fill(monkeypatch):
    monkeypatch.setattr(rl_config, "pageCompression", 0)
    pdf = render_pdf_document(ReportDocument(title="T", blocks=[Table(["a"], [["1"]], compact=True)]))
    for tone, fill in TONE_FILL.items():
        rgb = colors.HexColor(f"#{fill}")
        assert f"{fp_str(rgb.red, rgb.green, rgb.blue)} rg".encode() not in pdf, tone


# --- Text colour by contrast ------------------------------------------------------------


def _ratio(a: str, b: str) -> float:
    def luminance(hex_value: str) -> float:
        def lin(channel: float) -> float:
            return channel / 12.92 if channel <= 0.03928 else ((channel + 0.055) / 1.055) ** 2.4

        r, g, bl = (int(hex_value[i:i + 2], 16) / 255 for i in (0, 2, 4))
        return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(bl)

    hi, lo = sorted((luminance(a), luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


@pytest.mark.parametrize(("fill", "expected"), [
    ("006400", "FFFFFF"),  # dark green: white text
    ("0B3D91", "FFFFFF"),
    ("#475569", "FFFFFF"),  # the default header accent
    ("FFD700", "000000"),  # gold: black text
    ("90EE90", "000000"),  # light green
    ("FFFFFF", "000000"),
    ("000000", "FFFFFF"),
])
def test_text_on_picks_black_or_white_by_contrast(fill, expected):
    assert text_on(fill) == expected


def test_text_on_always_gives_the_better_of_black_and_white():
    for fill in ("2F855A", "B7791F", "C53030", "0E7490", "68D391", "7F7F7F", "808080", "777777"):
        chosen = text_on(fill)
        other = "FFFFFF" if chosen == "000000" else "000000"
        assert _ratio(fill, chosen) >= _ratio(fill, other), fill


def test_every_tone_fill_is_readable_with_its_text():
    for tone, fill in TONE_FILL.items():
        assert _ratio(fill, text_on(fill)) >= 7, tone  # WCAG AAA for normal text


def _header_run_colour(accent: str) -> str:
    document = ReportDocument(title="T", blocks=[Table(["a"], [["1"]], compact=True)], branding=ReportBranding(accent_color_hex=accent))
    rows = Document(io.BytesIO(render_docx_document(document))).tables[0].rows
    assert _fills(rows[0].cells[0]) == accent.lstrip("#").upper()
    return str(rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb)


def test_docx_header_text_follows_the_accent_colour():
    assert _header_run_colour("#0B3D91") == "FFFFFF"  # a dark accent keeps white
    assert _header_run_colour("#FFD700") == "000000"  # a light accent gets black, not unreadable white


@pytest.mark.parametrize("compact", [True, False])
def test_pdf_header_text_follows_the_accent_colour(monkeypatch, compact):
    monkeypatch.setattr(rl_config, "pageCompression", 0)
    white_text = b"1 1 1 rg"

    def render(accent: str) -> bytes:
        document = ReportDocument(
            title="T", blocks=[Table(["Heading"], [], compact=compact)], branding=ReportBranding(accent_color_hex=accent),
        )
        return render_pdf_document(document)

    assert white_text in render("#0B3D91")
    assert white_text not in render("#FFD700")


# --- Badge tones by column --------------------------------------------------------------


def test_tones_by_label_joins_enums_and_rejects_a_label_with_two_tones():
    from app.services.report_framework import tones_by_label

    merged = tones_by_label(({"a": "Draft", "b": "Done"}, {"a": "muted", "b": "accent"}), ({"c": "Draft"}, {"c": "muted"}))
    assert merged == {"Draft": "muted", "Done": "accent"}
    with pytest.raises(ValueError, match="two different tones"):
        tones_by_label(({"a": "Closed"}, {"a": "muted"}), ({"b": "Closed"}, {"b": "accent"}))
    with pytest.raises(ValueError, match="Unknown tone"):
        tones_by_label(({"a": "X"}, {"a": "purple"}))
    assert tones_by_label(({"a": "X"}, {"b": "muted"})) == {}  # a tone for a value with no label is skipped


def test_apply_cell_tones_fills_declared_columns_and_keeps_a_collectors_own_tone():
    from datetime import UTC, datetime

    from app.services.report_framework import ReportResult, apply_cell_tones

    result = ReportResult("k", "T", "S", datetime(2026, 1, 1, tzinfo=UTC), [], [
        ReportSection(
            "a", "A", ["Name", "Status"], [["x", "Approved"], ["y", "Other"], ["z", "Approved"]],
            tones=[[None, None], [None, None], [None, "danger"]],
        ),
        ReportSection("b", "B", ["Status", "Count"], [["Approved", "1"]]),
        ReportSection("c", "C", ["Name"], [["Approved"]]),
    ])
    apply_cell_tones(result, {"Status": {"Approved": "accent"}})
    a, b, c = result.sections
    assert a.tones == [[None, "accent"], [None, None], [None, "danger"]]  # unlisted text plain; a collector's tone wins
    assert b.tones == [["accent", None]]  # every section with the heading
    assert c.tones is None  # no such column: untouched
    assert apply_cell_tones(result, {}) is result


def test_backend_badge_tones_match_the_frontends():
    """Each backend tone map is the hand-kept twin of a frontend `*_TONE` map; they must not drift."""
    from app.modules.compliance.labels import COMPLIANCE_EVIDENCE_VALIDITY_STATE_TONE
    from app.modules.context_strategy.labels import LIFECYCLE_STATUS_TONE, OPEN_QUESTION_STATUS_TONE, PAIN_POINT_STATUS_TONE
    from app.modules.decisions.labels import DECISION_STATUS_TONE
    from app.modules.risk_management.labels import RISK_STATUS_TONE, TREATMENT_ASSESSMENT_STATUS_TONE
    from app.modules.stakeholders.labels import RECORD_STATUS_TONE
    from app.services.labels import CHANGE_REQUEST_STATUS_TONE, REQUIREMENT_STATUS_TONE

    src = THEME.parents[1]
    if not src.is_dir():
        pytest.skip("the frontend sources are not visible (the backend container mounts only backend/)")

    def frontend_map(path: str, name: str) -> dict[str, str]:
        text = (src / path).read_text()
        body = re.search(rf"export const {name}\b[^=]*=\s*{{(.*?)}};", text, re.S).group(1)
        return dict(re.findall(r"(\w+):\s*\"(\w+)\"", body))

    pairs = [
        ("api/types.ts", "REQUIREMENT_STATUS_TONE", REQUIREMENT_STATUS_TONE),
        ("api/types.ts", "CHANGE_REQUEST_STATUS_TONE", CHANGE_REQUEST_STATUS_TONE),
        ("modules/context_strategy/types.ts", "STRATEGY_STATUS_TONE", LIFECYCLE_STATUS_TONE),
        (
            "modules/context_strategy/types.ts", "GUIDING_PRINCIPLE_STATUS_TONE",
            {k: v for k, v in LIFECYCLE_STATUS_TONE.items() if k != "under_review"},  # principles have no review state
        ),
        ("modules/context_strategy/types.ts", "PAIN_POINT_STATUS_TONE", PAIN_POINT_STATUS_TONE),
        ("modules/context_strategy/types.ts", "OPEN_QUESTION_STATUS_TONE", OPEN_QUESTION_STATUS_TONE),
        ("modules/decisions/types.ts", "DECISION_STATUS_TONE", DECISION_STATUS_TONE),
        ("modules/risk_management/types.ts", "RISK_STATUS_TONE", RISK_STATUS_TONE),
        ("modules/risk_management/types.ts", "TREATMENT_ASSESSMENT_STATUS_TONE", TREATMENT_ASSESSMENT_STATUS_TONE),
        ("modules/compliance/types.ts", "COMPLIANCE_EVIDENCE_VALIDITY_STATE_TONE", COMPLIANCE_EVIDENCE_VALIDITY_STATE_TONE),
        ("modules/stakeholders/types.ts", "PERSONA_STATUS_TONE", RECORD_STATUS_TONE),
        ("modules/stakeholders/types.ts", "STAKEHOLDER_STATUS_TONE", RECORD_STATUS_TONE),
        ("modules/stakeholders/types.ts", "NEED_STATUS_TONE", RECORD_STATUS_TONE),
    ]
    for path, name, backend in pairs:
        assert frontend_map(path, name) == backend, name


# --- Cover title on white paper ---------------------------------------------------------


@pytest.mark.parametrize("accent", ["#FFD700", "#90EE90", "#FFFFFF", "#F5F5F5", "#00FFFF"])
def test_readable_on_white_darkens_a_light_accent_until_it_reads(accent):
    from app.services.branding import contrast_ratio, readable_on_white

    shown = readable_on_white(accent)
    assert contrast_ratio(shown, "#ffffff") >= 4.5
    assert shown != accent.lower()


def test_readable_on_white_leaves_a_dark_accent_alone_and_keeps_the_hue():
    from colorsys import rgb_to_hls

    from app.services.branding import readable_on_white

    assert readable_on_white("#0B3D91") == "#0b3d91"
    gold = readable_on_white("#FFD700")
    hue = lambda h: rgb_to_hls(*(int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)))[0]  # noqa: E731
    assert abs(hue(gold) - hue("#FFD700")) < 0.02  # darker, still gold


def _cover_document(accent: str) -> ReportDocument:
    return ReportDocument(
        title="Cover", blocks=[], branding=ReportBranding(accent_color_hex=accent, include_cover_page=True),
    )


def test_docx_cover_title_is_readable_on_the_white_page():
    from app.services.branding import contrast_ratio

    def title_colour(accent: str) -> str:
        doc = Document(io.BytesIO(render_docx_document(_cover_document(accent))))
        run = next(p for p in doc.paragraphs if p.text == "Cover").runs[0]
        return str(run.font.color.rgb)

    assert title_colour("#0B3D91") == "0B3D91"
    assert contrast_ratio(title_colour("#FFD700"), "FFFFFF") >= 4.5


def test_pdf_cover_title_is_not_drawn_in_a_faint_accent(monkeypatch):
    monkeypatch.setattr(rl_config, "pageCompression", 0)
    gold = colors.HexColor("#FFD700")
    gold_operator = f"{fp_str(gold.red, gold.green, gold.blue)} rg".encode()
    assert gold_operator not in render_pdf_document(_cover_document("#FFD700"))
    navy = colors.HexColor("#0B3D91")
    assert f"{fp_str(navy.red, navy.green, navy.blue)} rg".encode() in render_pdf_document(_cover_document("#0B3D91"))
