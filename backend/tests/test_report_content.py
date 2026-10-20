"""Tests for per-report-type report content across the template, project and
organisation tiers :
the resolution chain, ancestor fallback, placeholders, validation, deletion,
the data migration, and export/import round trips (including images).
"""

from __future__ import annotations

import io
import json
import struct
import uuid
import zipfile
import zlib
from contextlib import contextmanager
from datetime import UTC, date, datetime

from docx import Document
from sqlalchemy import select, text

from alembic import command
from app.database import SessionLocal, engine
from app.models.organization import Organization
from app.models.project import Project
from app.models.report_content import ReportContent
from app.services.report_content import (
    ALL_REPORTS_KEY,
    REQUIREMENTS_REPORT_KEY,
    build_placeholders,
    resolve_report_content,
)
from app.services.report_model import Heading, Paragraph, fill_placeholders, markdown_to_blocks
from tests.conftest import (
    auth_headers,
    build_alembic_config,
    create_component_and_category,
    create_org_admin_in,
    create_org_user,
    create_project,
    login,
)

REQ = REQUIREMENTS_REPORT_KEY


def _entry(key=REQ, intro="", chapters=(), appendices=()):
    return {
        "report_key": key, "intro": intro,
        "chapters": [{"title": t, "body": b} for t, b in chapters],
        "appendices": [{"title": t, "body": b} for t, b in appendices],
    }


def _put_org(client, token, org_id, *entries):
    resp = client.put(f"/api/v1/orgs/{org_id}/report-defaults", json={"content": list(entries)}, headers=auth_headers(token))
    assert resp.status_code == 200, resp.text
    return resp.json()


def _put_project(client, token, project_id, *entries, template_id=None):
    resp = client.put(
        f"/api/v1/projects/{project_id}/report-config",
        json={"content": list(entries), "default_report_template_id": template_id}, headers=auth_headers(token),
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def _make_template(client, token, org_id, name, *entries):
    resp = client.post(
        f"/api/v1/orgs/{org_id}/report-templates", json={"name": name, "content": list(entries)}, headers=auth_headers(token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def _effective(client, token, project_id, key=REQ, template_id=None, expect=200):
    params = {"report_key": key, **({"template_id": template_id} if template_id else {})}
    resp = client.get(f"/api/v1/projects/{project_id}/report-content/effective", params=params, headers=auth_headers(token))
    assert resp.status_code == expect, resp.text
    return resp.json()


def _docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)


def _tiny_png() -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    ihdr = chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
    return b"\x89PNG\r\n\x1a\n" + ihdr + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff")) + chunk(b"IEND", b"")


def _rows(scope_id: str) -> list[ReportContent]:
    with SessionLocal() as db:
        return list(db.scalars(select(ReportContent).where(ReportContent.scope_id == uuid.UUID(scope_id))).all())


# --- Resolution -------------------------------------------------------------------


def test_resolution_is_tier_first_then_specific_report_before_all_reports(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    template = _make_template(client, admin_token, org_id, "T", _entry(intro="template-req"))
    _put_org(client, admin_token, org_id, _entry(intro="org-req"), _entry(ALL_REPORTS_KEY, intro="org-all"))

    eff = _effective(client, admin_token, project["id"])
    assert (eff["intro"], eff["intro_source"]["scope_type"], eff["intro_source"]["report_key"]) == ("org-req", "organization", REQ)

    # A project's "all reports" entry beats the organisation's more specific one: the tier is decided first.
    _put_project(client, admin_token, project["id"], _entry(ALL_REPORTS_KEY, intro="proj-all"))
    eff = _effective(client, admin_token, project["id"])
    assert (eff["intro"], eff["intro_source"]["report_key"]) == ("proj-all", ALL_REPORTS_KEY)

    # Within the project tier the specific entry beats "all reports".
    _put_project(client, admin_token, project["id"], _entry(ALL_REPORTS_KEY, intro="proj-all"), _entry(intro="proj-req"))
    assert _effective(client, admin_token, project["id"])["intro"] == "proj-req"

    # A selected template is the most specific tier of all.
    eff = _effective(client, admin_token, project["id"], template_id=template["id"])
    assert (eff["intro"], eff["intro_source"]["scope_type"], eff["intro_source"]["scope_name"]) == ("template-req", "template", "T")


def test_each_field_resolves_independently(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    _put_org(client, admin_token, org_id, _entry(intro="org intro", chapters=[("Org ch", "o")], appendices=[("Org app", "o")]))
    _put_project(client, admin_token, project["id"], _entry(chapters=[("Proj ch", "p")]))

    eff = _effective(client, admin_token, project["id"])
    assert eff["intro"] == "org intro" and eff["intro_source"]["scope_type"] == "organization"
    assert eff["chapters"] == [{"title": "Proj ch", "body": "p"}] and eff["chapters_source"]["scope_type"] == "project"
    assert eff["appendices"] == [{"title": "Org app", "body": "o"}]


def test_a_child_project_inherits_its_nearest_ancestors_content(client, admin_token, org_id):
    root = create_project(client, admin_token, org_id, "Root", can_be_parent=True)
    mid = create_project(client, admin_token, org_id, "Mid", parent_project_id=root["id"], can_be_parent=True)
    leaf = create_project(client, admin_token, org_id, "Leaf", parent_project_id=mid["id"])
    sibling = create_project(client, admin_token, org_id, "Sibling", parent_project_id=root["id"])
    _put_org(client, admin_token, org_id, _entry(intro="org intro"))
    _put_project(client, admin_token, root["id"], _entry(intro="root intro", chapters=[("Root ch", "r")]))
    _put_project(client, admin_token, mid["id"], _entry(intro="mid intro"))

    eff = _effective(client, admin_token, leaf["id"])
    assert eff["intro"] == "mid intro"  # nearest ancestor wins over the root and the organisation
    assert eff["intro_source"]["scope_id"] == mid["id"]
    assert eff["chapters"] == [{"title": "Root ch", "body": "r"}]  # mid has none, so the root's is next
    assert eff["chapters_source"]["scope_name"] == "Root"

    # The child's own row wins and nothing leaks sideways to a sibling.
    _put_project(client, admin_token, leaf["id"], _entry(intro="leaf intro"))
    assert _effective(client, admin_token, leaf["id"])["intro"] == "leaf intro"
    assert _effective(client, admin_token, sibling["id"])["intro"] == "root intro"

    # The ancestors' own rows are untouched by the child's save (edits stay scoped to owned rows).
    assert [r.intro for r in _rows(mid["id"])] == ["mid intro"]
    assert _put_project(client, admin_token, leaf["id"])["content"] == []


def test_ancestor_walk_survives_a_cycle_and_the_organisation_still_applies(client, admin_token, org_id):
    """`get_project_chain_ids` is cycle-safe; a (corrupt) cycle must not hang resolution."""
    a = create_project(client, admin_token, org_id, "Cycle A", can_be_parent=True)
    b = create_project(client, admin_token, org_id, "Cycle B", parent_project_id=a["id"], can_be_parent=True)
    _put_org(client, admin_token, org_id, _entry(intro="org"))
    with engine.begin() as conn:  # the API refuses cycles, so corrupt the data directly
        conn.execute(text("UPDATE projects SET parent_project_id = :b WHERE id = :a"), {"a": a["id"], "b": b["id"]})
    assert _effective(client, admin_token, a["id"])["intro"] == "org"


def test_project_description_stands_in_for_a_blank_requirements_intro_only(client, admin_token, org_id):
    root = create_project(client, admin_token, org_id, "Root", can_be_parent=True)
    child = create_project(client, admin_token, org_id, "Child", parent_project_id=root["id"])
    _put_org(client, admin_token, org_id, _entry(intro="org intro"))
    _put_project(client, admin_token, root["id"], _entry(intro="root intro"))
    client.patch(f"/api/v1/projects/{child['id']}", json={"summary": "The child's own description."}, headers=auth_headers(admin_token))

    eff = _effective(client, admin_token, child["id"])
    # The description is the project's own text, so it beats an ancestor's intro, and is flagged as such.
    assert eff["intro"] == "The child's own description."
    assert eff["intro_source"]["scope_type"] == "project_summary"

    # An explicit intro of its own beats the description.
    _put_project(client, admin_token, child["id"], _entry(intro="explicit"))
    assert _effective(client, admin_token, child["id"])["intro"] == "explicit"

    # A report type other than the requirements one never uses the description.
    with SessionLocal() as db:
        project = db.get(Project, uuid.UUID(child["id"]))
        org = db.get(Organization, project.organization_id)
        other = resolve_report_content(db, "some_module.r1", org, project, None)
        assert other.intro.value == "" and other.intro.source is None


# --- Validation and access ---------------------------------------------------------


def test_unknown_and_repeated_report_types_are_rejected_at_every_tier(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    bad = [_entry("no_such_module.r9", intro="x")]
    dup = [_entry(intro="a"), _entry(intro="b")]
    for entries in (bad, dup):
        assert client.put(f"/api/v1/orgs/{org_id}/report-defaults", json={"content": entries}, headers=auth_headers(admin_token)).status_code == 400
        assert client.put(
            f"/api/v1/projects/{project['id']}/report-config", json={"content": entries}, headers=auth_headers(admin_token),
        ).status_code == 400
        assert client.post(
            f"/api/v1/orgs/{org_id}/report-templates", json={"name": "X", "content": entries}, headers=auth_headers(admin_token),
        ).status_code == 400
    assert _effective(client, admin_token, project["id"], key="no_such_module.r9", expect=400)
    assert _effective(client, admin_token, project["id"], key=ALL_REPORTS_KEY, expect=400)
    # The "all reports" entry itself is valid to store.
    _put_org(client, admin_token, org_id, _entry(ALL_REPORTS_KEY, intro="fine"))


def test_a_blank_entry_or_an_absent_report_type_removes_the_row(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    _put_project(client, admin_token, project["id"], _entry(intro="x"), _entry(ALL_REPORTS_KEY, intro="y"))
    assert len(_rows(project["id"])) == 2
    saved = _put_project(client, admin_token, project["id"], _entry(intro="   "), _entry(ALL_REPORTS_KEY, intro="y"))
    assert [e["report_key"] for e in saved["content"]] == [ALL_REPORTS_KEY]
    assert _put_project(client, admin_token, project["id"])["content"] == []
    assert _rows(project["id"]) == []


def test_oversized_content_is_rejected(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    resp = client.put(
        f"/api/v1/projects/{project['id']}/report-config",
        json={"content": [_entry(intro="x" * 100_001)]}, headers=auth_headers(admin_token),
    )
    assert resp.status_code == 422


def test_other_organisations_cannot_read_or_write_content_and_foreign_templates_are_rejected(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    other_org, other_token = create_org_admin_in(client, admin_token, "Content Isolation Org")
    other_template = _make_template(client, other_token, other_org["id"], "Foreign", _entry(intro="secret"))

    assert client.get(f"/api/v1/orgs/{org_id}/report-defaults", headers=auth_headers(other_token)).status_code == 403
    assert client.put(
        f"/api/v1/orgs/{org_id}/report-defaults", json={"content": [_entry(intro="x")]}, headers=auth_headers(other_token),
    ).status_code == 403
    assert client.get(f"/api/v1/projects/{project['id']}/report-config", headers=auth_headers(other_token)).status_code == 403
    assert client.get(
        f"/api/v1/projects/{project['id']}/report-content/effective", headers=auth_headers(other_token),
    ).status_code == 403
    _effective(client, admin_token, project["id"], template_id=other_template["id"], expect=400)
    assert client.post(
        f"/api/v1/projects/{project['id']}/reports",
        json={"format": "pdf", "report_template_id": other_template["id"]}, headers=auth_headers(admin_token),
    ).status_code == 400


def test_a_plain_stakeholder_can_read_effective_content_but_not_write(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    _put_org(client, admin_token, org_id, _entry(intro="visible"))
    user_id = create_org_user(client, admin_token, org_id, "content-stakeholder@example.com", role="member")
    client.post(f"/api/v1/projects/{project['id']}/roles", json={"user_id": user_id, "role": "stakeholder"}, headers=auth_headers(admin_token))
    token = login(client, "content-stakeholder@example.com", "Password123!")

    assert _effective(client, token, project["id"])["intro"] == "visible"
    types = client.get(f"/api/v1/projects/{project['id']}/report-types", headers=auth_headers(token))
    assert types.status_code == 200 and types.json()[0]["key"] == REQ
    assert client.put(
        f"/api/v1/projects/{project['id']}/report-config", json={"content": [_entry(intro="no")]}, headers=auth_headers(token),
    ).status_code == 403


def test_report_types_list_the_requirements_report_first(client, admin_token, org_id):
    resp = client.get(f"/api/v1/orgs/{org_id}/report-types", headers=auth_headers(admin_token))
    assert resp.status_code == 200
    keys = [t["key"] for t in resp.json()]
    assert keys[0] == REQ
    assert all("." in k for k in keys[1:])  # every module report is "<module>.<report>"


# --- Placeholders ---------------------------------------------------------------------


def test_fill_placeholders_substitutes_the_allow_list_only():
    values = build_placeholders(project_name="Alpha", report_title="Alpha", generated_on=date(2026, 10, 11), organisation="Acme")
    assert fill_placeholders("{{project_name}} / {{ organisation }} / {{generated_on}} / {{report_title}}", values) == (
        "Alpha / Acme / 2026-10-11 / Alpha"
    )
    # Unknown names and other braces stay literal; a name that is not on the list is never substituted.
    assert fill_placeholders("{{nope}} {single} {{project_name", values) == "{{nope}} {single} {{project_name"
    assert fill_placeholders("{{secret}}", {**values, "secret": "leak"}) == "{{secret}}"
    # Single pass: a value containing a placeholder is not expanded again.
    assert fill_placeholders("{{project_name}}", {**values, "project_name": "{{organisation}}"}) == "{{organisation}}"
    assert fill_placeholders("{{project_name}}", None) == "{{project_name}}"


def test_placeholder_values_are_plain_text_not_markdown():
    values = {"project_name": "# Not a heading\n\n![x](attachment:00000000-0000-0000-0000-000000000000)\n- item", "organisation": "Acme"}
    blocks = markdown_to_blocks("Report for {{project_name}} at {{organisation}}", {}, values)
    assert blocks[0] == Paragraph("Report for " + values["project_name"] + " at Acme")
    assert not any(isinstance(b, Heading) for b in blocks)
    assert [type(b).__name__ for b in markdown_to_blocks("# About {{organisation}}", {}, values)] == ["Heading"]


def test_requirements_docx_fills_placeholders_in_resolved_and_ad_hoc_content(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id, "Placeholder <b>Project</b>")
    create_component_and_category(client, admin_token, project["id"])
    org_name = client.get(f"/api/v1/orgs/{org_id}", headers=auth_headers(admin_token)).json()["name"]
    _put_org(client, admin_token, org_id, _entry(
        intro="Intro for {{project_name}} ({{organisation}}) on {{generated_on}} {{unknown}}",
        chapters=[("About {{report_title}}", "Body {{project_name}}")], appendices=[("Appendix", "End of {{project_name}}")],
    ))
    resp = client.post(f"/api/v1/projects/{project['id']}/reports", json={"format": "docx"}, headers=auth_headers(admin_token))
    assert resp.status_code == 200, resp.text
    text_out = _docx_text(resp.content)
    today = datetime.now(UTC).date().isoformat()  # the report stamps the UTC date
    assert f"Intro for Placeholder <b>Project</b> ({org_name}) on" in text_out
    assert "{{unknown}}" in text_out
    assert "About Placeholder <b>Project</b>" in text_out and "End of Placeholder <b>Project</b>" in text_out
    assert today in text_out

    # Ad-hoc text sent by the generation page is substituted too, and wins over the stored content.
    adhoc = client.post(
        f"/api/v1/projects/{project['id']}/reports",
        json={"format": "docx", "pre_markdown": "Ad hoc for {{project_name}}", "post_markdown": "Tail {{organisation}}"},
        headers=auth_headers(admin_token),
    )
    adhoc_text = _docx_text(adhoc.content)
    assert "Ad hoc for Placeholder <b>Project</b>" in adhoc_text and f"Tail {org_name}" in adhoc_text
    assert "Intro for" not in adhoc_text


def test_a_child_projects_report_uses_its_ancestors_chapters(client, admin_token, org_id):
    root = create_project(client, admin_token, org_id, "Report Root", can_be_parent=True)
    child = create_project(client, admin_token, org_id, "Report Child", parent_project_id=root["id"])
    create_component_and_category(client, admin_token, child["id"])
    _put_project(client, admin_token, root["id"], _entry(chapters=[("Inherited chapter", "From the root.")]))
    resp = client.post(f"/api/v1/projects/{child['id']}/reports", json={"format": "docx"}, headers=auth_headers(admin_token))
    assert "Inherited chapter" in _docx_text(resp.content)


# --- Deletion ------------------------------------------------------------------------------


def test_deleting_a_template_removes_its_content_and_leaves_other_scopes(client, admin_token, org_id):
    project = create_project(client, admin_token, org_id)
    template = _make_template(client, admin_token, org_id, "Doomed", _entry(intro="t"))
    _put_project(client, admin_token, project["id"], _entry(intro="p"))
    assert len(_rows(template["id"])) == 1
    assert client.delete(f"/api/v1/orgs/{org_id}/report-templates/{template['id']}", headers=auth_headers(admin_token)).status_code == 204
    assert _rows(template["id"]) == []
    assert len(_rows(project["id"])) == 1


def test_deleting_an_organisation_removes_content_at_all_three_tiers(client, admin_token):
    org, org_token = create_org_admin_in(client, admin_token, "Content Cascade Org")
    sibling, sibling_token = create_org_admin_in(client, admin_token, "Content Cascade Sibling")
    project = create_project(client, org_token, org["id"])
    template = _make_template(client, org_token, org["id"], "T", _entry(intro="t"))
    _put_org(client, org_token, org["id"], _entry(intro="o"))
    _put_project(client, org_token, project["id"], _entry(intro="p"))
    _put_org(client, sibling_token, sibling["id"], _entry(intro="keep"))

    resp = client.request("DELETE", f"/api/v1/orgs/{org['id']}", json={"confirm_name": org["name"]}, headers=auth_headers(admin_token))
    assert resp.status_code == 204, resp.text
    assert _rows(org["id"]) == [] and _rows(project["id"]) == [] and _rows(template["id"]) == []
    assert len(_rows(sibling["id"])) == 1


# --- Migration --------------------------------------------------------------------------------


@contextmanager
def _downgraded_to_0070():
    command.downgrade(build_alembic_config(), "0070")
    try:
        yield
    finally:
        command.upgrade(build_alembic_config(), "head")


def test_the_migration_moves_existing_content_into_requirements_rows_and_back(client, admin_token, org_id):
    filled = create_project(client, admin_token, org_id, "Migrate Filled")
    empty = create_project(client, admin_token, org_id, "Migrate Empty")
    template = _make_template(client, admin_token, org_id, "Migrate T")
    chapter = [{"title": "Ch", "body": "Body"}]

    with _downgraded_to_0070(), engine.begin() as conn:
        # Pre-migration shape: three sets of columns on templates, projects and the organisation.
        conn.execute(
            text("UPDATE report_templates SET intro = 'tpl intro', chapters = CAST(:c AS jsonb) WHERE id = :id"),
            {"c": json.dumps(chapter), "id": template["id"]},
        )
        conn.execute(
            text("UPDATE projects SET report_intro = 'proj intro', report_appendices = CAST(:c AS jsonb) WHERE id = :id"),
            {"c": json.dumps(chapter), "id": filled["id"]},
        )
        conn.execute(
            text("UPDATE organizations SET default_report_intro = 'org intro', default_report_chapters = CAST(:c AS jsonb) WHERE id = :id"),
            {"c": json.dumps(chapter), "id": org_id},
        )

    rows = {r.scope_type: r for scope in (template["id"], filled["id"], org_id) for r in _rows(scope)}
    assert (rows["template"].report_key, rows["template"].intro, rows["template"].chapters) == (REQ, "tpl intro", chapter)
    assert (rows["project"].intro, rows["project"].chapters, rows["project"].appendices) == ("proj intro", [], chapter)
    assert (rows["organization"].intro, rows["organization"].chapters) == ("org intro", chapter)
    assert _rows(empty["id"]) == []  # only non-empty content is copied

    # The downgrade puts the requirements row back into the old columns.
    _put_project(client, admin_token, empty["id"], _entry(intro="round trip"), _entry(ALL_REPORTS_KEY, intro="dropped"))
    with _downgraded_to_0070(), engine.begin() as conn:
        back = conn.execute(text("SELECT report_intro FROM projects WHERE id = :id"), {"id": empty["id"]}).scalar_one()
    assert back == "round trip"


# --- Export / import ---------------------------------------------------------------------------


def _upload_resource_image(client, token, org_id) -> str:
    resp = client.post(
        f"/api/v1/orgs/{org_id}/resources", files={"file": ("pixel.png", _tiny_png(), "image/png")}, headers=auth_headers(token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _set_branding(client, token, org_id, kind, filename, content):
    resp = client.post(
        f"/api/v1/orgs/{org_id}/branding-image", files={"file": (filename, content, "image/png")}, data={"kind": kind},
        headers=auth_headers(token),
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def _admin_of(client, admin_token, org_id) -> str:
    """A token for a fresh org admin of `org_id` (the server admin holds no org role in an imported org)."""
    email = f"admin-{uuid.uuid4().hex[:8]}@example.com"
    create_org_user(client, admin_token, org_id, email, role="org_admin")
    return login(client, email, "Password123!")


def _download(client, token, file_id) -> bytes:
    resp = client.get(f"/api/v1/files/{file_id}", headers=auth_headers(token))
    assert resp.status_code == 200, resp.text
    return resp.content


def test_org_export_import_round_trips_content_images_and_branding_files(client, admin_token):
    org, token = create_org_admin_in(client, admin_token, "Content Export Org")
    project = create_project(client, token, org["id"], "Exported Project")
    image_id = _upload_resource_image(client, token, org["id"])
    logo = _tiny_png() + b"logo"
    background = _tiny_png() + b"background"
    branded = _set_branding(client, token, org["id"], "logo", "logo.png", logo)
    branded = _set_branding(client, token, org["id"], "login_background", "bg.png", background)
    _make_template(client, token, org["id"], "Exported T", _entry(intro="tpl intro", chapters=[("T ch", f"![pix](attachment:{image_id})")]))
    _put_org(client, token, org["id"], _entry(intro="org intro"), _entry(ALL_REPORTS_KEY, intro="org all"))
    _put_project(client, token, project["id"], _entry(intro="proj intro", appendices=[("App", "tail")]))

    bundle = client.get(f"/api/v1/orgs/{org['id']}/export", headers=auth_headers(token))
    assert bundle.status_code == 200
    assert f"files/{image_id}_pixel.png" in zipfile.ZipFile(io.BytesIO(bundle.content)).namelist()

    resp = client.post(
        "/api/v1/orgs/import", data={"name": "Content Import Copy"}, files={"file": ("b.zip", bundle.content, "application/zip")},
        headers=auth_headers(admin_token),
    )
    assert resp.status_code == 201, resp.text
    new_org = resp.json()["organization"]
    assert new_org["logo_file_id"] and new_org["login_background_file_id"]
    assert new_org["logo_file_id"] != branded["logo_file_id"]
    new_token = _admin_of(client, admin_token, new_org["id"])
    assert _download(client, admin_token, new_org["logo_file_id"]) == logo
    assert _download(client, admin_token, new_org["login_background_file_id"]) == background

    defaults = client.get(f"/api/v1/orgs/{new_org['id']}/report-defaults", headers=auth_headers(new_token)).json()["content"]
    assert {e["report_key"]: e["intro"] for e in defaults} == {REQ: "org intro", ALL_REPORTS_KEY: "org all"}

    [new_template] = client.get(f"/api/v1/orgs/{new_org['id']}/report-templates", headers=auth_headers(new_token)).json()
    [tpl_entry] = new_template["content"]
    assert tpl_entry["intro"] == "tpl intro"
    body = tpl_entry["chapters"][0]["body"]
    new_image_id = body.split("attachment:")[1].rstrip(")")
    assert new_image_id != image_id
    assert _download(client, new_token, new_image_id) == _tiny_png()
    resources = client.get(f"/api/v1/orgs/{new_org['id']}/resources", headers=auth_headers(new_token)).json()
    assert [r["id"] for r in resources] == [new_image_id]

    with SessionLocal() as db:  # nobody holds a project role in the imported org yet, so look the project up directly
        new_project_id = db.scalar(select(Project.id).where(
            Project.organization_id == uuid.UUID(new_org["id"]), Project.name == "Exported Project",
        ))
    [own] = _rows(str(new_project_id))
    assert own.intro == "proj intro" and own.appendices == [{"title": "App", "body": "tail"}]


def test_merge_import_brings_template_and_project_content_but_leaves_the_org_profile(client, admin_token):
    source, source_token = create_org_admin_in(client, admin_token, "Content Merge Source")
    target, target_token = create_org_admin_in(client, admin_token, "Content Merge Target")
    project = create_project(client, source_token, source["id"], "Merged Project")
    _make_template(client, source_token, source["id"], "Merged T", _entry(intro="merged tpl"))
    _put_org(client, source_token, source["id"], _entry(intro="source org default"))
    _put_project(client, source_token, project["id"], _entry(intro="merged proj"))
    _put_org(client, target_token, target["id"], _entry(intro="target org default"))

    bundle = client.get(f"/api/v1/orgs/{source['id']}/export", headers=auth_headers(source_token)).content
    resp = client.post(
        f"/api/v1/orgs/{target['id']}/import/merge", files={"file": ("b.zip", bundle, "application/zip")},
        data={"resolutions": "{}"}, headers=auth_headers(target_token),
    )
    assert resp.status_code == 200, resp.text

    [template] = client.get(f"/api/v1/orgs/{target['id']}/report-templates", headers=auth_headers(target_token)).json()
    assert template["content"][0]["intro"] == "merged tpl"
    [merged] = [
        p for p in client.get(f"/api/v1/projects?organization_id={target['id']}", headers=auth_headers(target_token)).json()
        if p["name"] == "Merged Project"
    ]
    merged_config = client.get(f"/api/v1/projects/{merged['id']}/report-config", headers=auth_headers(target_token)).json()
    assert merged_config["content"][0]["intro"] == "merged proj"
    # The target's own organisation-tier content is untouched: merge never writes the org profile.
    defaults = client.get(f"/api/v1/orgs/{target['id']}/report-defaults", headers=auth_headers(target_token)).json()["content"]
    assert [e["intro"] for e in defaults] == ["target org default"]


def test_project_export_import_carries_content_and_remaps_images(client, admin_token):
    source, source_token = create_org_admin_in(client, admin_token, "Project Content Source")
    target, target_token = create_org_admin_in(client, admin_token, "Project Content Target")
    project = create_project(client, source_token, source["id"], "Bundled Project")
    image_id = _upload_resource_image(client, source_token, source["id"])
    _put_project(
        client, source_token, project["id"],
        _entry(intro=f"![pix](attachment:{image_id})", chapters=[("Ch", "body")]), _entry(ALL_REPORTS_KEY, intro="all"),
    )

    bundle = client.get(f"/api/v1/projects/{project['id']}/export", headers=auth_headers(source_token))
    assert bundle.status_code == 200
    resp = client.post(
        "/api/v1/projects/import", data={"organization_id": target["id"], "name": "Bundled Copy"},
        files={"file": ("b.zip", bundle.content, "application/zip")}, headers=auth_headers(target_token),
    )
    assert resp.status_code == 201, resp.text
    new_project = resp.json()["project"]
    config = client.get(f"/api/v1/projects/{new_project['id']}/report-config", headers=auth_headers(target_token)).json()
    entries = {e["report_key"]: e for e in config["content"]}
    assert entries[ALL_REPORTS_KEY]["intro"] == "all" and entries[REQ]["chapters"] == [{"title": "Ch", "body": "body"}]
    new_image_id = entries[REQ]["intro"].split("attachment:")[1].rstrip(")")
    assert new_image_id != image_id
    assert _download(client, target_token, new_image_id) == _tiny_png()

    # Importing the same bundle again reuses the identical shared resource rather than duplicating it.
    client.post(
        "/api/v1/projects/import", data={"organization_id": target["id"], "name": "Bundled Copy Two"},
        files={"file": ("b.zip", bundle.content, "application/zip")}, headers=auth_headers(target_token),
    )
    resources = client.get(f"/api/v1/orgs/{target['id']}/resources", headers=auth_headers(target_token)).json()
    assert [r["id"] for r in resources] == [new_image_id]


def test_bundles_written_before_per_type_content_still_import_and_unknown_types_are_skipped(client, admin_token):
    source, source_token = create_org_admin_in(client, admin_token, "Legacy Bundle Source")
    create_project(client, source_token, source["id"], "Legacy Project")
    bundle = client.get(f"/api/v1/orgs/{source['id']}/export", headers=auth_headers(source_token)).content

    # Rewrite org.json into the old shape (flat org columns, template intro/chapters/appendices)
    # plus one entry for a report type this build does not have.
    src = zipfile.ZipFile(io.BytesIO(bundle))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        for name in src.namelist():
            data = src.read(name)
            if name == "org.json":
                org_json = json.loads(data)
                org_json.pop("report_content", None)
                org_json["default_report_intro"] = "legacy org intro"
                org_json["default_report_chapters"] = [{"title": "Legacy", "body": "x"}]
                org_json["report_templates"] = [
                    {"name": "Legacy T", "intro": "legacy tpl intro", "chapters": [], "appendices": [{"title": "A", "body": "b"}]},
                    {"name": "Future T", "content": [_entry("future_module.r1", intro="gone"), _entry(intro="kept")]},
                ]
                data = json.dumps(org_json).encode()
            zf.writestr(name, data)
    resp = client.post(
        "/api/v1/orgs/import", data={"name": "Legacy Bundle Copy"}, files={"file": ("b.zip", out.getvalue(), "application/zip")},
        headers=auth_headers(admin_token),
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    new_org = body["organization"]
    assert any("future_module.r1" in w for w in body["warnings"])

    new_token = _admin_of(client, admin_token, new_org["id"])
    defaults = client.get(f"/api/v1/orgs/{new_org['id']}/report-defaults", headers=auth_headers(new_token)).json()["content"]
    assert defaults == [_entry(intro="legacy org intro", chapters=[("Legacy", "x")])]
    listed = client.get(f"/api/v1/orgs/{new_org['id']}/report-templates", headers=auth_headers(new_token)).json()
    templates = {t["name"]: t["content"] for t in listed}
    assert templates["Legacy T"][0]["intro"] == "legacy tpl intro" and templates["Legacy T"][0]["appendices"] == [{"title": "A", "body": "b"}]
    assert [e["report_key"] for e in templates["Future T"]] == [REQ]
