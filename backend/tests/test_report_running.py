"""Tests for a document layout's running headers and footers: the tree check (structure, placeholders, fit in the margin),
the text view, the neutral model and the template defaults, the PDF, Word, OpenDocument text and HTML renderers (page
numbers, the first page, mirrored pages, sections), the evaluator through the preview API, and the logo in the bundles.
"""

from __future__ import annotations

import base64
import io
import json
import re
import uuid
import zipfile
import zlib
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from docx import Document as DocxDocument
from PIL import Image as PilImage

from app.services.report_document import ReportBranding
from app.services.report_docx import render_docx_document
from app.services.report_html import render_html_document
from app.services.report_layout_text import format_tree, parse_text
from app.services.report_layout_tree import InvalidLayoutError, assert_valid, empty_tree
from app.services.report_model import (
    Image,
    Margins,
    PageBreak,
    PageField,
    PageStyle,
    PageStyleChange,
    Paragraph,
    ReportDocument,
    RunningBand,
)
from app.services.report_odt import render_odt_document
from app.services.report_pdf import render_pdf_document
from app.services.report_running import apply_default_bands, cell_lines
from tests.conftest import auth_headers
from tests.link_world import World

FILE = "3e5e58a1-f297-4727-971c-c01d4db79697"
HEADER = {"left": "{{project_code}}", "centre": "{{report_title}}", "right": "Page {{page}} of {{pages}}"}


def png(colour: str = "red") -> bytes:
    out = io.BytesIO()
    PilImage.new("RGB", (200, 100), colour).save(out, "PNG")
    return out.getvalue()


def tree_with(**keys: Any) -> dict[str, Any]:
    return {**empty_tree("document"), **keys}


def errors_of(tree: dict[str, Any]) -> list[tuple[str, str]]:
    with pytest.raises(InvalidLayoutError) as caught:
        assert_valid(tree)
    return [(e.path, e.message) for e in caught.value.errors]


# ---- the tree -----------------------------------------------------------------------------------------------------


def test_a_header_and_footer_are_part_of_a_valid_tree_on_the_document_and_on_a_page_style():
    tree = tree_with(
        header={**HEADER, "logo": {"file": FILE, "height": 0.8}, "first_page": "none"}, footer={"centre": "{{generated_on}}"},
        page_styles=[{"name": "Wide", "page_size": "A3", "header": {}, "footer": {"left": "Wide: {{layout_version}}"}}],
    )
    assert_valid(tree)


@pytest.mark.parametrize(("band", "path", "fragment"), [
    ({"colour": "red"}, "header.colour", "Unknown key"),
    ({"left": 3}, "header.left", "Must be text"),
    ({"left": "x" * 301}, "header.left", "Too long"),
    ({"left": "a\nb\nc\nd"}, "header.left", "Too long"),
    ({"first_page": "odd"}, "header.first_page", "same, none"),
    ({"logo": "x"}, "header.logo", "Must be an object"),
    ({"logo": {"file": FILE}}, "header.logo.height", "Required"),
    ({"logo": {"file": FILE, "height": 9}}, "header.logo.height", "from 0.3 to 5"),
    ({"logo": {"file": FILE, "height": 1, "slot": "top"}}, "header.logo.slot", "left, centre, right"),
    ({"left": "{{nonsense}}"}, "header.left", 'Unknown placeholder "{{nonsense}}"'),
    ({"left": "{{req.title}}"}, "header.left", "names a record field"),
    ({"left": "{{var.nope}}"}, "header.left", "not a variable"),
])
def test_a_mistake_in_a_header_is_refused_with_its_path(band, path, fragment):
    errors = dict(errors_of(tree_with(header=band)))
    assert fragment in errors[path]


def test_a_header_taller_than_its_margin_is_refused_at_save_not_clipped():
    # 0.5 cm to the page edge + a 2 cm logo needs a top margin of 2.5 cm; the default is 1.5 cm.
    band = {"logo": {"file": FILE, "height": 2}}
    errors = dict(errors_of(tree_with(header=band)))
    assert "needs a top margin of at least 2.5 cm (it is 1.5 cm)" in errors["header"]
    assert_valid(tree_with(header=band, margins={"top": 3}))
    # A footer is checked against the bottom margin, and three lines need 0.5 + 1.2 cm.
    errors = dict(errors_of(tree_with(footer={"left": "a\nb\nc"}, margins={"bottom": 1})))
    assert "bottom margin of at least 1.7 cm" in errors["footer"]


def test_a_page_style_cannot_inherit_a_header_that_does_not_fit_its_smaller_margin():
    tree = tree_with(header={"logo": {"file": FILE, "height": 0.9}}, page_styles=[{"name": "Tight", "margins": {"top": 1}}])
    errors = dict(errors_of(tree))
    assert "the document's" in errors["page_styles[0].margins.top"]
    assert_valid({**tree, "page_styles": [{"name": "Tight", "margins": {"top": 1}, "header": {}}]})  # its own empty header fits


def test_a_table_layout_has_no_pages_to_put_a_header_on():
    tree = {**empty_tree("table"), "header": {"left": "x"}}
    assert dict(errors_of(tree))["header"] == "Only a document layout has pages."


# ---- the text view ------------------------------------------------------------------------------------------------

TEXT = f"""@layout document
@header left "{{{{project_code}}}}" centre "Line 1\\nLine 2" right "Page {{{{page}}}}" logo "{FILE}" height 0.8 slot centre first none
@footer
@pagestyle "Wide" A3 landscape
@footer left "Wide"

Intro
"""


def test_the_text_view_round_trips_headers_and_footers_on_the_document_and_a_page_style():
    parsed = parse_text(TEXT)
    assert parsed.ok, parsed.errors
    assert parsed.tree["header"] == {
        "left": "{{project_code}}", "centre": "Line 1\nLine 2", "right": "Page {{page}}",
        "logo": {"file": FILE, "height": 0.8, "slot": "centre"}, "first_page": "none",
    }
    assert parsed.tree["footer"] == {}  # "@footer" alone is an empty footer: print none
    assert parsed.tree["page_styles"][0]["footer"] == {"left": "Wide"}
    written = format_tree(parsed.tree)
    assert written.splitlines()[1].startswith('@header left "{{project_code}}"') and written.splitlines()[2] == "@footer"
    assert parse_text(written).tree == parsed.tree


@pytest.mark.parametrize(("text", "fragment"), [
    ('@layout document\n@header left "a" left "b"\n', "left is written once"),
    ('@layout document\n@header colour "a"\n', "Unexpected"),
    ('@layout document\n@header height 1\n', "after the logo"),
    ('@layout document\n@header\n@header\n', "written once"),
    ('@layout document\n\nText\n\n@header\n', "before the first block"),
    ('@layout document\n@header logo\n', "file id"),
])
def test_a_mistake_in_a_header_line_is_a_positioned_error(text, fragment):
    parsed = parse_text(text)
    assert not parsed.ok and any(fragment in e.message for e in parsed.errors)
    assert all(e.line >= 1 for e in parsed.errors)


def test_a_placeholder_error_in_a_page_styles_footer_is_reported_on_its_line():
    parsed = parse_text('@layout document\n@pagestyle "A"\n@footer left "{{nope}}"\n')
    assert [e.line for e in parsed.errors] == [3]


# ---- the model and the template defaults ---------------------------------------------------------------------------


def test_a_cells_lines_split_at_line_breaks_and_keep_page_fields_in_place():
    assert cell_lines(("a\nb", PageField.PAGE, " of ", PageField.PAGES)) == [["a"], ["b", PageField.PAGE, " of ", PageField.PAGES]]
    assert cell_lines(()) == [] and cell_lines(("",)) == []


def test_a_mirrored_band_swaps_its_left_and_right_cells_and_a_logo_in_one_of_them():
    band = RunningBand(left=("L",), centre=("C",), right=("R",), logo=Image(data=png()), logo_slot="left")
    swapped = band.mirrored()
    assert (swapped.left, swapped.centre, swapped.right, swapped.logo_slot) == (("R",), ("C",), ("L",), "right")


def test_a_page_number_footer_is_supplied_only_where_the_page_style_says_nothing():
    own = RunningBand(left=("mine",))
    document = ReportDocument(
        title="T", footer=own, blocks=[Paragraph("a"), PageStyleChange(PageStyle(), section_start=True)],
    )
    result = apply_default_bands(document)
    assert result.footer is own  # the layout's own footer wins
    assert result.header is None  # the organisation's logo is never put in a layout's header for it
    section = result.blocks[1]
    assert isinstance(section, PageStyleChange) and section.style.footer.centre == ("Page ", PageField.PAGE)
    assert section.style.footer.first_page is True


def test_with_nothing_written_a_layout_gets_a_page_number_footer_and_no_header():
    result = apply_default_bands(ReportDocument(title="T"))
    assert result.header is None and result.footer.centre == ("Page ", PageField.PAGE) and result.footer.first_page is True


def test_an_explicitly_empty_band_is_kept_and_prints_nothing():
    empty = RunningBand()
    result = apply_default_bands(ReportDocument(title="T", footer=empty))
    assert result.footer is empty and result.footer.empty


# ---- the renderers ------------------------------------------------------------------------------------------------

BAND = RunningBand(left=("ACME-1",), centre=("Quarterly review",), right=("Page ", PageField.PAGE, " of ", PageField.PAGES))
HIDDEN_FIRST = RunningBand(left=("Header text",), first_page=False)
A3 = PageStyle("A3", True, Margins(), False, HIDDEN_FIRST, BAND)


def document(*blocks, header=None, footer=BAND, **options: Any) -> ReportDocument:
    return ReportDocument(title="T", blocks=list(blocks), header=header, footer=footer, **options)


def streams(pdf: bytes) -> list[str]:
    found = []
    for raw in re.findall(rb"stream\r?\n(.*?)endstream", pdf, re.S):
        try:
            content = zlib.decompress(base64.a85decode(raw.strip(), adobe=True)).decode("latin-1")
        except Exception:  # noqa: BLE001 - a font or the document info, not a page
            continue
        if " Tj" in content:
            found.append(content)
    return found


def page_texts(pdf: bytes) -> list[list[str]]:
    return [re.findall(r"\((.*?)\) Tj", content) for content in streams(pdf)]


def test_the_pdf_prints_the_footer_with_the_page_number_and_the_total_on_every_page():
    pdf = render_pdf_document(document(Paragraph("a"), PageBreak(), Paragraph("b"), PageBreak(), Paragraph("c")))
    pages = page_texts(pdf)
    assert [p[-1] for p in pages] == ["Page 1 of 3", "Page 2 of 3", "Page 3 of 3"]
    assert all("ACME-1" in p and "Quarterly review" in p for p in pages)


def test_the_pdf_skips_a_header_on_the_first_page_and_on_a_sections_first_page_but_not_on_the_return():
    home = PageStyle(header=HIDDEN_FIRST, footer=BAND)
    blocks = [Paragraph("one"), PageBreak(), Paragraph("two"), PageStyleChange(A3), Paragraph("wide"), PageBreak(), Paragraph("wide2"),
              PageStyleChange(home, section_start=False), Paragraph("back")]
    pages = page_texts(render_pdf_document(document(*blocks, header=HIDDEN_FIRST)))
    shown = ["Header text" in p for p in pages]
    assert shown == [False, True, False, True, True]  # page 1, page 3 (section start) and no others skip it
    assert [p[-1] for p in pages] == [f"Page {n} of 5" for n in range(1, 6)]  # numbering continues across sections


def test_the_pdf_draws_the_logo_and_survives_a_logo_it_cannot_read():
    logo = RunningBand(logo=Image(data=png(), decorative=True), logo_height_cm=0.8)
    assert b"/Subtype /Image" in render_pdf_document(document(Paragraph("a"), header=logo))
    bad = RunningBand(logo=Image(data=b"junk", decorative=True), logo_height_cm=0.8)
    assert render_pdf_document(document(Paragraph("a"), header=bad)).startswith(b"%PDF")


def test_the_pdf_swaps_the_cells_on_left_hand_pages_when_the_margins_mirror():
    mirrored = PageStyle(margins=Margins(), mirror=True, footer=RunningBand(left=("INNER",), right=("OUTER",)))
    pages = page_texts(render_pdf_document(ReportDocument(
        title="T", blocks=[Paragraph("a"), PageBreak(), Paragraph("b")], mirror=True, margins=mirrored.margins, footer=mirrored.footer,
    )))
    assert pages[0][-2:] == ["INNER", "OUTER"] and pages[1][-2:] == ["OUTER", "INNER"]


def test_a_plain_report_keeps_the_branding_footer_text_without_a_band():
    pdf = render_pdf_document(ReportDocument(title="T", blocks=[Paragraph("a")], branding=ReportBranding(footer_text="Legacy footer")))
    assert "Legacy footer" in page_texts(pdf)[0]


def package(data: bytes) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(data))


def test_word_has_the_band_as_a_borderless_table_with_live_page_fields():
    pkg = package(render_docx_document(document(Paragraph("a"), header=HIDDEN_FIRST)))
    footer = next(pkg.read(n).decode() for n in pkg.namelist() if n.startswith("word/footer"))
    assert "w:tbl" in footer and 'w:instr="PAGE"' in footer and 'w:instr="NUMPAGES"' in footer and "Quarterly review" in footer
    assert "w:tblBorders" not in footer
    section = DocxDocument(io.BytesIO(render_docx_document(document(Paragraph("a"), header=HIDDEN_FIRST)))).sections[0]
    assert section.different_first_page_header_footer is True
    assert section.first_page_header.paragraphs[0].text == "" and section.header.tables[0].rows[0].cells[0].text == "Header text"


def test_word_gives_each_section_its_own_bands_and_the_first_page_rule_only_to_a_section_start():
    home = PageStyle(header=HIDDEN_FIRST, footer=BAND)
    data = render_docx_document(document(
        Paragraph("a"), PageStyleChange(A3), Paragraph("b"), PageStyleChange(home, section_start=False), Paragraph("c"), header=HIDDEN_FIRST,
    ))
    sections = DocxDocument(io.BytesIO(data)).sections
    assert [s.different_first_page_header_footer for s in sections] == [True, True, False]
    assert all(not s.header.is_linked_to_previous and not s.footer.is_linked_to_previous for s in sections)


def test_word_mirrors_the_bands_with_odd_and_even_headers():
    data = render_docx_document(document(Paragraph("a"), mirror=True, footer=RunningBand(left=("INNER",), right=("OUTER",))))
    doc = DocxDocument(io.BytesIO(data))
    assert doc.settings.odd_and_even_pages_header_footer is True
    section = doc.sections[0]
    assert [c.text for c in section.footer.tables[0].rows[0].cells] == ["INNER", "", "OUTER"]
    assert [c.text for c in section.even_page_footer.tables[0].rows[0].cells] == ["OUTER", "", "INNER"]


def test_a_word_band_with_a_hostile_cell_is_stored_as_text():
    hostile = RunningBand(left=('<w:p>"&</w:p>\x00',))
    pkg = package(render_docx_document(document(Paragraph("a"), footer=hostile)))
    footer = next(pkg.read(n).decode() for n in pkg.namelist() if n.startswith("word/footer"))
    assert "&lt;w:p&gt;" in footer and "\x00" not in footer


def test_opendocument_masters_carry_the_header_the_first_page_the_left_hand_page_and_a_continuation():
    home = PageStyle(header=HIDDEN_FIRST, footer=BAND, mirror=True)
    pkg = package(render_odt_document(document(
        Paragraph("a"), PageStyleChange(A3), Paragraph("b"), PageStyleChange(home, section_start=False), Paragraph("c"),
        header=HIDDEN_FIRST, mirror=True,
    )))
    styles, content = pkg.read("styles.xml").decode(), pkg.read("content.xml").decode()
    names = re.findall(r'<style:master-page style:name="(\w+)"', styles)
    assert names == ["Standard", "Style1", "Continue", "Cover"]
    standard = re.search(r'<style:master-page style:name="Standard".*?</style:master-page>', styles, re.S).group(0)
    for tag in ("header", "header-left", "header-first", "footer", "footer-left", "footer-first"):
        assert f"<style:{tag}>" in standard
    assert "<style:header-first><text:p/></style:header-first>" in standard  # the first page skips the header
    assert "<text:page-count>" in standard and "<text:page-number" in standard
    continuation = re.search(r'<style:master-page style:name="Continue".*?</style:master-page>', styles, re.S).group(0)
    assert "-first>" not in continuation
    assert 'text:style-name="P_Continue"' in content and 'master-page-name="Continue"' in content


def test_the_opendocument_layout_leaves_the_body_where_the_margin_says():
    pkg = package(render_odt_document(document(Paragraph("a"), header=HIDDEN_FIRST, margins=Margins(3, 2.5, 2, 2))))
    styles = pkg.read("styles.xml").decode()
    layout = re.search(r'<style:page-layout style:name="PL0">.*?</style:page-layout>', styles, re.S).group(0)
    # The page margin is the gap to the edge, and the header and footer fill the rest of the margin the layout asked for.
    assert 'fo:margin-top="0.500cm"' in layout and 'fo:margin-bottom="0.500cm"' in layout
    assert 'fo:min-height="2.500cm"' in layout and 'fo:min-height="2.000cm"' in layout
    assert 'fo:margin-top="3.000cm"' in re.search(r'<style:page-layout style:name="PLC">.*?</style:page-layout>', styles, re.S).group(0)


def test_html_shows_the_bands_on_screen_and_in_the_print_margin_boxes_and_escapes_the_text():
    hostile = RunningBand(left=('</style><script>alert("x")</script>',), right=("Page ", PageField.PAGE, " of ", PageField.PAGES), first_page=False)
    page = render_html_document(document(Paragraph("a"), footer=hostile, header=BAND)).decode()
    assert page.count('<div class="band') == 2 and "<script>" not in page and "&lt;script&gt;" in page
    assert "@bottom-right { content: \"Page \" counter(page) \" of \" counter(pages);" in page
    assert "@page:first { @bottom-left { content: none; } }" in page
    assert "</style><script>" not in page and "\\3c " in page  # the angle bracket is escaped inside the CSS string


def test_html_swaps_the_margin_boxes_on_left_hand_pages_when_they_mirror():
    page = render_html_document(document(Paragraph("a"), footer=RunningBand(left=("INNER",), right=("OUTER",)), mirror=True)).decode()
    assert '@page:left { @bottom-right { content: "INNER";' in page and '@page:left { @bottom-left { content: "OUTER";' in page


# ---- the evaluator, the API and the bundles ------------------------------------------------------------------------


@pytest.fixture
def world(client, admin_token) -> World:
    return World(client, admin_token, "Running Org")


def preview(world: World, project: dict, text: str, **headers: str) -> dict:
    resp = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts/preview", json={"text": text},
        headers={**auth_headers(world.token), **headers},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_the_preview_fills_the_placeholders_in_a_header_and_leaves_the_page_fields_to_the_renderer(world):
    project = world.project("Alpha")
    text = (
        '@layout document\n@header left "{{project_name}}" centre "{{organisation}}" right "v{{layout_version}}"\n'
        '@footer left "{{generated_on}}" right "Page {{page}} of {{pages}}"\n\nBody\n'
    )
    html = preview(world, project, text, **{"X-Client-Date": "2026-10-30"})["html"]
    assert ">Alpha<" in html and f">{world.org['name']}<" in html
    assert "2026-10-30" in html and 'title="Filled in when printed"' in html  # the caller's own day, and fields left for printing
    assert "{{" not in html


@pytest.mark.parametrize("offset_days", [-1, 1])
def test_the_run_date_is_the_callers_own_calendar_day_not_the_servers(world, offset_days):
    # 09:00 at UTC+10 is the previous day in UTC and 23:00 at UTC-8 is the next: the browser sends its own day,
    # which the server accepts within a day of its own (relative to today so the test does not age).
    local_date = expected = (datetime.now(UTC).date() + timedelta(days=offset_days)).isoformat()
    project = world.project("Alpha")
    html = preview(world, project, '@layout document\n@footer left "{{generated_on}}"\n\nBody\n', **{"X-Client-Date": local_date})["html"]
    assert expected in html


def test_a_header_value_never_becomes_markup(world):
    project = world.project("Alpha")
    html = preview(world, project, '@layout document\n@header left "<b>x</b> & \\"y\\"" right "{{organisation}}"\n\nBody\n')["html"]
    assert "<b>x</b>" not in html and "&lt;b&gt;x&lt;/b&gt; &amp;" in html


def upload(world: World, url: str, data: bytes, name: str = "logo.png"):
    return world.client.post(url, files={"file": (name, data, "application/octet-stream")}, headers=auth_headers(world.token))


def layout_body(name: str, file_id: str, **keys: Any) -> dict[str, Any]:
    return {
        "name": name, "description": "d", "family": "document",
        "tree": {**empty_tree("document"), "header": {"logo": {"file": file_id, "height": 0.8}}, "margins": {"top": 2}, **keys},
    }


def test_a_logo_follows_the_picture_tier_rule_at_save_and_a_project_logo_cannot_be_in_an_organisation_layout(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    own = upload(world, f"/api/v1/projects/{project['id']}/report-images", png()).json()["id"]
    shared = upload(world, f"/api/v1/orgs/{world.org['id']}/report-images", png("green"), "shared.png").json()["id"]
    url = f"/api/v1/projects/{project['id']}/report-layouts"
    assert world.client.post(url, json=layout_body("Own", own), headers=h).status_code == 201
    assert world.client.post(url, json=layout_body("Shared", shared), headers=h).status_code == 201

    refused = world.client.post(f"/api/v1/orgs/{world.org['id']}/report-layouts", json=layout_body("Org", own), headers=h)
    assert refused.status_code == 422 and "header.logo.file" in refused.text
    assert world.client.post(f"/api/v1/orgs/{world.org['id']}/report-layouts", json=layout_body("Org", shared), headers=h).status_code == 201
    images = world.client.get(f"/api/v1/orgs/{world.org['id']}/report-images", headers=h).json()["images"]
    assert sorted(images[0]["used_by"]) == ["Org", "Shared"]  # a logo counts as a use, so deleting it shows what breaks


def test_a_project_bundle_carries_a_logo_and_rewrites_its_id(world):
    project = world.project("Source")
    h = auth_headers(world.token)
    own = upload(world, f"/api/v1/projects/{project['id']}/report-images", png("purple"), "mark.png").json()["id"]
    body = layout_body("Branded", own, page_styles=[{"name": "Wide", "footer": {"logo": {"file": own, "height": 0.6}}}])
    assert world.client.post(f"/api/v1/projects/{project['id']}/report-layouts", json=body, headers=h).status_code == 201
    bundle = world.client.get(f"/api/v1/projects/{project['id']}/export", headers=h).content
    with zipfile.ZipFile(io.BytesIO(bundle)) as z:
        section = json.loads(z.read("project.json"))["report_layout_images"]
        assert [e["filename"] for e in section] == ["mark.png"]  # one file, though two places use it

    copy = world.client.post(
        "/api/v1/projects/import", data={"organization_id": world.org["id"], "name": f"Copy {uuid.uuid4().hex[:6]}"},
        files={"file": ("bundle.zip", bundle, "application/zip")}, headers=h,
    )
    assert copy.status_code == 201, copy.text
    imported = copy.json()["project"]
    layouts = world.client.get(f"/api/v1/projects/{imported['id']}/report-layouts", headers=h).json()
    detail = world.client.get(f"/api/v1/projects/{imported['id']}/report-layouts/{layouts[0]['id']}", headers=h).json()
    tree = detail["current"]["tree"]
    new_id = tree["header"]["logo"]["file"]
    assert new_id != own and tree["page_styles"][0]["footer"]["logo"]["file"] == new_id
    assert world.client.get(f"/api/v1/files/{new_id}", headers=h).content == png("purple")
