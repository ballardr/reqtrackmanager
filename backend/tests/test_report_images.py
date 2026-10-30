"""Tests for a layout's image block and the picture libraries behind it: the tree check, the text view, sizing, the safe
embedding of a picture (metadata stripped, caps), the PDF, Word, OpenDocument text and HTML renderers, the library API,
the tier visibility rule (save time and run time), downloads and the project and organisation bundles.
"""

from __future__ import annotations

import io
import json
import re
import uuid
import zipfile
from typing import Any

import pytest
from docx import Document as DocxDocument
from PIL import Image as PilImage
from sqlalchemy import select

from app.database import SessionLocal
from app.models.generated_report import GeneratedReport
from app.services.report_docx import render_docx_document
from app.services.report_html import render_html_document
from app.services.report_image_layout import PIXELS_PER_CM, placed_size_cm
from app.services.report_images import (
    ImageRejectedError,
    max_image_bytes,
    max_image_pixels,
    normalise_raster,
    sniff_image,
)
from app.services.report_layout_text import format_tree, parse_text
from app.services.report_layout_tree import InvalidLayoutError, assert_valid, empty_tree
from app.services.report_model import Image, Paragraph, ReportDocument
from app.services.report_odt import render_odt_document
from app.services.report_pdf import render_pdf_document
from tests.conftest import auth_headers, create_org_user, login
from tests.link_world import World

SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50"><rect width="100" height="50" fill="#3366cc"/></svg>'
FILE = "3e5e58a1-f297-4727-971c-c01d4db79697"


def png(width: int = 200, height: int = 100, colour: str = "red") -> bytes:
    out = io.BytesIO()
    PilImage.new("RGB", (width, height), colour).save(out, "PNG")
    return out.getvalue()


def jpeg_with_metadata() -> bytes:
    """A 40x20 JPEG that claims to be rotated and names a camera owner (and so should print as 20x40, with neither)."""
    image = PilImage.new("RGB", (40, 20), "blue")
    exif = image.getexif()
    exif[0x010F] = "Camera Co"
    exif[0x0112] = 6
    out = io.BytesIO()
    image.save(out, "JPEG", exif=exif)
    return out.getvalue()


def image_tree(**keys: Any) -> dict[str, Any]:
    return {**empty_tree("document"), "nodes": [{"type": "image", "file": FILE, "alt": "A picture", **keys}]}


def errors_of(tree: dict[str, Any]) -> list[tuple[str, str]]:
    with pytest.raises(InvalidLayoutError) as caught:
        assert_valid(tree)
    return [(e.path, e.message) for e in caught.value.errors]


# ---- the tree ---------------------------------------------------------------------------------------------------


def test_an_image_block_with_every_property_is_valid():
    assert_valid(image_tree(
        width=8.5, width_unit="cm", height=3, align="left", caption="The {{project_name}} logo", caption_position="above",
        numbering="never", decorative=False,
    ))


def test_a_decorative_image_needs_no_alternative_text_but_any_other_does():
    assert_valid({**empty_tree("document"), "nodes": [{"type": "image", "file": FILE, "decorative": True}]})
    assert errors_of({**empty_tree("document"), "nodes": [{"type": "image", "file": FILE}]})[0][0] == "nodes[0].alt"


@pytest.mark.parametrize(
    ("keys", "path"),
    [
        ({"file": "not-an-id"}, "nodes[0].file"),
        ({"width": 500}, "nodes[0].width"),
        ({"width": 0}, "nodes[0].width"),
        ({"width": 150, "width_unit": "percent"}, "nodes[0].width"),
        ({"width": True}, "nodes[0].width"),
        ({"width_unit": "px"}, "nodes[0].width_unit"),
        ({"width_unit": "cm"}, "nodes[0].width_unit"),
        ({"height": 0.1}, "nodes[0].height"),
        ({"align": "justify"}, "nodes[0].align"),
        ({"caption": "x", "caption_position": "beside"}, "nodes[0].caption_position"),
        ({"numbering": "sometimes"}, "nodes[0].numbering"),
        ({"width": 4, "width_unit": "mm"}, "nodes[0].width"),
        ({"width": 1200, "width_unit": "mm"}, "nodes[0].width"),
        ({"height": 3, "height_unit": "percent"}, "nodes[0].height_unit"),
        ({"height_unit": "mm"}, "nodes[0].height_unit"),
        ({"caption": "{{record.title}}"}, "nodes[0].caption"),
        ({"caption": "x" * 501}, "nodes[0].caption"),
        ({"colour": "red"}, "nodes[0].colour"),
    ],
)
def test_a_bad_image_property_is_refused_with_its_path(keys, path):
    tree = image_tree(**keys)
    if keys.get("file") == "not-an-id":
        tree["nodes"][0]["file"] = "not-an-id"
    assert path in [p for p, _ in errors_of(tree)]


def test_an_image_block_is_a_document_block_only():
    tree = {**empty_tree("table"), "nodes": [{"type": "image", "file": FILE, "alt": "x"}]}
    assert "cannot be used in a table layout" in errors_of(tree)[0][1]


# ---- the text view ----------------------------------------------------------------------------------------------


def test_the_text_view_round_trips_every_image_property_and_leaves_absent_ones_absent():
    tree = {**empty_tree("document"), "nodes": [
        {"type": "image", "file": FILE, "alt": 'Logo "A"', "width": 8.5, "width_unit": "cm", "height": 3, "align": "left",
         "caption": "The {{project_name}} logo", "caption_position": "above", "numbering": "always", "decorative": False},
        {"type": "image", "file": FILE, "decorative": True},
        {"type": "image", "file": FILE, "alt": "a", "width": 50, "width_unit": "percent"},
        {"type": "image", "file": FILE, "alt": "a", "width": 85, "width_unit": "mm", "height": 30, "height_unit": "mm", "numbering": "auto"},
        {"type": "image", "file": FILE, "alt": "a", "width": 5},
    ]}
    text = format_tree(tree)
    parsed = parse_text(text)
    assert parsed.ok, parsed.errors
    assert parsed.tree == tree
    assert format_tree(parsed.tree) == text


@pytest.mark.parametrize(
    ("line", "fragment"),
    [
        (f'@image "{FILE}" alt', "Expected"),
        (f'@image "{FILE}" width wide', "number"),
        (f'@image "{FILE}" decorative maybe', "true or false"),
        (f'@image "{FILE}" alt "a" alt "b"', "once"),
        (f'@image "{FILE}" glow', "Unexpected"),
        ("@image", "Expected"),
    ],
)
def test_a_mistake_in_an_image_line_is_a_positioned_error(line, fragment):
    result = parse_text(f"@layout document\n\n{line}\n")
    assert not result.ok
    assert result.errors[0].line == 3 and fragment in result.errors[0].message


def test_a_checked_problem_in_an_image_block_is_reported_on_its_line():
    result = parse_text(f'@layout document\n\n@image "{FILE}"\n')
    assert not result.ok and result.errors[0].line == 3


# ---- sizing -----------------------------------------------------------------------------------------------------


def block(**keys: Any) -> Image:
    return Image(data=png(), **keys)  # 200 x 100 px: 5.29 x 2.65 cm at 96 dpi


def test_a_picture_with_no_size_prints_at_its_natural_size():
    width, height = placed_size_cm(block(), 17)
    assert width == pytest.approx(200 / PIXELS_PER_CM) and height == pytest.approx(100 / PIXELS_PER_CM)


def test_a_width_keeps_the_shape_a_height_alone_sizes_by_height_and_both_stretch():
    assert placed_size_cm(block(width=8, width_unit="cm"), 17) == pytest.approx((8, 4))
    assert placed_size_cm(block(height_cm=3), 17) == pytest.approx((6, 3))
    assert placed_size_cm(block(width=8, width_unit="cm", height_cm=8), 17) == pytest.approx((8, 8))


def test_a_percentage_is_of_the_text_width_of_the_page_it_is_on():
    assert placed_size_cm(block(width=50, width_unit="percent"), 17) == pytest.approx((8.5, 4.25))
    assert placed_size_cm(block(width=50, width_unit="percent"), 40) == pytest.approx((20, 10))


def test_a_picture_is_never_wider_than_the_text_and_shrinks_in_height_with_it():
    assert placed_size_cm(block(width=30, width_unit="cm"), 17) == pytest.approx((17, 8.5))
    assert placed_size_cm(block(width=30, width_unit="cm", height_cm=30), 17) == pytest.approx((17, 17))


def test_a_picture_that_cannot_be_read_has_no_size():
    assert placed_size_cm(Image(data=b"not an image"), 17) is None


def test_an_svg_is_sized_from_its_own_declared_size():
    width, height = placed_size_cm(Image(svg=SVG, width=10, width_unit="cm"), 17)
    assert (width, height) == pytest.approx((10, 5))


# ---- embedding safely ---------------------------------------------------------------------------------------------


def test_embedding_strips_metadata_and_applies_the_orientation():
    original = jpeg_with_metadata()
    assert b"Camera Co" in original
    normal = normalise_raster(original)
    assert (normal.width_px, normal.height_px) == (20, 40)
    assert normal.content_type == "image/jpeg" and b"Camera Co" not in normal.data
    assert not PilImage.open(io.BytesIO(normal.data)).getexif()


def test_a_gif_keeps_its_first_frame_and_a_transparent_png_keeps_its_transparency():
    frames = [PilImage.new("RGB", (10, 10), c) for c in ("red", "blue")]
    out = io.BytesIO()
    frames[0].save(out, "GIF", save_all=True, append_images=frames[1:])
    assert normalise_raster(out.getvalue()).content_type == "image/png"
    clear = io.BytesIO()
    PilImage.new("RGBA", (4, 4), (0, 0, 0, 0)).save(clear, "PNG")
    assert PilImage.open(io.BytesIO(normalise_raster(clear.getvalue()).data)).mode == "RGBA"


def test_a_picture_is_judged_by_its_bytes_not_its_name_or_declared_type():
    assert sniff_image(png()) == "image/png" and sniff_image(SVG.encode()) == "image/svg+xml"
    for bad in (b"", b"<html><script>alert(1)</script></html>", b"%PDF-1.4", b"GIF89a-but-not"):
        with pytest.raises(ImageRejectedError):
            sniff_image(bad)


def test_an_svg_with_a_script_or_external_reference_is_refused():
    with pytest.raises(ImageRejectedError):
        sniff_image(b'<svg xmlns="http://www.w3.org/2000/svg"><rect style="fill:url(http://evil/x)"/></svg>')
    with pytest.raises(ImageRejectedError):
        sniff_image(b'<!DOCTYPE svg [<!ENTITY x "y">]><svg xmlns="http://www.w3.org/2000/svg"/>')


def test_an_image_with_too_many_pixels_is_refused_before_it_is_decoded(monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "report_image_max_pixels", 1_000_000)
    big = io.BytesIO()
    PilImage.new("1", (max_image_pixels() // 1000 + 1, 1000)).save(big, "PNG")
    with pytest.raises(ImageRejectedError, match="too many pixels"):
        sniff_image(big.getvalue())
    with pytest.raises(ImageRejectedError):
        normalise_raster(big.getvalue())


# ---- the renderers ----------------------------------------------------------------------------------------------


def document_with(*blocks: Image) -> ReportDocument:
    return ReportDocument(title="Pictures", blocks=list(blocks))


CAPTIONED = dict(width=8, width_unit="cm", align="right", caption="Figure 1: A diagram", alt="A red box")


def test_html_shows_a_figure_with_its_width_alignment_alt_text_and_caption_below():
    html = render_html_document(document_with(block(**CAPTIONED))).decode()
    figure = re.search(r'<figure class="image align-right">(.*?)</figure>', html, re.S)
    assert figure and 'alt="A red box"' in figure.group(1) and "width:8cm" in figure.group(1)
    assert figure.group(1).index("<img") < figure.group(1).index("<figcaption>Figure 1: A diagram</figcaption>")


def test_html_puts_the_caption_above_when_asked_and_hides_a_decorative_picture_from_assistive_technology():
    html = render_html_document(document_with(block(caption="Up", caption_above=True, decorative=True))).decode()
    assert html.index("<figcaption>Up") < html.index("<img") and 'alt=""' in html


def test_html_inlines_an_svg_after_sanitising_it_and_labels_it():
    html = render_html_document(document_with(Image(svg=SVG, alt="A chart", width=10, width_unit="cm"))).decode()
    assert '<svg' in html and 'role="img" aria-label="A chart"' in html and "width:10cm" in html


def test_html_text_in_a_caption_or_alt_is_escaped():
    html = render_html_document(document_with(block(caption="<script>x</script>", alt='"><b>'))).decode()
    assert "<script>x" not in html and "&lt;script&gt;" in html and '"><b>' not in html


def test_word_embeds_the_picture_at_its_placed_size_with_alt_text_alignment_and_a_caption_after_it():
    doc = DocxDocument(io.BytesIO(render_docx_document(document_with(block(**CAPTIONED)))))
    shape = doc.inline_shapes[0]
    assert shape.width.cm == pytest.approx(8, abs=0.01) and shape.height.cm == pytest.approx(4, abs=0.01)
    assert shape._inline.docPr.get("descr") == "A red box"
    paragraphs = [p for p in doc.paragraphs]
    picture = next(i for i, p in enumerate(paragraphs) if p._p.xpath(".//w:drawing"))
    assert paragraphs[picture].alignment == 2  # right
    assert paragraphs[picture + 1].text == "Figure 1: A diagram" and paragraphs[picture + 1].style.name == "Caption"


def test_word_puts_a_caption_above_the_picture_when_asked():
    doc = DocxDocument(io.BytesIO(render_docx_document(document_with(block(caption="Top", caption_above=True)))))
    texts = [p.text for p in doc.paragraphs]
    assert texts.index("Top") < next(i for i, p in enumerate(doc.paragraphs) if p._p.xpath(".//w:drawing"))


def test_word_renders_an_svg_as_a_picture_and_skips_an_unreadable_image_without_failing():
    doc = DocxDocument(io.BytesIO(render_docx_document(document_with(Image(svg=SVG, alt="chart"), Image(data=b"junk", caption="x")))))
    assert len(doc.inline_shapes) == 1


def odt_content(document: ReportDocument) -> str:
    with zipfile.ZipFile(io.BytesIO(render_odt_document(document))) as odt:
        return odt.read("content.xml").decode()


def test_opendocument_places_the_picture_at_its_size_with_alt_text_and_a_caption():
    content = odt_content(document_with(block(**CAPTIONED)))
    assert 'svg:width="8.000cm"' in content and 'svg:height="4.000cm"' in content
    assert "<svg:title>A red box</svg:title>" in content
    assert content.index("<draw:frame") < content.index("Figure 1: A diagram")
    assert 'text:style-name="P_Image_right"' in content and 'text:style-name="P_Caption_right"' in content


def test_opendocument_puts_a_caption_above_the_picture_when_asked_and_gives_a_decorative_one_no_title():
    content = odt_content(document_with(block(caption="Top", caption_above=True, decorative=True)))
    assert content.index("Top") < content.index("<draw:frame") and "<svg:title>" not in content


def test_pdf_embeds_each_picture_and_survives_an_unreadable_one():
    pdf = render_pdf_document(document_with(block(**CAPTIONED), Image(svg=SVG, caption="vector"), Image(data=b"junk")))
    assert pdf.startswith(b"%PDF") and pdf.count(b"/Subtype /Image") >= 1


# ---- the library and the run ------------------------------------------------------------------------------------


@pytest.fixture
def world(client, admin_token) -> World:
    return World(client, admin_token, "Report Images Org")


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


def upload(world: World, url: str, data: bytes, name: str = "logo.png", token: str | None = None):
    return world.client.post(url, files={"file": (name, data, "application/octet-stream")}, headers=auth_headers(token or world.token))


def org_images(world: World, suffix: str = "") -> str:
    return f"/api/v1/orgs/{world.org['id']}/report-images{suffix}"


def project_images(project: dict, suffix: str = "") -> str:
    return f"/api/v1/projects/{project['id']}/report-images{suffix}"


def layout_body(name: str, file_id: str, **keys: Any) -> dict[str, Any]:
    return {
        "name": name, "description": "d", "family": "document",
        "tree": {**empty_tree("document"), "nodes": [{"type": "image", "file": file_id, "alt": "A logo", **keys}]},
    }


def test_an_image_is_uploaded_listed_downloaded_and_deleted_in_the_project_library(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    made = upload(world, project_images(project), png())
    assert made.status_code == 201, made.text
    row = made.json()
    assert row["tier"] == "project" and row["owned"] is True and row["content_type"] == "image/png" and row["used_by"] == []

    assert [r["id"] for r in world.client.get(project_images(project), headers=h).json()["images"]] == [row["id"]]
    assert world.client.get(f"/api/v1/files/{row['id']}", headers=h).content == png()
    assert world.client.delete(project_images(project, f"/{row['id']}"), headers=h).status_code == 204
    assert world.client.get(project_images(project), headers=h).json()["images"] == []


def test_the_library_tells_a_picture_in_use_from_one_only_older_versions_use_and_one_nobody_uses(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    now, old, never = (upload(world, project_images(project), png(colour=c), f"{c}.png").json() for c in ("red", "green", "blue"))
    layout = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts", json=layout_body("Report", old["id"]), headers=h,
    ).json()
    saved = world.client.put(
        f"/api/v1/projects/{project['id']}/report-layouts/{layout['id']}/tree",
        json={"tree": layout_body("Report", now["id"])["tree"], "base_version": 1}, headers=h,
    )
    assert saved.status_code == 200, saved.text

    rows = {r["id"]: r for r in world.client.get(project_images(project), headers=h).json()["images"]}

    assert (rows[now["id"]]["used_by"], rows[now["id"]]["in_older_versions"]) == (["Report"], False)
    assert (rows[old["id"]]["used_by"], rows[old["id"]]["in_older_versions"]) == ([], True)
    assert (rows[never["id"]]["used_by"], rows[never["id"]]["in_older_versions"]) == ([], False)


def test_an_upload_is_judged_by_its_bytes_and_a_refusal_says_why(world):
    project = world.project("Alpha")
    resp = upload(world, project_images(project), b"<html><script>1</script></html>", "logo.png")
    assert resp.status_code == 422 and "PNG" in resp.json()["detail"]
    assert upload(world, project_images(project), b"", "empty.png").status_code == 422


def test_a_project_lists_the_organisations_pictures_too_but_can_delete_only_its_own(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    shared = upload(world, org_images(world), png(colour="green")).json()
    own = upload(world, project_images(project), png()).json()

    rows = {r["id"]: r for r in world.client.get(project_images(project), headers=h).json()["images"]}
    assert rows[shared["id"]]["tier"] == "organisation" and rows[shared["id"]]["owned"] is False
    assert rows[own["id"]]["owned"] is True
    assert world.client.delete(project_images(project, f"/{shared['id']}"), headers=h).status_code == 404
    assert [r["id"] for r in world.client.get(org_images(world), headers=h).json()["images"]] == [shared["id"]]
    assert world.client.delete(org_images(world, f"/{own['id']}"), headers=h).status_code == 404


def test_the_library_needs_the_layout_manage_permission(client, world):
    project = world.project("Alpha")
    email = f"plain-{uuid.uuid4().hex[:6]}@example.com"
    create_org_user(client, world.token, world.org["id"], email)
    token = login(client, email, "Password123!")
    assert client.get(project_images(project), headers=auth_headers(token)).status_code in (403, 404)
    assert upload(world, project_images(project), png(), token=token).status_code in (403, 404)
    assert client.get(org_images(world), headers=auth_headers(token)).status_code in (403, 404)
    assert client.get(project_images(project)).status_code in (401, 403)


def test_a_project_picture_is_downloadable_by_the_projects_members_only(client, admin_token, world):
    project = world.project("Alpha")
    other = world.project("Beta")
    row = upload(world, project_images(project), png()).json()
    outsider = f"outsider-{uuid.uuid4().hex[:6]}@example.com"
    create_org_user(client, world.token, world.org["id"], outsider)
    token = login(client, outsider, "Password123!")
    assert client.get(f"/api/v1/files/{row['id']}", headers=auth_headers(token)).status_code == 403
    assert other["id"] != project["id"]


def test_a_project_layout_may_use_its_own_pictures_and_the_organisations(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    own = upload(world, project_images(project), png()).json()["id"]
    shared = upload(world, org_images(world), png(colour="green")).json()["id"]
    url = f"/api/v1/projects/{project['id']}/report-layouts"
    assert world.client.post(url, json=layout_body("Own", own), headers=h).status_code == 201
    assert world.client.post(url, json=layout_body("Shared", shared), headers=h).status_code == 201


def test_an_organisation_layout_may_use_organisation_pictures_only(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    own = upload(world, project_images(project), png()).json()["id"]
    shared = upload(world, org_images(world), png(colour="green")).json()["id"]
    url = f"/api/v1/orgs/{world.org['id']}/report-layouts"
    assert world.client.post(url, json=layout_body("Shared", shared), headers=h).status_code == 201
    refused = world.client.post(url, json=layout_body("Leaky", own), headers=h)
    assert refused.status_code in (400, 422), refused.text
    assert "organisation" in refused.text and "logo.png" in refused.text


def test_a_project_may_not_use_another_projects_picture(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    h = auth_headers(world.token)
    theirs = upload(world, project_images(beta), png()).json()["id"]
    resp = world.client.post(f"/api/v1/projects/{alpha['id']}/report-layouts", json=layout_body("Cross", theirs), headers=h)
    assert resp.status_code in (400, 422), resp.text


def test_a_picture_of_another_organisation_is_refused(client, admin_token, world):
    other = World(client, admin_token, "Other Images Org")
    foreign = upload(other, org_images(other), png()).json()["id"]
    resp = world.client.post(
        f"/api/v1/orgs/{world.org['id']}/report-layouts", json=layout_body("Foreign", foreign), headers=auth_headers(world.token),
    )
    assert resp.status_code in (400, 422), resp.text


def test_the_validate_endpoint_reports_a_picture_the_layout_may_not_use(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    own = upload(world, project_images(project), png()).json()["id"]
    tree = layout_body("x", own)["tree"]
    org = world.client.post(f"/api/v1/orgs/{world.org['id']}/report-layouts/validate", json={"tree": tree}, headers=h).json()
    assert org["ok"] is False and org["errors"][0]["path"] == "nodes[0].file"
    mine = world.client.post(f"/api/v1/projects/{project['id']}/report-layouts/validate", json={"tree": tree}, headers=h).json()
    assert mine["ok"] is True


def preview(world: World, project: dict, tree: dict) -> dict:
    resp = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts/preview", json={"tree": tree}, headers=auth_headers(world.token),
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_a_run_draws_the_picture_with_numbered_captions_in_order(world):
    project = world.project("Alpha")
    first = upload(world, project_images(project), png()).json()["id"]
    second = upload(world, project_images(project), png(colour="green"), "two.png").json()["id"]
    nodes = [
        {"type": "image", "file": first, "alt": "a", "caption": "Context for {{project_name}}"},
        {"type": "image", "file": second, "alt": "b", "caption": "Plain", "numbering": "never"},
        {"type": "image", "file": first, "alt": "c", "caption": "Again"},
        {"type": "image", "file": second, "decorative": True},
    ]
    result = preview(world, project, {**empty_tree("document"), "nodes": nodes})
    captions = re.findall(r"<figcaption>(.*?)</figcaption>", result["html"])
    assert captions == ["Figure 1: Context for Alpha", "Plain", "Figure 2: Again"]
    assert result["html"].count("<img") == 4


def test_numbering_is_automatic_by_default_and_can_be_forced_either_way(world):
    project = world.project("Alpha")
    file_id = upload(world, project_images(project), png()).json()["id"]
    nodes = [
        {"type": "image", "file": file_id, "decorative": True},  # a bare logo: automatic leaves it alone
        {"type": "image", "file": file_id, "decorative": True, "numbering": "always"},  # forced: "Figure 1" with no text
        {"type": "image", "file": file_id, "alt": "a", "caption": "Captioned"},  # automatic: numbered because it has a caption
        {"type": "image", "file": file_id, "alt": "b", "caption": "Kept plain", "numbering": "never"},
        {"type": "image", "file": file_id, "alt": "c", "caption": "Forced", "numbering": "always"},
    ]
    html = preview(world, project, {**empty_tree("document"), "nodes": nodes})["html"]
    assert re.findall(r"<figcaption>(.*?)</figcaption>", html) == ["Figure 1", "Figure 2: Captioned", "Kept plain", "Figure 3: Forced"]


def test_a_size_in_millimetres_is_the_same_length_as_the_same_size_in_centimetres(world):
    project = world.project("Alpha")
    file_id = upload(world, project_images(project), png()).json()["id"]
    nodes = [
        {"type": "image", "file": file_id, "alt": "a", "width": 85, "width_unit": "mm"},
        {"type": "image", "file": file_id, "alt": "b", "width": 8.5},
        {"type": "image", "file": file_id, "alt": "c", "width": 40, "width_unit": "mm", "height": 20, "height_unit": "mm"},
    ]
    html = preview(world, project, {**empty_tree("document"), "nodes": nodes})["html"]
    widths = re.findall(r"<img [^>]*style=\"width:([0-9.]+)cm;[^\"]*aspect-ratio:([0-9./]+)", html)
    assert [w for w, _ in widths] == ["8.5", "8.5", "4"]
    assert widths[2][1] == "4/2"


def test_the_image_size_limit_is_a_deployment_setting_defaulting_to_fifty_megabytes(monkeypatch):
    from app.config import get_settings

    assert max_image_bytes() == 50 * 1024 * 1024
    settings = get_settings()
    monkeypatch.setattr(settings, "report_image_max_bytes", 20)
    assert max_image_bytes() == 20
    with pytest.raises(ImageRejectedError, match="too large"):
        sniff_image(png(100, 100))
    monkeypatch.setattr(settings, "report_image_max_bytes", 500 * 1024 * 1024)  # never above the general upload limit
    assert max_image_bytes() == settings.max_upload_bytes


def test_the_pixel_limit_is_a_deployment_setting(monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "report_image_max_pixels", 50)
    with pytest.raises(ImageRejectedError, match="too many pixels"):
        sniff_image(png(10, 10))
    assert sniff_image(png(7, 7)) == "image/png"


def test_the_library_listing_tells_the_designer_the_size_limit(world, monkeypatch):
    from app.config import get_settings

    project = world.project("Alpha")
    monkeypatch.setattr(get_settings(), "report_image_max_bytes", 7 * 1024 * 1024)
    for url in (project_images(project), org_images(world)):
        assert world.client.get(url, headers=auth_headers(world.token)).json()["max_image_bytes"] == 7 * 1024 * 1024


def test_an_svg_in_the_library_is_drawn_as_a_labelled_vector_figure_in_a_run(world):
    project = world.project("Alpha")
    uploaded = upload(world, project_images(project), SVG.encode(), "diagram.svg")
    assert uploaded.status_code == 201 and uploaded.json()["content_type"] == "image/svg+xml"
    tree = layout_body("Vector", uploaded.json()["id"], width=10, width_unit="mm", caption="A diagram")["tree"]
    html = preview(world, project, tree)["html"]
    assert 'role="img" aria-label="A logo"' in html and "<svg" in html and "width:1cm" in html
    assert "<figcaption>Figure 1: A diagram</figcaption>" in html


def test_an_svg_with_a_script_is_refused_on_upload(world):
    project = world.project("Alpha")
    hostile = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script><rect/></svg>'
    resp = upload(world, project_images(project), hostile, "bad.svg")
    # A script element is dropped by the sanitiser rather than refused, so what is stored and drawn has none of it.
    assert resp.status_code in (201, 422)
    if resp.status_code == 201:
        html = preview(world, project, layout_body("Bad", resp.json()["id"])["tree"])["html"]
        assert "<script" not in html and "alert(1)" not in html


def test_a_markdown_picture_that_is_an_svg_is_drawn_as_a_vector_figure():
    from app.services.report_model import markdown_to_blocks

    blocks = markdown_to_blocks("![chart](attachment:abc)", images={"attachment:abc": SVG.encode()})
    picture = next(b for b in blocks if isinstance(b, Image))
    assert picture.svg == SVG and picture.data == b""
    raster = next(b for b in markdown_to_blocks("![p](attachment:x)", images={"attachment:x": png()}) if isinstance(b, Image))
    assert raster.svg is None and raster.data == png()


def test_a_deleted_picture_is_skipped_with_a_warning_and_the_run_still_succeeds(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    file_id = upload(world, project_images(project), png()).json()["id"]
    tree = layout_body("Gone", file_id)["tree"]
    tree["nodes"].append({"type": "heading", "text": "After"})
    world.client.delete(project_images(project, f"/{file_id}"), headers=h)
    result = preview(world, project, tree)
    assert "After" in result["html"] and "<img" not in result["html"]
    assert any("image" in w.lower() and "deleted" in w.lower() for w in result["warnings"])


def test_a_hand_typed_id_cannot_cross_the_tier_rule_at_run_time(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    theirs = upload(world, project_images(beta), png()).json()["id"]
    result = preview(world, alpha, layout_body("x", theirs)["tree"])  # a preview does not save, so only the run checks
    assert "<img" not in result["html"] and any("may use" in w for w in result["warnings"])


def test_the_library_names_the_layouts_that_use_a_picture(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    file_id = upload(world, project_images(project), png()).json()["id"]
    world.client.post(f"/api/v1/projects/{project['id']}/report-layouts", json=layout_body("Uses it", file_id), headers=h)
    rows = world.client.get(project_images(project), headers=h).json()["images"]
    assert rows[0]["used_by"] == ["Uses it"]


def test_the_picture_a_run_embeds_has_no_metadata(world):
    project = world.project("Alpha")
    file_id = upload(world, project_images(project), jpeg_with_metadata(), "photo.jpg").json()["id"]
    html = preview(world, project, layout_body("Photo", file_id)["tree"])["html"]
    import base64

    embedded = base64.b64decode(re.search(r"data:image/jpeg;base64,([A-Za-z0-9+/=]+)", html).group(1))
    assert b"Camera Co" not in embedded and PilImage.open(io.BytesIO(embedded)).size == (20, 40)


def test_the_run_records_the_pictures_it_used_by_id_and_count_never_bytes(world, db):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    file_id = upload(world, project_images(project), png()).json()["id"]
    tree = {**empty_tree("document"), "nodes": [{"type": "image", "file": file_id, "alt": "a"}, {"type": "image", "file": file_id, "alt": "b"}]}
    created = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts",
        json={"name": "Prov", "description": "d", "family": "document", "tree": tree}, headers=h,
    ).json()
    run = world.client.get(f"/api/v1/projects/{project['id']}/custom-reports/{created['id']}/run", headers=h)
    assert run.status_code == 200, run.text
    row = db.scalars(select(GeneratedReport).where(GeneratedReport.layout_id == uuid.UUID(created["id"]))).one()
    assert row.stats["images"] == 2 and row.stats["image_files"] == [file_id]


# ---- bundles ------------------------------------------------------------------------------------------------------


def export_project(world: World, project: dict) -> bytes:
    resp = world.client.get(f"/api/v1/projects/{project['id']}/export", headers=auth_headers(world.token))
    assert resp.status_code == 200, resp.text
    return resp.content


def import_project(world: World, bundle: bytes, name: str = "Imported") -> dict:
    resp = world.client.post(
        "/api/v1/projects/import", data={"organization_id": world.org["id"], "name": name},
        files={"file": ("bundle.zip", bundle, "application/zip")}, headers=auth_headers(world.token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def test_a_project_bundle_carries_the_pictures_its_layouts_use_and_rewrites_their_ids(world):
    project = world.project("Source")
    h = auth_headers(world.token)
    own = upload(world, project_images(project), png(colour="purple"), "own.png").json()["id"]
    shared = upload(world, org_images(world), png(colour="green"), "shared.png").json()["id"]
    url = f"/api/v1/projects/{project['id']}/report-layouts"
    world.client.post(url, json=layout_body("Own picture", own), headers=h)
    world.client.post(url, json=layout_body("Shared picture", shared), headers=h)

    bundle = export_project(world, project)
    with zipfile.ZipFile(io.BytesIO(bundle)) as z:
        section = json.loads(z.read("project.json"))["report_layout_images"]
        assert sorted(e["filename"] for e in section) == ["own.png", "shared.png"]
        assert any(n.startswith("files/") and n.endswith("own.png") for n in z.namelist())

    copy = import_project(world, bundle)
    assert copy["warnings"] == []
    imported = copy["project"]
    mine = {r["id"]: r for r in world.client.get(project_images(imported), headers=h).json()["images"]}
    assert [r["filename"] for r in mine.values() if r["owned"]] == ["own.png"]  # the picture copied into the new project
    layouts = {r["name"]: r for r in world.client.get(f"/api/v1/projects/{imported['id']}/report-layouts", headers=h).json()}
    own_copy = world.client.get(f"/api/v1/projects/{imported['id']}/report-layouts/{layouts['Own picture']['id']}", headers=h).json()
    new_id = own_copy["current"]["tree"]["nodes"][0]["file"]
    assert new_id != own and mine[new_id]["owned"] is True
    assert world.client.get(f"/api/v1/files/{new_id}", headers=h).content == png(colour="purple")
    shared_copy = world.client.get(f"/api/v1/projects/{imported['id']}/report-layouts/{layouts['Shared picture']['id']}", headers=h).json()
    assert shared_copy["current"]["tree"]["nodes"][0]["file"] == shared  # an organisation picture stays a reference


def test_an_organisation_bundle_carries_organisation_pictures_into_the_new_organisations_library(client, admin_token, world):
    h = auth_headers(world.token)
    shared = upload(world, org_images(world), png(colour="orange"), "brand.png").json()["id"]
    world.client.post(f"/api/v1/orgs/{world.org['id']}/report-layouts", json=layout_body("Branded", shared), headers=h)
    bundle = client.get(f"/api/v1/orgs/{world.org['id']}/export", headers=h).content

    resp = client.post(
        "/api/v1/orgs/import", data={"name": f"Copy {uuid.uuid4().hex[:6]}"},
        files={"file": ("bundle.zip", bundle, "application/zip")}, headers=auth_headers(admin_token),
    )
    assert resp.status_code == 201, resp.text
    new_org = resp.json()["organization"]["id"]
    images = client.get(f"/api/v1/orgs/{new_org}/report-images", headers=h).json()["images"]
    assert [i["filename"] for i in images] == ["brand.png"] and images[0]["id"] != shared
    layouts = client.get(f"/api/v1/orgs/{new_org}/report-layouts", headers=h).json()
    detail = client.get(f"/api/v1/orgs/{new_org}/report-layouts/{layouts[0]['id']}", headers=h).json()
    assert detail["current"]["tree"]["nodes"][0]["file"] == images[0]["id"]
    assert images[0]["used_by"] == ["Branded"]


def test_a_bundle_whose_picture_bytes_are_missing_drops_that_version_with_a_warning(world):
    project = world.project("Source")
    h = auth_headers(world.token)
    own = upload(world, project_images(project), png(), "gone.png").json()["id"]
    world.client.post(f"/api/v1/projects/{project['id']}/report-layouts", json=layout_body("Needs it", own), headers=h)
    bundle = export_project(world, project)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(bundle)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            if not (item.filename.startswith("files/") and item.filename.endswith("gone.png")):
                dst.writestr(item, src.read(item.filename))
    copy = import_project(world, out.getvalue())
    joined = " ".join(copy["warnings"])
    assert "gone.png" in joined and "Needs it" in joined


# ---- placement: wrap, behind and in front ----------------------------------------------------------------------


def pdf_content(pdf: bytes) -> str:
    """The page content streams of a ReportLab PDF (ASCII85 + Flate), joined, with image draws shown as `<IMG>`."""
    import base64
    import zlib

    out = []
    for stream in re.findall(rb"stream\r?\n(.*?)endstream", pdf, re.S):
        body = stream.strip()
        try:
            out.append(zlib.decompress(base64.a85decode(body[:-2] if body.endswith(b"~>") else body)).decode("latin1"))
        except Exception:  # noqa: BLE001 - an image or font stream: not content
            continue
    return re.sub(r"/FormXob\.\w+ Do", "<IMG>", "\n".join(s for s in out if "BT" in s or "Do" in s))


def layered(**keys: Any) -> Image:
    return Image(**{"data": png(), "alt": "A picture", "width": 4, **keys})


@pytest.mark.parametrize(
    "keys",
    [
        {"placement": "wrap", "align": "left"},
        {"placement": "wrap", "align": "right", "width": 4},
        {"placement": "behind"},
        {"placement": "front", "anchor": "margin", "offset_x": -1.5, "offset_y": 0},
        {"placement": "behind", "anchor": "page", "offset_x": 2, "offset_y": 3},
        {"placement": "inline", "align": "center", "caption": "Still captioned"},
    ],
)
def test_every_placement_mode_with_its_own_settings_is_valid(keys):
    assert_valid(image_tree(**keys))


@pytest.mark.parametrize(
    ("keys", "path"),
    [
        ({"placement": "floating"}, "nodes[0].placement"),
        ({"placement": "wrap"}, "nodes[0].align"),  # no side chosen: the default is centred
        ({"placement": "wrap", "align": "center"}, "nodes[0].align"),
        ({"placement": "wrap", "align": "left", "caption": "x"}, "nodes[0].caption"),
        ({"placement": "wrap", "align": "left", "numbering": "always"}, "nodes[0].numbering"),
        ({"placement": "wrap", "align": "left", "offset_x": 1}, "nodes[0].offset_x"),
        ({"placement": "wrap", "align": "left", "anchor": "page"}, "nodes[0].anchor"),
        ({"placement": "behind", "align": "left"}, "nodes[0].align"),
        ({"placement": "front", "caption": "x"}, "nodes[0].caption"),
        ({"placement": "behind", "anchor": "sheet"}, "nodes[0].anchor"),
        ({"placement": "behind", "offset_x": 101}, "nodes[0].offset_x"),
        ({"placement": "behind", "offset_y": "far"}, "nodes[0].offset_y"),
        ({"placement": "behind", "offset_y": True}, "nodes[0].offset_y"),
        ({"offset_x": 1}, "nodes[0].offset_x"),
        ({"anchor": "page"}, "nodes[0].anchor"),
    ],
)
def test_a_placement_setting_its_mode_does_not_use_is_refused_with_its_path(keys, path):
    assert path in [p for p, _ in errors_of(image_tree(**keys))]


def test_the_text_view_round_trips_placement_anchor_and_offsets():
    tree = {**empty_tree("document"), "nodes": [
        {"type": "image", "file": FILE, "alt": "a", "placement": "wrap", "align": "right", "width": 4},
        {"type": "image", "file": FILE, "alt": "a", "placement": "behind", "anchor": "page", "offset_x": 2.5, "offset_y": -3},
        {"type": "image", "file": FILE, "alt": "a", "placement": "front", "offset_x": 0, "offset_y": 1},
    ]}
    text = format_tree(tree)
    assert "placement behind anchor page x 2.5 y -3" in text
    parsed = parse_text(text)
    assert parsed.ok, parsed.errors
    assert parsed.tree == tree and format_tree(parsed.tree) == text


@pytest.mark.parametrize(
    ("clause", "fragment"),
    [("x 1 x 2", "x is written once"), ("y far", "down distance"), ("x", "across distance")],
)
def test_a_bad_position_clause_is_a_positioned_error(clause, fragment):
    parsed = parse_text(f'@layout document\n@image "{FILE}" alt "a" placement behind {clause}')
    assert not parsed.ok and fragment in parsed.errors[0].message


# PDF: a wrapped picture shares lines with its text; a layered one is drawn first (behind) or last (in front).


def test_pdf_runs_the_text_beside_a_wrapped_picture_then_across_the_page():
    text = "word " * 400
    wrapped = render_pdf_document(document_with(Image(data=png(), alt="p", width=5, placement="wrap", align="right"), Paragraph(text)))
    plain = render_pdf_document(document_with(Image(data=png(), alt="p", width=5), Paragraph(text)))

    def first_lines(pdf: bytes) -> list[int]:
        return [len(m) for m in re.findall(r"\((word[ word]*?)\) Tj", pdf_content(pdf))]

    beside = first_lines(wrapped)
    assert beside[0] < max(beside)  # narrower lines beside the picture, then full-width ones
    assert len(set(first_lines(plain))) == 1  # in line: every line is full width
    assert wrapped.count(b"/Subtype /Image") == 1


def test_pdf_wrap_carries_on_over_a_page_break_and_ends_at_a_table():
    from app.services.report_model import Table

    blocks = [Image(data=png(300, 1200), alt="p", width=5, placement="wrap", align="left"), Paragraph("word " * 3000),
              Table(columns=["A"], rows=[["cell"]]), Paragraph("after")]
    pdf = render_pdf_document(document_with(*blocks))
    assert pdf.count(b"/Type /Page\n") >= 2 and b"(after)" in pdf_content(pdf).encode("latin1")


def test_pdf_draws_a_picture_behind_the_text_before_it_and_one_in_front_after_it():
    pdf = render_pdf_document(document_with(
        Paragraph("first words"), layered(placement="front", anchor="page", offset_x_cm=1, offset_y_cm=1),
        layered(placement="behind", anchor="page", offset_x_cm=5, offset_y_cm=5), Paragraph("last words"),
    ))
    content = pdf_content(pdf)
    behind, front = content.index("<IMG>"), content.rindex("<IMG>")
    assert content.count("<IMG>") == 2 and behind < content.index("first words") and front > content.index("last words")


@pytest.mark.parametrize(
    ("anchor", "x_pt", "y_pt"),
    # A4 is 595.28 x 841.89 pt; the default margin is 1.5 cm (42.52 pt); the picture is 4 cm x 2 cm (113.39 x 56.69 pt).
    [("page", 56.69, 841.89 - 85.04 - 56.69), ("margin", 42.52 + 56.69, 841.89 - 42.52 - 85.04 - 56.69)],
)
def test_pdf_measures_a_layered_picture_from_the_page_or_margin_corner(anchor, x_pt, y_pt):
    pdf = render_pdf_document(document_with(layered(placement="behind", anchor=anchor, offset_x_cm=2, offset_y_cm=3), Paragraph("text")))
    match = re.search(r"([\d.]+) 0 0 ([\d.]+) ([\d.]+) ([\d.]+) cm\n<IMG>", pdf_content(pdf))
    assert match, pdf_content(pdf)[:400]
    assert (float(match.group(3)), float(match.group(4))) == pytest.approx((x_pt, y_pt), abs=0.1)


def test_pdf_measures_a_paragraph_anchored_picture_from_where_the_block_sits():
    pdf = render_pdf_document(document_with(
        Paragraph("above"), layered(placement="front", anchor="paragraph", offset_x_cm=1, offset_y_cm=0), Paragraph("below"),
    ))
    content = pdf_content(pdf)
    match = re.search(r"([\d.]+) 0 0 [\d.]+ ([\d.]+) ([\d.]+) cm\n<IMG>", content)
    assert match and float(match.group(2)) == pytest.approx(48.52 + 28.35, abs=0.1)  # frame left + 1 cm
    top = float(match.group(3)) + 56.69  # the picture's top edge: where the block sits, just under "above", not the page corner
    assert 700 < top < 760


def test_pdf_leaves_out_a_layered_or_wrapped_picture_it_cannot_draw():
    pdf = render_pdf_document(document_with(
        Image(data=b"junk", placement="behind"), Image(data=b"junk", placement="wrap", align="left"), Paragraph("kept"),
    ))
    assert pdf.startswith(b"%PDF") and b"(kept)" in pdf_content(pdf).encode("latin1")


# Word: floating drawings with the wrap, layer and anchor each mode names.


def word_anchors(document: ReportDocument) -> list[Any]:
    doc = DocxDocument(io.BytesIO(render_docx_document(document)))
    return doc.element.body.xpath(".//wp:anchor")


def test_word_wraps_text_around_a_picture_anchored_to_the_column_on_its_side():
    (anchor,) = word_anchors(document_with(Image(data=png(), alt="A box", width=5, placement="wrap", align="left"), Paragraph("text")))
    assert anchor.xpath("./wp:wrapSquare") and anchor.get("behindDoc") == "0"
    assert anchor.xpath("string(./wp:positionH/@relativeFrom)") == "column" and anchor.xpath("string(./wp:positionH/wp:align)") == "left"
    assert anchor.xpath("string(./wp:docPr/@descr)") == "A box"


@pytest.mark.parametrize(("placement", "behind"), [("behind", "1"), ("front", "0")])
@pytest.mark.parametrize("anchor_name", ["paragraph", "margin", "page"])
def test_word_layers_a_picture_with_its_anchor_and_offsets_in_emu(placement, behind, anchor_name):
    (anchor,) = word_anchors(document_with(layered(placement=placement, anchor=anchor_name, offset_x_cm=2, offset_y_cm=-1.5)))
    assert anchor.get("behindDoc") == behind and anchor.xpath("./wp:wrapNone")
    assert anchor.xpath("string(./wp:positionH/@relativeFrom)") == anchor_name == anchor.xpath("string(./wp:positionV/@relativeFrom)")
    assert anchor.xpath("string(./wp:positionH/wp:posOffset)") == "720000" and anchor.xpath("string(./wp:positionV/wp:posOffset)") == "-540000"


def test_word_keeps_the_anchor_paragraph_tiny_and_with_the_text_and_orders_the_layers():
    doc = DocxDocument(io.BytesIO(render_docx_document(document_with(
        layered(placement="behind"), layered(placement="front"), Paragraph("text"),
    ))))
    anchors = doc.element.body.xpath(".//wp:anchor")
    heights = [int(a.get("relativeHeight")) for a in anchors]
    assert heights == sorted(set(heights)) and len(heights) == 2
    holder = anchors[0].getparent().getparent().getparent()
    assert holder.xpath("string(./w:pPr/w:keepNext/@w:val)") in ("", "1", "true") and holder.xpath("./w:pPr/w:keepNext")


def test_word_floating_pictures_are_not_inline_shapes_and_a_bad_one_is_skipped():
    doc = DocxDocument(io.BytesIO(render_docx_document(document_with(
        layered(placement="behind"), Image(data=b"junk", placement="wrap", align="right"), Paragraph("text"),
    ))))
    assert len(doc.inline_shapes) == 0 and len(doc.element.body.xpath(".//wp:anchor")) == 1


# OpenDocument text: frames whose graphic style carries the wrap, layer and anchor.


def odt_styles(document: ReportDocument) -> tuple[str, str]:
    with zipfile.ZipFile(io.BytesIO(render_odt_document(document))) as odt:
        content = odt.read("content.xml").decode()
    return content, content[: content.index("</office:automatic-styles>")]


def test_opendocument_wraps_text_in_parallel_around_a_frame_on_the_picture_side():
    content, styles = odt_styles(document_with(Image(data=png(), alt="A box", width=5, placement="wrap", align="right"), Paragraph("t")))
    assert 'draw:style-name="fr_wrap_right"' in content and 'text:anchor-type="paragraph"' in content
    style = re.search(r'name="fr_wrap_right".*?</style:style>', styles, re.S).group(0)
    assert 'style:wrap="parallel"' in style and 'style:horizontal-pos="right"' in style
    assert "<svg:title>A box</svg:title>" in content


@pytest.mark.parametrize(("placement", "run_through"), [("behind", "background"), ("front", "foreground")])
@pytest.mark.parametrize(("anchor", "rel"), [("paragraph", "paragraph"), ("margin", "page-content"), ("page", "page")])
def test_opendocument_layers_a_frame_with_its_anchor_and_offsets(placement, run_through, anchor, rel):
    content, styles = odt_styles(document_with(layered(placement=placement, anchor=anchor, offset_x_cm=2, offset_y_cm=3)))
    name = f"fr_{placement}_{anchor}"
    assert f'draw:style-name="{name}"' in content and 'svg:x="2.000cm" svg:y="3.000cm"' in content
    style = re.search(rf'name="{name}".*?</style:style>', styles, re.S).group(0)
    assert f'style:run-through="{run_through}"' in style and 'style:wrap="run-through"' in style
    assert f'style:horizontal-rel="{rel}"' in style and f'style:vertical-rel="{rel}"' in style


def test_opendocument_floating_frames_are_in_a_tiny_paragraph_kept_with_the_next_and_a_bad_one_is_skipped():
    content, styles = odt_styles(document_with(layered(placement="behind"), Image(data=b"junk", placement="wrap", align="left")))
    assert content.count("<draw:frame") == 1 and 'text:style-name="P_Float"' in content
    assert 'fo:keep-with-next="always"' in re.search(r'name="P_Float".*?</style:style>', styles, re.S).group(0)


# HTML: a float for a wrapped picture, a positioned layer for the others.


def test_html_floats_a_wrapped_picture_to_the_side_it_is_aligned_to():
    html = render_html_document(document_with(Image(data=png(), alt="A box", width=5, placement="wrap", align="left"), Paragraph("t"))).decode()
    assert '<figure class="image wrap-left" style="width:5cm;max-width:60%"><img' in html and "figure.image.wrap-left { float: left" in html
    assert "<figcaption>" not in html


@pytest.mark.parametrize("placement", ["behind", "front"])
def test_html_lays_a_layered_picture_over_the_text_at_its_offsets(placement):
    html = render_html_document(document_with(layered(placement=placement, anchor="page", offset_x_cm=2, offset_y_cm=-1.5))).decode()
    assert f'<div class="layer-anchor {placement}"><figure class="image" style="left:2cm;top:-1.5cm"><img' in html
    assert "div.layer-anchor.behind > figure.image { z-index: -1; }" in html and "body { isolation: isolate;" in html


def test_html_text_in_a_layered_pictures_alt_is_escaped():
    html = render_html_document(document_with(layered(placement="front", alt='"><script>x</script>'))).decode()
    assert "<script>x" not in html and "&lt;script&gt;" in html


# A run: the evaluator carries the settings into the document and gives a floating picture no gap.


def test_a_run_carries_placement_anchor_and_offsets_into_every_format(world):
    project = world.project("Alpha")
    file_id = upload(world, project_images(project), png()).json()["id"]
    nodes = [
        {"type": "image", "file": file_id, "alt": "a", "placement": "wrap", "align": "right", "width": 4},
        {"type": "text", "text": "Beside the picture."},
        {"type": "image", "file": file_id, "alt": "b", "placement": "behind", "anchor": "page", "offset_x": 2, "offset_y": 3, "width": 6},
    ]
    html = preview(world, project, {**empty_tree("document"), "nodes": nodes})["html"]
    assert 'class="image wrap-right"' in html and '<div class="layer-anchor behind"><figure class="image" style="left:2cm;top:3cm"' in html


def test_saving_a_layout_with_a_caption_on_a_floating_picture_is_refused_with_the_path(world):
    project = world.project("Alpha")
    file_id = upload(world, project_images(project), png()).json()["id"]
    body = layout_body("Bad", file_id, placement="behind", caption="no")
    resp = world.client.post(f"/api/v1/projects/{project['id']}/report-layouts", json=body, headers=auth_headers(world.token))
    assert resp.status_code == 422 and "nodes[0].caption" in resp.text
