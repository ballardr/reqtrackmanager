"""Tests for the renderer registry and the table and HTML renderers (`services.report_render`,
`services.report_html`): families and formats, CSV and its zip, XLSX cell safety and sheet naming, HTML
escaping, tones, branding and page breaks. No database."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import UTC, date, datetime, timedelta, timezone
from decimal import Decimal

import openpyxl
import pytest

from app.services import report_render
from app.services.report_document import ReportBranding
from app.services.report_html import CSP, render_html_document
from app.services.report_model import (
    TONE_FILL,
    BulletList,
    CellLink,
    DataCell,
    Gap,
    Heading,
    Image,
    LinkedText,
    Note,
    PageBreak,
    Paragraph,
    ReportDocument,
    Sheet,
    Table,
    typed_cell,
)
from app.services.report_render import MAX_CELL_CHARS, RenderedFile, Renderer, RenderInput, formats_for, render, unique_names


def test_formats_are_read_from_the_registry_per_family():
    assert formats_for("document") == ["pdf", "docx", "html", "odt"]
    assert formats_for("table") == ["csv", "xlsx", "ods"]
    assert formats_for("nothing") == []


def test_a_format_is_not_rendered_for_a_family_it_does_not_serve():
    with pytest.raises(KeyError):
        render("pdf", "table", RenderInput("t", sheets=[Sheet("s", ["a"], [])]))
    with pytest.raises(KeyError):
        render("xlsx", "document", RenderInput("t", document=ReportDocument("t")))
    with pytest.raises(KeyError):
        render("ods", "document", RenderInput("t", document=ReportDocument("t")))
    with pytest.raises(KeyError):
        render("rtf", "document", RenderInput("t", document=ReportDocument("t")))


def test_a_new_format_is_one_registration(monkeypatch):
    monkeypatch.setattr(report_render, "_RENDERERS", dict(report_render._RENDERERS))
    report_render.register_renderer(Renderer("txt", ("table",), lambda s: RenderedFile(b"x", "csv")))

    assert formats_for("table") == ["csv", "xlsx", "ods", "txt"]
    with pytest.raises(ValueError):
        report_render.register_renderer(Renderer("txt", ("table",), lambda s: RenderedFile(b"", "csv")))


def test_unique_names_are_distinct_case_insensitively_and_cut_to_the_limit():
    assert unique_names(["A", "a", "A"], clean=str) == ["A", "a (2)", "A (3)"]
    assert unique_names(["", ""], clean=str) == ["Sheet", "Sheet (2)"]
    long = "x" * 40
    names = unique_names([long, long], clean=str, limit=31)
    assert all(len(n) <= 31 for n in names) and names[0] != names[1] and names[1].endswith(" (2)")


def rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode())))


def test_one_sheet_is_a_plain_csv_with_formula_neutralised_cells():
    out = render("csv", "table", RenderInput("t", sheets=[Sheet("S", ["Code", "=cmd"], [["R-1", "@SUM(A1)"], ["R-2", "ok"]])]))

    assert out.extension == "csv"
    assert rows(out.content) == [["Code", "'=cmd"], ["R-1", "'@SUM(A1)"], ["R-2", "ok"]]


def test_several_sheets_are_a_zip_with_one_csv_each_and_safe_unique_names():
    out = render("csv", "table", RenderInput("t", sheets=[Sheet("a/b", ["x"], [["1"]]), Sheet("A/B", ["x"], [["2"]]), Sheet("..\\evil", ["x"], [])]))

    assert out.extension == "zip"
    with zipfile.ZipFile(io.BytesIO(out.content)) as archive:
        names = archive.namelist()
        assert len({n.lower() for n in names}) == 3
        assert all(n.endswith(".csv") and "/" not in n and "\\" not in n for n in names)
        assert rows(archive.read(names[0])) == [["x"], ["1"]]


def test_xlsx_cells_are_text_guarded_and_cut_at_the_excel_limit():
    long = "y" * (MAX_CELL_CHARS + 50)
    out = render("xlsx", "table", RenderInput("t", sheets=[Sheet("S", ["h"], [["=1+1"], ["+cmd"], ["-5"], [long], ["plain"]])]))

    sheet = openpyxl.load_workbook(io.BytesIO(out.content))["S"]
    values = [sheet.cell(r, 1).value for r in range(1, 7)]
    assert out.extension == "xlsx" and values[0] == "h"
    assert values[1:4] == ["'=1+1", "'+cmd", "'-5"] and values[5] == "plain"
    assert len(values[4]) == MAX_CELL_CHARS
    assert all(sheet.cell(r, 1).data_type == "s" for r in range(1, 7))
    assert sheet.cell(1, 1).font.bold and sheet.freeze_panes == "A2"


def test_xlsx_sheet_names_are_cleaned_unique_and_at_most_31_characters():
    names = ["Plan: Q1/Q2 [draft]?", "plan: q1/q2 [draft]?", "n" * 50, ""]
    out = render("xlsx", "table", RenderInput("t", sheets=[Sheet(n, ["a"], [["1"]]) for n in names]))

    sheetnames = openpyxl.load_workbook(io.BytesIO(out.content)).sheetnames
    assert len(sheetnames) == 4 and len({n.lower() for n in sheetnames}) == 4
    assert all(len(n) <= 31 and not set("[]:*?/\\") & set(n) for n in sheetnames)


def test_xlsx_with_no_sheet_is_still_a_valid_workbook():
    out = render("xlsx", "table", RenderInput("t", sheets=[]))

    assert openpyxl.load_workbook(io.BytesIO(out.content)).sheetnames == ["Sheet"]


# ---- html ------------------------------------------------------------------------------------------------------


def page(document: ReportDocument) -> str:
    return render_html_document(document).decode()


def test_html_escapes_every_string_and_carries_no_script():
    hostile = '<img src=x onerror="alert(1)"> & "q"'
    out = page(ReportDocument(
        title=hostile,
        blocks=[Heading(hostile, 2), Paragraph(hostile), Note(hostile), BulletList([hostile]), Table([hostile], [[hostile]])],
        branding=ReportBranding(footer_text=hostile),
    ))

    assert "<img" not in out and "onerror=\"" not in out and "<script" not in out
    assert out.count("&lt;img src=x onerror=&quot;alert(1)&quot;&gt; &amp; &quot;q&quot;") >= 7
    assert f'content="{CSP}"' in out


def test_html_headings_tones_page_breaks_and_gaps():
    out = page(ReportDocument(
        title="T", blocks=[
            Heading("One", 1), Heading("Deep", 9), Gap(1), PageBreak(),
            Table(["s"], [["bad"], ["none"]], tones=[["danger"], [None]], compact=True),
        ],
    ))

    assert '<h2 id="one">One</h2>' in out and '<h5 id="deep">Deep</h5>' in out  # a heading carries its anchor
    assert out.count('class="page-break"') == 1 and 'class="compact"' in out
    assert f'style="background-color:#{TONE_FILL["danger"]}' in out
    assert out.count('style="background-color:#') == 1  # only the toned cell carries a style


def test_html_branding_accent_logo_footer_and_a_bad_colour():
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360000002000001e221bc330000000049454e44ae426082"
    )
    branded = page(ReportDocument(
        title="T", blocks=[Image(png, alt="")],
        branding=ReportBranding(accent_color_hex="#112233", footer_text="Confidential", logo_bytes=png),
    ))
    bad = page(ReportDocument(title="T", branding=ReportBranding(accent_color_hex="red;} body{display:none")))

    assert "--accent: #112233" in branded and "<footer>Confidential</footer>" in branded
    assert branded.count("data:image/png;base64,") == 2
    assert "display:none" not in bad and "--accent: #" in bad


def test_html_does_not_inline_bytes_that_are_not_an_image():
    out = page(ReportDocument(title="T", blocks=[Image(b"<script>alert(1)</script>")], branding=ReportBranding(logo_bytes=b"junk")))

    assert "base64" not in out and "<script" not in out and "<img" not in out


# --- typed cells ---------------------------------------------------------------------------------------------


def test_typed_cell_types_numbers_dates_and_instants_and_keeps_everything_else_text():
    assert typed_cell("3", 3) == DataCell("3", 3)
    assert typed_cell("-2.5", -2.5) == DataCell("-2.5", -2.5)
    assert typed_cell("4.20", Decimal("4.20")) == DataCell("4.20", Decimal("4.20"))
    assert typed_cell("2026-10-30", date(2026, 10, 30)) == DataCell("2026-10-30", date(2026, 10, 30))
    # An instant is stored as naive UTC: 09:00 at UTC+10 is 23:00 the day before.
    sydney = datetime(2026, 10, 30, 9, 0, tzinfo=timezone(timedelta(hours=10)))
    assert typed_cell("x", sydney) == DataCell("x", datetime(2026, 10, 29, 23, 0))
    la = datetime(2026, 10, 30, 23, 0, tzinfo=timezone(timedelta(hours=-8)))
    assert typed_cell("x", la) == DataCell("x", datetime(2026, 10, 31, 7, 0))
    # Not held exactly, not finite, empty, boolean or not a number at all: text.
    assert typed_cell("1000000000000000", 10**15) == "1000000000000000"
    assert typed_cell("nan", float("nan")) == "nan" and typed_cell("inf", float("inf")) == "inf"
    assert typed_cell("NaN", Decimal("NaN")) == "NaN"
    assert typed_cell("", 5) == "" and typed_cell("Yes", True) == "Yes" and typed_cell("=1+1", "=1+1") == "=1+1"


SAMPLE = Sheet("S", ["n", "d", "at", "t"], [
    [DataCell("-3", -3), DataCell("2026-10-30", date(2026, 10, 30)), DataCell("x", datetime(2026, 10, 30, 9, 15, 30)), "=1+1"],
    [DataCell("2.5", 2.5), "", "", "-7"],
])


def test_xlsx_stores_typed_cells_as_numbers_and_dates_and_text_stays_guarded():
    out = render("xlsx", "table", RenderInput("t", sheets=[SAMPLE]))

    sheet = openpyxl.load_workbook(io.BytesIO(out.content))["S"]
    assert sheet["A2"].value == -3 and sheet["A2"].data_type == "n"
    assert sheet["A3"].value == 2.5
    assert sheet["B2"].value == datetime(2026, 10, 30) and sheet["B2"].number_format == "yyyy-mm-dd"
    assert sheet["C2"].value == datetime(2026, 10, 30, 9, 15, 30) and sheet["C2"].number_format.startswith("yyyy-mm-dd hh:mm:ss")
    assert sheet["D2"].value == "'=1+1" and sheet["D2"].data_type == "s" and sheet["D3"].value == "'-7"


def test_ods_stores_typed_cells_with_value_types_and_text_stays_guarded():
    out = render("ods", "table", RenderInput("t", sheets=[SAMPLE]))

    with zipfile.ZipFile(io.BytesIO(out.content)) as archive:
        content = archive.read("content.xml").decode()
    assert 'office:value-type="float" office:value="-3"' in content and 'office:value="2.5"' in content
    assert 'office:value-type="date" office:date-value="2026-10-30"' in content
    assert 'office:date-value="2026-10-30T09:15:30"' in content
    assert "<text:p>&apos;=1+1</text:p>" in content or "<text:p>'=1+1</text:p>" in content
    assert 'style:data-style-name="nf_date"' in content and 'style:data-style-name="nf_instant"' in content
    from xml.dom import minidom
    minidom.parseString(content.encode())  # well formed


def test_csv_writes_a_typed_cell_as_its_text_so_a_negative_number_stays_a_number():
    out = render("csv", "table", RenderInput("t", sheets=[SAMPLE]))

    assert rows(out.content) == [["n", "d", "at", "t"], ["-3", "2026-10-30", "x", "'=1+1"], ["2.5", "", "", "'-7"]]


def test_xlsx_reads_back_in_utc_for_an_aware_instant():
    cell = typed_cell("t", datetime(2026, 10, 30, 9, 0, tzinfo=UTC))
    out = render("xlsx", "table", RenderInput("t", sheets=[Sheet("S", ["at"], [[cell]])]))

    assert openpyxl.load_workbook(io.BytesIO(out.content))["S"]["A2"].value == datetime(2026, 10, 30, 9, 0)


# --- links between sheets ---------------------------------------------------------------------------------


def linked_sheets() -> list[Sheet]:
    return [
        Sheet("Owner's list", ["Code"], [["R-1"], ["R-2"]]),
        Sheet("Pain/Points", ["Title", "Needs", "Count"], [
            ["Late", LinkedText("R-2", CellLink(0, 3)), DataCell("2", 2, CellLink(0, 2))],
            ["=evil", LinkedText("=cmd", CellLink(0, 2)), ""],
            ["Gone", LinkedText("R-9", CellLink(7, 2)), ""],  # a sheet that is not in the workbook
        ]),
    ]


def test_xlsx_links_open_the_row_on_the_tab_the_sheet_really_got():
    out = render("xlsx", "table", RenderInput("t", sheets=linked_sheets()))

    workbook = openpyxl.load_workbook(io.BytesIO(out.content))
    pain = workbook["Pain Points"]  # the slash was cleaned out of the tab name
    assert pain["B2"].hyperlink.location == "'Owner''s list'!A3" and pain["B2"].value == "R-2"
    assert pain["C2"].hyperlink.location == "'Owner''s list'!A2" and pain["C2"].value == 2
    assert pain["B3"].hyperlink.location == "'Owner''s list'!A2" and pain["B3"].value == "'=cmd" and pain["B3"].data_type == "s"
    assert pain["B4"].hyperlink is None and pain["B4"].value == "R-9"
    assert pain["B2"].font.underline == "single" and pain["A2"].hyperlink is None


def test_ods_links_open_the_row_on_the_tab_the_sheet_really_got():
    out = render("ods", "table", RenderInput("t", sheets=linked_sheets()))

    with zipfile.ZipFile(io.BytesIO(out.content)) as archive:
        content = archive.read("content.xml").decode()
    assert 'xlink:href="#&apos;Owner&apos;&apos;s list&apos;.A3"' in content or "xlink:href=\"#'Owner''s list'.A3\"" in content
    assert content.count("<text:a ") == 3  # the unknown sheet is plain text
    assert "table:name=\"Pain Points\"" in content
    from xml.dom import minidom
    minidom.parseString(content.encode())


def test_csv_drops_links_and_keeps_the_guarded_text():
    out = render("csv", "table", RenderInput("t", sheets=linked_sheets()))

    with zipfile.ZipFile(io.BytesIO(out.content)) as archive:
        pain = rows(archive.read(next(n for n in archive.namelist() if n.startswith("Pain"))))
    assert pain == [["Title", "Needs", "Count"], ["Late", "R-2", "2"], ["'=evil", "'=cmd", ""], ["Gone", "R-9", ""]]
