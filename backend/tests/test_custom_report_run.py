"""Tests for running a custom report layout (`services.custom_report_run`, `routers.custom_reports`): the
project and organisation-wide runs, the formats a layout family allows, who can run what, the provenance row,
the work limits and the catalogue entries. Each test builds its own organisation.
"""

from __future__ import annotations

import csv
import io
import uuid
from typing import Any

import pytest
from sqlalchemy import select

from app.database import SessionLocal
from app.models.generated_report import GeneratedReport
from app.services import custom_report_run, link_graph, report_layout_eval
from app.services.report_layout_tree import empty_tree
from tests.conftest import auth_headers, create_org_user, login
from tests.link_world import REQ, World


@pytest.fixture
def world(client, admin_token) -> World:
    return World(client, admin_token, "Run Org")


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


def document_tree(text: str = "{{record.title}}") -> dict[str, Any]:
    """A document layout: a heading, then one paragraph per requirement."""
    return {
        **empty_tree("document"),
        "nodes": [
            {"type": "heading", "text": "Requirements of {{project_name}}"},
            {
                "type": "each", "source": {"record_type": REQ, "sort": [{"field": "title"}]},
                "children": [{"type": "text", "text": text}],
            },
        ],
    }


def sheet_tree(*names: str) -> dict[str, Any]:
    """A table layout with one sheet per name, each listing requirement codes and titles."""
    return {
        **empty_tree("table"),
        "nodes": [
            {
                "type": "sheet", "name": name, "source": {"record_type": REQ, "sort": [{"field": "title"}]},
                "columns": [{"field": "code"}, {"field": "title"}],
            }
            for name in names
        ],
    }


def make_layout(world: World, tree: dict[str, Any], name: str, *, project: dict | None = None) -> dict:
    url = (
        f"/api/v1/projects/{project['id']}/report-layouts" if project else f"/api/v1/orgs/{world.org['id']}/report-layouts"
    )
    resp = world.client.post(
        url, json={"name": name, "description": "d", "family": tree["family"], "tree": tree}, headers=auth_headers(world.token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def project_run(world: World, project: dict, layout: dict, token: str | None = None, **params: Any):
    return world.client.get(
        f"/api/v1/projects/{project['id']}/custom-reports/{layout['id']}/run", params=params, headers=auth_headers(token or world.token),
    )


def org_run(world: World, layout: dict, token: str | None = None, **params: Any):
    return world.client.get(
        f"/api/v1/orgs/{world.org['id']}/custom-reports/{layout['id']}/run", params=params, headers=auth_headers(token or world.token),
    )


def rows_of(report: dict, section: int = 0) -> list[list[str]]:
    return report["sections"][section]["rows"]


def member_in(world: World, project: dict, role: str = "member") -> str:
    """A fresh organisation member holding `role` in `project` only; returns their token."""
    email = f"{role}-{uuid.uuid4().hex[:6]}@example.com"
    user_id = create_org_user(world.client, world.token, world.org["id"], email)
    resp = world.client.post(
        f"/api/v1/projects/{project['id']}/roles", json={"user_id": user_id, "role": role}, headers=auth_headers(world.token),
    )
    assert resp.status_code in (200, 201, 204), resp.text
    return login(world.client, email, "Password123!")


def provenance(db, layout: dict) -> list[GeneratedReport]:
    return list(db.scalars(
        select(GeneratedReport).where(GeneratedReport.layout_id == uuid.UUID(layout["id"])).order_by(GeneratedReport.created_at)
    ).all())


# ---- a project run ---------------------------------------------------------------------------------------------


def test_a_project_run_returns_the_report_as_json(world):
    project = world.project("Alpha")
    world.requirement(project, "Zeta rule")
    world.requirement(project, "Alpha rule")
    layout = make_layout(world, document_tree("{{record.code}} {{record.title}}"), "Rules", project=project)

    resp = project_run(world, project, layout)

    assert resp.status_code == 200, resp.text
    report = resp.json()
    assert report["title"] == "Rules" and report["scope_label"] == "Alpha" and report["eligible_projects"] == 1
    assert report["data"]["family"] == "document" and report["data"]["layout_version"] == 1
    assert report["metrics"] == [{"label": "Records read", "value": 2, "gap": False, "group": None, "link": None}]
    # No table in this layout, so its text (heading first) is carried as the report's notes, in order.
    assert report["notes"][0] == "Requirements of Alpha"
    assert report["notes"][1].endswith("Alpha rule") and report["notes"][2].endswith("Zeta rule")
    assert report["sections"] == []


def test_a_table_in_the_document_becomes_a_section_titled_by_its_heading(world):
    project = world.project()
    world.requirement(project, "One")
    tree = {
        **empty_tree("document"),
        "nodes": [
            {"type": "heading", "text": "All requirements"},
            {"type": "text", "text": "Everything we track."},
            {"type": "table", "source": {"record_type": REQ}, "columns": [{"field": "title"}]},
        ],
    }
    layout = make_layout(world, tree, "Listing", project=project)

    report = project_run(world, project, layout).json()

    assert [s["title"] for s in report["sections"]] == ["All requirements"]
    assert report["sections"][0]["note"] == "Everything we track."
    assert rows_of(report) == [["One"]]


def test_pdf_and_word_are_produced_for_a_document_layout(world):
    project = world.project()
    world.requirement(project, "One")
    layout = make_layout(world, document_tree(), "Doc", project=project)

    pdf = project_run(world, project, layout, format="pdf")
    docx = project_run(world, project, layout, format="docx")

    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == "application/pdf"
    assert docx.status_code == 200 and docx.content.startswith(b"PK")
    assert "attachment" in pdf.headers["content-disposition"]


def test_a_single_sheet_table_layout_downloads_as_csv(world):
    project = world.project()
    world.requirement(project, "=1+1 formula")
    world.requirement(project, "Plain")
    layout = make_layout(world, sheet_tree("Reqs"), "Sheet", project=project)

    resp = project_run(world, project, layout, format="csv")

    assert resp.status_code == 200 and resp.headers["content-type"].startswith("text/csv")
    parsed = list(csv.reader(io.StringIO(resp.text)))
    assert parsed[0] == ["Code", "Title"]
    assert sorted(row[1] for row in parsed[1:]) == ["'=1+1 formula", "Plain"]


def test_a_layout_is_only_produced_in_formats_its_family_allows(world):
    project = world.project()
    document = make_layout(world, document_tree(), "Doc", project=project)
    sheet = make_layout(world, sheet_tree("A"), "Sheet", project=project)

    for fmt in ("csv", "xlsx", "ods"):
        refused = project_run(world, project, document, format=fmt)
        assert refused.status_code == 400 and "json, pdf, docx, html, odt" in refused.json()["detail"]
    for fmt in ("pdf", "docx", "html", "odt"):
        refused = project_run(world, project, sheet, format=fmt)
        assert refused.status_code == 400 and "json, csv, xlsx, ods" in refused.json()["detail"]
    assert project_run(world, project, document, format="rtf").status_code == 422  # not a known format at all
    assert custom_report_run.available_formats("table") == ["json", "csv", "xlsx", "ods"]
    assert custom_report_run.available_formats("document") == ["json", "pdf", "docx", "html", "odt"]


def test_html_is_a_standalone_script_free_page_that_escapes_record_text(world):
    project = world.project()
    world.requirement(project, "<script>alert(1)</script> & co")
    layout = make_layout(world, document_tree("{{record.title}}"), "Doc", project=project)

    resp = project_run(world, project, layout, format="html")

    assert resp.status_code == 200 and resp.headers["content-type"].startswith("text/html")
    assert "attachment" in resp.headers["content-disposition"] and resp.headers["content-disposition"].endswith('.html"')
    page = resp.text
    assert "<script" not in page and "&lt;script&gt;alert(1)&lt;/script&gt; &amp; co" in page
    assert "Content-Security-Policy" in page and "default-src 'none'" in page


def test_a_table_layout_downloads_as_xlsx_with_text_cells_and_one_tab_per_sheet(world):
    import openpyxl

    project = world.project()
    world.requirement(project, "=1+1 formula")
    layout = make_layout(world, sheet_tree("Reqs", "Other"), "Two", project=project)

    resp = project_run(world, project, layout, format="xlsx")

    assert resp.status_code == 200 and resp.content.startswith(b"PK")
    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Reqs", "Other"]
    rows = list(workbook["Reqs"].iter_rows(values_only=True))
    assert rows[0] == ("Code", "Title") and rows[1][1] == "'=1+1 formula"
    assert workbook["Reqs"].cell(2, 2).data_type == "s"


def test_a_document_layout_downloads_as_odt_with_escaped_record_text(world):
    import zipfile

    project = world.project()
    world.requirement(project, "<script>alert(1)</script> & co")
    layout = make_layout(world, document_tree("{{record.title}}"), "Doc", project=project)

    resp = project_run(world, project, layout, format="odt")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.oasis.opendocument.text"
    assert resp.headers["content-disposition"].endswith('.odt"')
    with zipfile.ZipFile(io.BytesIO(resp.content)) as archive:
        assert archive.read("mimetype") == b"application/vnd.oasis.opendocument.text"
        content = archive.read("content.xml").decode()
    assert "<script" not in content and "&lt;script&gt;alert(1)&lt;/script&gt; &amp; co" in content


def test_a_table_layout_downloads_as_ods_with_text_cells_and_one_tab_per_sheet(world):
    import zipfile

    from lxml import etree

    project = world.project()
    world.requirement(project, "=1+1 formula")
    layout = make_layout(world, sheet_tree("Reqs", "Other"), "Two", project=project)

    resp = project_run(world, project, layout, format="ods")

    assert resp.status_code == 200 and resp.headers["content-type"] == "application/vnd.oasis.opendocument.spreadsheet"
    assert resp.headers["content-disposition"].endswith('.ods"')
    with zipfile.ZipFile(io.BytesIO(resp.content)) as archive:
        content = etree.fromstring(archive.read("content.xml"))
    ns = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0"}
    assert [t.get(f"{{{ns['table']}}}name") for t in content.xpath("//table:table", namespaces=ns)] == ["Reqs", "Other"]
    assert "'=1+1 formula" in "".join(content.itertext()) and not content.xpath("//@table:formula", namespaces=ns)


def test_a_table_layout_with_several_sheets_downloads_as_a_zip_of_csv_files(world):
    import zipfile

    project = world.project()
    world.requirement(project, "One")
    layout = make_layout(world, sheet_tree("First", "Second"), "Two", project=project)

    resp = project_run(world, project, layout, format="csv")

    assert resp.status_code == 200 and resp.headers["content-type"] == "application/zip"
    assert resp.headers["content-disposition"].endswith('.zip"')
    with zipfile.ZipFile(io.BytesIO(resp.content)) as archive:
        assert sorted(archive.namelist()) == ["First.csv", "Second.csv"]
        assert archive.read("First.csv").decode().splitlines()[0] == "Code,Title"


def test_an_invalid_branding_template_is_refused(world):
    project = world.project()
    layout = make_layout(world, document_tree(), "Doc", project=project)

    resp = project_run(world, project, layout, format="pdf", report_template_id=str(uuid.uuid4()))

    assert resp.status_code == 400


def test_include_children_reads_the_readable_child_projects(world):
    parent = world.project("Parent")
    child = world.project("Child", parent=parent)
    world.requirement(parent, "In parent")
    world.requirement(child, "In child")
    layout = make_layout(world, document_tree(), "Doc", project=parent)

    alone = project_run(world, parent, layout).json()
    with_children = project_run(world, parent, layout, include_children=True).json()

    assert alone["metrics"][0]["value"] == 1 and with_children["metrics"][0]["value"] == 2
    assert with_children["eligible_projects"] == 2


# ---- which layouts a project can run ------------------------------------------------------------------------------


def test_a_project_runs_its_own_ancestors_and_the_organisations_layouts(world):
    parent = world.project("Parent")
    child = world.project("Child", parent=parent)
    world.requirement(child, "Child rule")
    org_layout = make_layout(world, document_tree(), "Org layout")
    parent_layout = make_layout(world, document_tree(), "Parent layout", project=parent)
    own = make_layout(world, document_tree(), "Own layout", project=child)

    for layout in (org_layout, parent_layout, own):
        assert project_run(world, child, layout).status_code == 200, layout["name"]
    assert project_run(world, parent, own).status_code == 404  # a descendant's layout is not inherited upwards


def test_another_projects_or_organisations_layout_is_not_found(world, client, admin_token):
    mine, other = world.project("Mine"), world.project("Other")
    theirs = make_layout(world, document_tree(), "Theirs", project=other)
    foreign_world = World(client, admin_token, "Foreign Org")
    foreign_project = foreign_world.project()
    foreign = make_layout(foreign_world, document_tree(), "Foreign")

    assert project_run(world, mine, theirs).status_code == 404
    assert project_run(world, mine, foreign).status_code == 404
    assert project_run(foreign_world, foreign_project, foreign).status_code == 200
    assert org_run(world, foreign).status_code == 404


def test_a_layout_a_project_hid_cannot_be_run_there_and_is_not_catalogued(world):
    project = world.project()
    layout = make_layout(world, document_tree(), "Org layout")
    hide = world.client.put(
        f"/api/v1/projects/{project['id']}/report-layouts/{layout['id']}/visibility", json={"hidden": True},
        headers=auth_headers(world.token),
    )
    assert hide.status_code == 200, hide.text

    assert project_run(world, project, layout).status_code == 404
    catalogue = world.client.get(f"/api/v1/projects/{project['id']}/report-catalogue", headers=auth_headers(world.token)).json()
    assert not [e for e in catalogue if e["module_key"] == "custom_reports"]


def test_a_non_member_cannot_run_a_project_layout(world):
    project = world.project()
    layout = make_layout(world, document_tree(), "Doc", project=project)
    outsider = login(world.client, *_outsider(world))

    assert project_run(world, project, layout, token=outsider).status_code in (403, 404)


def _outsider(world: World) -> tuple[str, str]:
    email = f"outsider-{uuid.uuid4().hex[:6]}@example.com"
    create_org_user(world.client, world.token, world.org["id"], email)
    return email, "Password123!"


# ---- what a run shows -------------------------------------------------------------------------------------------


def test_a_record_the_caller_cannot_view_is_counted_and_never_named(world, monkeypatch):
    project = world.project()
    world.requirement(project, "Secret rule")
    layout = make_layout(world, document_tree(), "Doc", project=project)
    real = link_graph.get_effective_permissions
    monkeypatch.setattr(
        link_graph, "get_effective_permissions",
        lambda *a, **k: {p for p in real(*a, **k) if not p.startswith("requirement:view")},
    )

    report = project_run(world, project, layout).json()

    assert "Secret rule" not in str(report)
    assert report["metrics"][0]["value"] == 0


# ---- the organisation-wide run -----------------------------------------------------------------------------------


def test_an_organisation_run_covers_every_project_the_caller_can_read(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    world.requirement(alpha, "Alpha rule")
    world.requirement(beta, "Beta rule")
    layout = make_layout(world, sheet_tree("Reqs"), "Everywhere")

    report = org_run(world, layout).json()

    assert report["scope_label"] == world.org["name"] and report["eligible_projects"] == 2
    assert report["sections"][0]["columns"] == ["Project", "Code", "Title"]
    assert sorted((r[0], r[2]) for r in rows_of(report)) == [("Alpha", "Alpha rule"), ("Beta", "Beta rule")]


def test_the_project_selection_narrows_the_run(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    world.requirement(alpha, "Alpha rule")
    world.requirement(beta, "Beta rule")
    layout = make_layout(world, sheet_tree("Reqs"), "Everywhere")

    report = org_run(world, layout, project_ids=beta["id"]).json()

    assert [r[0] for r in rows_of(report)] == ["Beta"] and report["eligible_projects"] == 1
    both = org_run(world, layout, project_ids=f"{alpha['id']}, {beta['id']}").json()
    assert both["eligible_projects"] == 2
    assert org_run(world, layout, project_ids="").json()["eligible_projects"] == 2


def test_a_selection_can_only_narrow_what_the_caller_can_read(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    world.requirement(alpha, "Alpha rule")
    world.requirement(beta, "Beta rule")
    layout = make_layout(world, sheet_tree("Reqs"), "Everywhere")
    only_alpha = member_in(world, alpha)

    report = org_run(world, layout, token=only_alpha).json()

    assert [r[0] for r in rows_of(report)] == ["Alpha"] and "Beta rule" not in str(report)
    assert org_run(world, layout, token=only_alpha, project_ids=beta["id"]).status_code == 404
    assert org_run(world, layout, token=only_alpha, project_ids=f"{alpha['id']},{uuid.uuid4()}").status_code == 404
    assert org_run(world, layout, project_ids="not-an-id").status_code == 422


def test_a_member_with_no_project_role_has_nothing_to_run(world):
    world.project("Alpha")
    layout = make_layout(world, sheet_tree("Reqs"), "Everywhere")
    email, password = _outsider(world)

    resp = org_run(world, layout, token=login(world.client, email, password))

    assert resp.status_code == 404


def test_a_project_that_hid_the_layout_is_left_out_of_the_organisation_run(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    world.requirement(alpha, "Alpha rule")
    world.requirement(beta, "Beta rule")
    layout = make_layout(world, sheet_tree("Reqs"), "Everywhere")
    world.client.put(
        f"/api/v1/projects/{beta['id']}/report-layouts/{layout['id']}/visibility", json={"hidden": True},
        headers=auth_headers(world.token),
    )

    report = org_run(world, layout).json()

    assert [r[0] for r in rows_of(report)] == ["Alpha"]
    assert any("left out" in n for n in report["notes"])
    assert org_run(world, layout, project_ids=beta["id"]).status_code == 404  # nothing left to run


def test_only_an_organisation_layout_runs_organisation_wide(world):
    project = world.project()
    own = make_layout(world, sheet_tree("Reqs"), "Project owned", project=project)

    assert org_run(world, own).status_code == 404


def test_an_organisation_document_has_a_heading_and_a_page_break_per_project(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    world.requirement(alpha, "Alpha rule")
    world.requirement(beta, "Beta rule")
    layout = make_layout(world, document_tree(), "Everywhere")

    run = _run_org(world, layout)

    from app.services.report_model import Heading, PageBreak

    blocks = run.evaluated.document.blocks
    top = [b.text for b in blocks if isinstance(b, Heading) and b.level == 1]
    assert top == ["Alpha", "Beta"]
    assert [b.text for b in blocks if isinstance(b, Heading) and b.level == 2] == ["Requirements of Alpha", "Requirements of Beta"]
    assert sum(isinstance(b, PageBreak) for b in blocks) == 1
    assert org_run(world, layout, format="pdf").content.startswith(b"%PDF")


def test_the_on_screen_sections_of_an_organisation_run_are_titled_by_project_then_heading(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    world.requirement(alpha, "Alpha rule")
    world.requirement(beta, "Beta rule")
    tree = {
        **empty_tree("document"),
        "nodes": [
            {"type": "heading", "text": "Requirements"},
            {"type": "table", "source": {"record_type": REQ}, "columns": [{"field": "title"}]},
        ],
    }
    layout = make_layout(world, tree, "Everywhere")

    report = org_run(world, layout).json()

    assert [s["title"] for s in report["sections"]] == ["Alpha / Requirements", "Beta / Requirements"]
    assert [s["rows"] for s in report["sections"]] == [[["Alpha rule"]], [["Beta rule"]]]


def _run_org(world: World, layout: dict):
    session = SessionLocal()
    try:
        from app.models.organization import Organization
        from app.models.report_layout import ReportLayout
        from app.models.user import User
        from app.services import report_layouts

        row = session.get(ReportLayout, uuid.UUID(layout["id"]))
        version = report_layouts.get_version(session, row.id, row.current_version)
        return custom_report_run.run_organization_layout(
            session, session.get(User, world.user_id), session.get(Organization, uuid.UUID(world.org["id"])), row, version,
        )
    finally:
        session.close()


def test_the_organisation_run_stops_at_the_project_limit(world, monkeypatch):
    for name in ("A", "B", "C"):
        world.project(name)
    layout = make_layout(world, sheet_tree("Reqs"), "Everywhere")
    monkeypatch.setattr(custom_report_run, "MAX_FANOUT_PROJECTS", 2)

    assert org_run(world, layout).status_code == 422
    assert org_run(world, layout, project_ids=",".join(
        p["id"] for p in world.client.get(f"/api/v1/orgs/{world.org['id']}/projects", headers=auth_headers(world.token)).json()[:2]
    )).status_code == 200


# ---- limits ----------------------------------------------------------------------------------------------------


def test_a_run_over_the_output_limit_is_refused_and_leaves_no_provenance(world, monkeypatch, db):
    project = world.project()
    for i in range(5):
        world.requirement(project, f"Rule {i}")
    layout = make_layout(world, document_tree(), "Doc", project=project)
    monkeypatch.setattr(report_layout_eval, "MAX_OUTPUT", 3)

    resp = project_run(world, project, layout)

    assert resp.status_code == 413
    assert provenance(db, layout) == []


def test_an_organisation_run_shares_one_output_budget_across_projects(world, monkeypatch, db):
    for name in ("A", "B", "C"):
        world.requirement(world.project(name), f"Rule {name}")
    layout = make_layout(world, document_tree(), "Everywhere")
    # Each project alone stays under the cap (heading + 1 text = 2 blocks); three together do not.
    monkeypatch.setattr(report_layout_eval, "MAX_OUTPUT", 5)
    monkeypatch.setattr(custom_report_run, "MAX_OUTPUT", 5)

    assert org_run(world, layout).status_code == 413
    assert provenance(db, layout) == []


def test_a_run_past_its_deadline_answers_408(world, monkeypatch):
    project = world.project()
    world.requirement(project, "One")
    layout = make_layout(world, document_tree(), "Doc", project=project)
    monkeypatch.setattr(report_layout_eval, "DEFAULT_TIMEOUT_SECONDS", -1.0)
    monkeypatch.setattr(custom_report_run, "DEFAULT_TIMEOUT_SECONDS", -1.0)

    assert project_run(world, project, layout).status_code == 408


# ---- provenance ------------------------------------------------------------------------------------------------


def test_every_run_records_its_provenance_without_content(world, db):
    project = world.project()
    world.requirement(project, "Confidential widget spec")
    layout = make_layout(world, document_tree(), "Doc", project=project)

    assert project_run(world, project, layout, format="pdf").status_code == 200
    assert project_run(world, project, layout, include_children=True).status_code == 200

    rows = provenance(db, layout)
    assert [r.format for r in rows] == ["pdf", "json"]
    first = rows[0]
    assert first.organization_id == uuid.UUID(world.org["id"]) and first.project_id == uuid.UUID(project["id"])
    assert first.scope == "project" and first.requested_by == world.user_id
    assert first.layout_name == "Doc" and first.layout_version == 1 and len(first.tree_hash) == 64
    assert first.engine_version == custom_report_run.ENGINE_VERSION and first.created_at.tzinfo is not None
    assert first.scope_project_ids == [project["id"]] and first.widgets == []
    assert first.record_counts == {"requirement": 1} and first.stats["rows"] == 0 and first.stats["blocks"] >= 2
    assert first.params == {
        "report_template_id": None, "include_children": False, "requested_project_ids": None, "variables": {},
    }
    assert rows[1].params["include_children"] is True
    for row in rows:
        assert "Confidential widget spec" not in str(
            (row.params, row.stats, row.record_counts, row.layout_name, row.scope_project_ids, row.widgets)
        )


def test_the_provenance_pins_the_version_that_ran(world, db):
    project = world.project()
    world.requirement(project, "One")
    layout = make_layout(world, document_tree(), "Doc", project=project)
    project_run(world, project, layout)
    saved = world.client.put(
        f"/api/v1/projects/{project['id']}/report-layouts/{layout['id']}/tree",
        json={"tree": document_tree("Changed {{record.title}}"), "base_version": 1, "comment": "v2"},
        headers=auth_headers(world.token),
    )
    assert saved.status_code == 200, saved.text
    project_run(world, project, layout)

    first, second = provenance(db, layout)

    assert (first.layout_version, second.layout_version) == (1, 2) and first.tree_hash != second.tree_hash


def test_the_source_digest_changes_only_when_the_data_does(world, db):
    project = world.project()
    req_id = world.requirement(project, "Original")
    layout = make_layout(world, document_tree(), "Doc", project=project)
    project_run(world, project, layout)
    project_run(world, project, layout)
    current = world.client.get(f"/api/v1/projects/{project['id']}/requirements/{req_id}", headers=auth_headers(world.token)).json()
    edit = world.client.put(
        f"/api/v1/projects/{project['id']}/requirements/{req_id}",
        json={
            "name": "Edited", "component_id": project["component_id"], "category_id": project["category_id"],
            "owner": current["owner"],
        },
        headers=auth_headers(world.token),
    )
    assert edit.status_code == 200, edit.text
    project_run(world, project, layout)

    a, b, c = provenance(db, layout)

    assert len(a.source_digest) == 64 and a.source_digest == b.source_digest and b.source_digest != c.source_digest


def test_the_provenance_names_the_version_of_every_versioned_record(world, db):
    project = world.project()
    req_id = world.requirement(project, "Original")
    world.action(project, "An action")  # a type with no versions is counted but not listed
    tree = {
        **empty_tree("document"),
        "nodes": [
            {"type": "each", "source": {"record_type": REQ}, "children": [{"type": "text", "text": "{{record.title}}"}]},
            {"type": "each", "source": {"record_type": "requirement_action"}, "children": [{"type": "text", "text": "{{record.title}}"}]},
        ],
    }
    layout = make_layout(world, tree, "Doc", project=project)
    project_run(world, project, layout)
    current = world.client.get(f"/api/v1/projects/{project['id']}/requirements/{req_id}", headers=auth_headers(world.token)).json()
    edit = world.client.put(
        f"/api/v1/projects/{project['id']}/requirements/{req_id}",
        json={"name": "Edited", "component_id": project["component_id"], "category_id": project["category_id"], "owner": current["owner"]},
        headers=auth_headers(world.token),
    )
    assert edit.status_code == 200, edit.text
    project_run(world, project, layout)

    first, second = provenance(db, layout)

    assert first.record_versions == {"requirement": {str(req_id): 1}} and first.record_versions_truncated is False
    assert second.record_versions == {"requirement": {str(req_id): 2}}
    assert second.record_counts == {"requirement": 1, "requirement_action": 1}


def test_the_recorded_versions_are_bounded_and_say_when_they_were_cut(world, db, monkeypatch):
    project = world.project()
    for i in range(3):
        world.requirement(project, f"Rule {i}")
    layout = make_layout(world, document_tree(), "Doc", project=project)
    monkeypatch.setattr(report_layout_eval, "MAX_RECORDED_VERSIONS", 2)

    project_run(world, project, layout)

    (row,) = provenance(db, layout)
    assert sum(len(v) for v in row.record_versions.values()) == 2 and row.record_versions_truncated is True
    assert row.record_counts == {"requirement": 3} and len(row.source_digest) == 64  # the digest still covers all three


def test_an_organisation_run_lists_each_versioned_record_once(world, db):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    a, b = world.requirement(alpha, "A"), world.requirement(beta, "B")
    layout = make_layout(world, document_tree(), "Everywhere")
    assert org_run(world, layout).status_code == 200

    (row,) = provenance(db, layout)

    assert row.record_versions == {"requirement": {str(a): 1, str(b): 1}} and row.record_versions_truncated is False


def test_every_versioned_report_type_exposes_its_current_version():
    from app.modules.registry import get_report_field_providers

    versioned = {t: p for t, p in get_report_field_providers().items() if p.where}  # `valid_to IS NULL` marks a versioned type
    assert "requirement" in versioned and "risk" in versioned
    assert [t for t, p in versioned.items() if p.columns.version is None] == []


def test_an_organisation_run_records_the_selection_and_the_projects_covered(world, db):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    layout = make_layout(world, sheet_tree("Reqs"), "Everywhere")

    assert org_run(world, layout, project_ids=beta["id"], format="csv").status_code == 200
    assert org_run(world, layout).status_code == 200

    narrowed, everyone = provenance(db, layout)
    assert narrowed.scope == "organization" and narrowed.project_id is None and narrowed.format == "csv"
    assert narrowed.scope_project_ids == [beta["id"]] and narrowed.params["requested_project_ids"] == [beta["id"]]
    assert sorted(everyone.scope_project_ids) == sorted([alpha["id"], beta["id"]])
    assert everyone.params["requested_project_ids"] is None


def test_provenance_outlives_the_layout_it_describes(world, db):
    project = world.project()
    layout = make_layout(world, document_tree(), "Doomed", project=project)
    project_run(world, project, layout)
    gone = world.client.delete(
        f"/api/v1/projects/{project['id']}/report-layouts/{layout['id']}", headers=auth_headers(world.token),
    )
    assert gone.status_code == 204

    (row,) = provenance(db, layout)

    assert row.layout_name == "Doomed" and row.layout_version == 1


# ---- the catalogue ---------------------------------------------------------------------------------------------


def test_the_project_catalogue_lists_runnable_layouts_with_their_formats(world):
    project = world.project()
    document = make_layout(world, document_tree(), "Doc", project=project)
    sheet = make_layout(world, sheet_tree("A"), "Sheet")

    entries = world.client.get(f"/api/v1/projects/{project['id']}/report-catalogue", headers=auth_headers(world.token)).json()
    custom = {e["title"]: e for e in entries if e["module_key"] == "custom_reports"}

    assert set(custom) == {"Doc", "Sheet"}
    assert custom["Doc"]["formats"] == ["json", "pdf", "docx", "html", "odt"] and custom["Doc"]["scope"] == "project"
    assert custom["Doc"]["path"] == f"/api/v1/projects/{project['id']}/custom-reports/{document['id']}/run"
    assert custom["Sheet"]["formats"] == ["json", "csv", "xlsx", "ods"] and custom["Doc"]["supports_include_children"] is True
    assert custom["Doc"]["module_name"] == "Custom reports" and custom["Doc"]["params"] == []
    assert len({e["slug"] for e in entries}) == len(entries)
    assert custom["Sheet"]["slug"].startswith("custom-sheet-") and custom["Sheet"]["key"] == sheet["id"]


def test_the_organisation_catalogue_offers_a_project_selection_of_what_the_caller_can_read(world):
    alpha, beta = world.project("Alpha"), world.project("Beta")
    make_layout(world, sheet_tree("A"), "Everywhere")
    make_layout(world, document_tree(), "Project owned", project=alpha)
    only_alpha = member_in(world, alpha)
    url = f"/api/v1/orgs/{world.org['id']}/report-catalogue"

    admin_view = [e for e in world.client.get(url, headers=auth_headers(world.token)).json() if e["module_key"] == "custom_reports"]
    member_view = [e for e in world.client.get(url, headers=auth_headers(only_alpha)).json() if e["module_key"] == "custom_reports"]

    assert [e["title"] for e in admin_view] == ["Everywhere"]  # a project's own layout is not organisation-wide
    entry = admin_view[0]
    assert entry["scope"] == "organization" and entry["supports_project_multi_filter"] is True
    assert entry["supports_project_filter"] is False and entry["supports_include_children"] is False
    assert sorted(p["name"] for p in entry["projects"]) == ["Alpha", "Beta"]
    assert [p["name"] for p in member_view[0]["projects"]] == ["Alpha"]
    assert beta["id"] not in str(member_view)


# ---- report content is gone -------------------------------------------------------------------------------------


def test_the_report_content_endpoints_no_longer_exist(world):
    project = world.project("Alpha")
    h = auth_headers(world.token)
    urls = [
        f"/api/v1/projects/{project['id']}/report-config",
        f"/api/v1/projects/{project['id']}/report-types",
        f"/api/v1/projects/{project['id']}/report-content/effective?report_key=requirements",
        f"/api/v1/orgs/{world.org['id']}/report-defaults",
        f"/api/v1/orgs/{world.org['id']}/report-types",
    ]

    assert [world.client.get(url, headers=h).status_code for url in urls] == [404] * len(urls)


# ---- bundles ----------------------------------------------------------------------------------------------------


def export_project(world: World, project: dict) -> bytes:
    resp = world.client.get(f"/api/v1/projects/{project['id']}/export", headers=auth_headers(world.token))
    assert resp.status_code == 200, resp.text
    return resp.content


def import_project(world: World, bundle: bytes, name: str) -> dict:
    resp = world.client.post(
        "/api/v1/projects/import", data={"organization_id": world.org["id"], "name": name},
        files={"file": ("bundle.zip", bundle, "application/zip")}, headers=auth_headers(world.token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def with_json_edit(bundle: bytes, member: str, edit) -> bytes:
    """A copy of `bundle` whose `member` JSON document has been changed by `edit(parsed)`."""
    import json
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(bundle)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == member:
                parsed = json.loads(data)
                edit(parsed)
                data = json.dumps(parsed).encode()
            dst.writestr(item, data)
    return out.getvalue()


def test_a_project_bundle_carries_layouts_and_no_report_content(world):
    import json
    import zipfile

    source = world.project("Source")
    make_layout(world, document_tree(), "Own report", project=source)

    bundle = export_project(world, source)
    with zipfile.ZipFile(io.BytesIO(bundle)) as z:
        project_json = json.loads(z.read("project.json"))

    assert not {"report_content", "report_images", "default_report_template_name"} & project_json.keys()
    copy = import_project(world, bundle, "Copy")
    assert copy["warnings"] == []
    layouts_url = f"/api/v1/projects/{copy['project']['id']}/report-layouts"
    names = [r["name"] for r in world.client.get(layouts_url, headers=auth_headers(world.token)).json()]
    assert "Own report" in names


def test_an_older_project_bundle_s_report_content_is_not_imported_and_warns_once(world):
    source = world.project("Source")
    bundle = with_json_edit(
        export_project(world, source), "project.json",
        lambda parsed: parsed.update(report_content=[{"report_key": "requirements", "intro": "Old intro"}]),
    )

    result = import_project(world, bundle, "Copy")

    assert [w for w in result["warnings"] if "report content has been removed" in w] != []
    assert len([w for w in result["warnings"] if "report content has been removed" in w]) == 1


def test_an_older_organisation_bundle_s_report_content_is_not_imported_and_warns_once(client, admin_token, world):
    h = auth_headers(world.token)
    client.post(f"/api/v1/orgs/{world.org['id']}/report-templates", json={"name": "Brand"}, headers=h)
    bundle = client.get(f"/api/v1/orgs/{world.org['id']}/export", headers=h).content

    def add_legacy(parsed: dict) -> None:
        parsed["report_content"] = [{"report_key": "requirements", "intro": "Old org intro"}]
        parsed["report_templates"][0]["content"] = [{"report_key": "requirements", "intro": "Old template intro"}]

    resp = client.post(
        "/api/v1/orgs/import", data={"name": f"Copy {uuid.uuid4().hex[:6]}"},
        files={"file": ("bundle.zip", with_json_edit(bundle, "org.json", add_legacy), "application/zip")},
        headers=auth_headers(admin_token),
    )

    assert resp.status_code == 201, resp.text
    assert len([w for w in resp.json()["warnings"] if "report content has been removed" in w]) == 1
    new_org = resp.json()["organization"]["id"]
    [template] = client.get(f"/api/v1/orgs/{new_org}/report-templates", headers=h).json()
    assert template["name"] == "Brand" and "content" not in template
