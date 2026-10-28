"""Tests for a document layout's page size (`page_size` in the tree, the `@layout document landscape A3` text form),
its effect on the PDF, Word, OpenDocument text and HTML renderers, and the project placeholders a layout can use
(`{{project_code}}`, `{{project_description}}`).
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from docx import Document as DocxDocument

from app.services.report_docx import render_docx_document
from app.services.report_html import render_html_document
from app.services.report_layout_text import format_tree, parse_text
from app.services.report_layout_tree import InvalidLayoutError, assert_valid, empty_tree
from app.services.report_model import PAGE_SIZES_CM, Paragraph, ReportDocument, page_dimensions_cm
from app.services.report_odt import render_odt_document
from app.services.report_pdf import render_pdf_document
from tests.conftest import auth_headers
from tests.link_world import World


def document(size: str = "A4", landscape: bool = False) -> ReportDocument:
    return ReportDocument(title="T", blocks=[Paragraph("Body")], landscape=landscape, page_size=size)


# ---- the tree and the text view --------------------------------------------------------------------------------


def test_a_page_size_is_part_of_a_valid_document_tree():
    tree = empty_tree("document")
    for size in ("A3", "A4", "A5", "Letter", "Legal"):
        tree["page_size"] = size
        assert_valid(tree)


@pytest.mark.parametrize("value", ["B5", "a4", 4, None])
def test_an_unknown_page_size_is_refused_with_its_path(value):
    tree = empty_tree("document")
    tree["page_size"] = value
    with pytest.raises(InvalidLayoutError) as caught:
        assert_valid(tree)
    assert [e.path for e in caught.value.errors] == ["page_size"]


def test_a_table_layout_has_no_page_size():
    tree = empty_tree("table")
    tree["page_size"] = "A3"
    with pytest.raises(InvalidLayoutError) as caught:
        assert_valid(tree)
    assert caught.value.errors[0].path == "page_size"


@pytest.mark.parametrize("header", [
    "@layout document landscape A3", "@layout document A3 landscape", "@layout document Letter", "@layout document landscape",
])
def test_the_header_accepts_orientation_and_size_in_either_order_and_round_trips(header):
    parsed = parse_text(f"{header}\n")
    assert parsed.ok, parsed.errors
    again = parse_text(format_tree(parsed.tree))
    assert again.tree == parsed.tree


def test_the_header_writes_orientation_then_size():
    parsed = parse_text("@layout document A3 landscape\n")
    assert format_tree(parsed.tree).splitlines()[0] == "@layout document landscape A3"


@pytest.mark.parametrize("header", ["@layout document A3 A4", "@layout document portrait landscape", "@layout document A9"])
def test_a_repeated_or_unknown_header_word_is_a_positioned_error(header):
    parsed = parse_text(f"{header}\n")
    assert not parsed.ok and parsed.errors[0].line == 1


# ---- the renderers ---------------------------------------------------------------------------------------------


def test_the_dimensions_turn_on_their_side_for_landscape_and_fall_back_to_a4():
    assert page_dimensions_cm(document("Letter")) == (21.59, 27.94)
    assert page_dimensions_cm(document("Letter", landscape=True)) == (27.94, 21.59)
    assert page_dimensions_cm(document("nonsense")) == PAGE_SIZES_CM["A4"]


@pytest.mark.parametrize(("size", "landscape", "width_pt", "height_pt"), [
    ("A4", False, 595, 842), ("Letter", False, 612, 792), ("A3", True, 1191, 842), ("Legal", False, 612, 1008),
])
def test_the_pdf_page_has_the_chosen_size(size, landscape, width_pt, height_pt):
    pdf = render_pdf_document(document(size, landscape))
    box = re.search(rb"/MediaBox \[ ?0 0 ([\d.]+) ([\d.]+) ?\]", pdf)
    assert box is not None
    assert (round(float(box.group(1))), round(float(box.group(2)))) == (width_pt, height_pt)


def test_the_word_section_has_the_chosen_size():
    section = DocxDocument(io.BytesIO(render_docx_document(document("A3")))).sections[0]
    assert round(section.page_width.cm, 1) == 29.7 and round(section.page_height.cm, 1) == 42.0
    landscape = DocxDocument(io.BytesIO(render_docx_document(document("Letter", landscape=True)))).sections[0]
    assert round(landscape.page_width.cm, 2) == 27.94 and round(landscape.page_height.cm, 2) == 21.59


def test_the_opendocument_page_layout_has_the_chosen_size():
    with zipfile.ZipFile(io.BytesIO(render_odt_document(document("Legal")))) as package:
        styles = package.read("styles.xml").decode()
    assert 'fo:page-width="21.590cm"' in styles and 'fo:page-height="35.560cm"' in styles


def test_the_html_view_is_as_wide_as_the_page_is_for_a4_and_scales_for_other_sizes():
    def html(size: str, landscape: bool = False) -> str:
        return render_html_document(document(size, landscape)).decode()

    assert "max-width: 50rem" in html("A4") and "max-width: 70rem" in html("A4", landscape=True)
    assert "max-width: 99rem" in html("A3", landscape=True)


# ---- project placeholders --------------------------------------------------------------------------------------


@pytest.fixture
def world(client, admin_token) -> World:
    return World(client, admin_token, "Page Size Org")


def test_a_layout_can_print_the_projects_code_and_description(world):
    project = world.project("Alpha")
    patched = world.client.patch(
        f"/api/v1/projects/{project['id']}", json={"summary": "Autonomous inspection drone"}, headers=auth_headers(world.token),
    )
    assert patched.status_code == 200, patched.text
    text = "@layout document A3\n\n@heading {{project_code}}: {{project_name}}\n\n{{project_description}}\n"
    resp = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts/preview", json={"text": text}, headers=auth_headers(world.token),
    )
    assert resp.status_code == 200, resp.text
    html = resp.json()["html"]
    assert "Autonomous inspection drone" in html and f"{project['code']}: Alpha" in html
