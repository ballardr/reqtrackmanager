"""Tests for a document layout's margins, mirrored margins and page styles (`margins`, `mirror_margins`, `page_styles`
and the `section` block): the tree check, the text view, the neutral model, the evaluator and the PDF, Word,
OpenDocument text and HTML renderers.
"""

from __future__ import annotations

import base64
import io
import re
import zipfile
import zlib
from typing import Any

import pytest
from docx import Document as DocxDocument

from app.services.report_docx import render_docx_document
from app.services.report_html import render_html_document
from app.services.report_layout_text import format_tree, parse_text
from app.services.report_layout_tree import InvalidLayoutError, assert_valid, empty_tree
from app.services.report_model import Margins, PageBreak, PageStyle, PageStyleChange, Paragraph, ReportDocument, Table
from app.services.report_odt import render_odt_document
from app.services.report_page_styles import resolve_page_styles
from app.services.report_pdf import render_pdf_document
from tests.conftest import auth_headers
from tests.link_world import World

WIDE = {"name": "Wide tables", "page_size": "A3", "orientation": "landscape", "margins": {"left": 1, "right": 1}}


def tree_with(**keys: Any) -> dict[str, Any]:
    return {**empty_tree("document"), **keys}


def errors_of(tree: dict[str, Any]) -> list[tuple[str, str]]:
    with pytest.raises(InvalidLayoutError) as caught:
        assert_valid(tree)
    return [(e.path, e.message) for e in caught.value.errors]


# ---- the model --------------------------------------------------------------------------------------------------


def test_the_usable_width_leaves_out_the_margins_and_the_gutter():
    style = PageStyle("A4", False, Margins(2, 2, 2.5, 1.5, 1))
    assert style.usable_width_cm == pytest.approx(21.0 - 2.5 - 1.5 - 1)
    assert style.usable_height_cm == pytest.approx(29.7 - 4)


def test_the_gutter_is_on_the_left_unless_margins_mirror_and_then_on_the_inner_edge_of_each_page():
    plain = PageStyle(margins=Margins(2, 2, 2, 1, 1))
    assert plain.edges_cm() == plain.edges_cm(verso=True) == (3, 1)
    mirrored = PageStyle(margins=Margins(2, 2, 2, 1, 1), mirror=True)
    assert mirrored.edges_cm() == (3, 1)
    assert mirrored.edges_cm(verso=True) == (1, 3)


def test_a_document_with_no_margins_keeps_one_and_a_half_cm_on_every_side():
    document = ReportDocument(title="T")
    assert document.style.margins == Margins(1.5, 1.5, 1.5, 1.5, 0)


# ---- the tree ---------------------------------------------------------------------------------------------------


def test_margins_and_page_styles_are_part_of_a_valid_document_tree():
    tree = tree_with(margins={"top": 2, "gutter": 1.5}, mirror_margins=True, page_styles=[WIDE], nodes=[
        {"type": "text", "text": "Before"},
        {"type": "section", "page_style": "Wide tables", "children": [{"type": "text", "text": "Inside"}]},
    ])
    assert_valid(tree)


def test_a_style_leaves_out_what_it_shares_with_the_layout():
    tree = tree_with(page_size="A5", margins={"top": 3}, page_styles=[{"name": "Plain"}, {"name": "Land", "orientation": "landscape"}])
    document, styles = resolve_page_styles(tree)
    assert document.page_size == "A5" and document.margins.top == 3
    assert styles["Plain"] == document
    assert styles["Land"].landscape and styles["Land"].page_size == "A5" and styles["Land"].margins.top == 3


@pytest.mark.parametrize("margins", [{"top": 0.5}, {"left": 10.5}, {"gutter": 5.5}, {"gutter": -1}, {"top": "2"}, {"top": True}, {"x": 1}, []])
def test_a_margin_outside_its_range_or_of_the_wrong_kind_is_refused_with_its_path(margins):
    errors = errors_of(tree_with(margins=margins))
    assert errors and all(path.startswith("margins") for path, _ in errors)


def test_margins_that_leave_no_room_for_text_are_refused():
    errors = errors_of(tree_with(page_size="A5", margins={"left": 7, "right": 7}))
    assert errors == [("margins", "The margins leave less than 5 cm for text on a A5 page.")]
    errors = errors_of(tree_with(page_styles=[{"name": "Tight", "margins": {"top": 10, "bottom": 10, "left": 1, "right": 1}},
                                              {"name": "Fine"}], page_size="A5"))
    assert [p for p, _ in errors] == ["page_styles[0].margins"]


def test_a_page_style_needs_a_unique_name_and_known_keys():
    errors = dict(errors_of(tree_with(page_styles=[{"name": "A"}, {"name": "A"}, {"orientation": "sideways"}, {"name": "B", "colour": 1}])))
    assert "used by another" in errors["page_styles[1].name"]
    assert errors["page_styles[2].name"].startswith("Required")
    assert "page_styles[2].orientation" in errors and errors["page_styles[3].colour"] == "Unknown key."


def test_there_is_a_limit_on_the_number_of_page_styles():
    errors = errors_of(tree_with(page_styles=[{"name": f"S{i}"} for i in range(11)]))
    assert errors == [("page_styles", "Too many page styles (the limit is 10).")]


def test_mirror_margins_must_be_true_or_false():
    assert errors_of(tree_with(mirror_margins="yes")) == [("mirror_margins", "Must be true or false.")]


def test_a_table_layout_has_no_pages_to_style():
    tree = {**empty_tree("table"), "margins": {"top": 2}, "mirror_margins": True, "page_styles": [WIDE]}
    assert {path for path, _ in errors_of(tree)} == {"margins", "mirror_margins", "page_styles"}


def test_a_section_must_name_a_page_style_of_the_layout():
    tree = tree_with(page_styles=[WIDE], nodes=[{"type": "section", "page_style": "Nope", "children": []}])
    errors = errors_of(tree)
    assert errors[0][0] == "nodes[0].page_style" and "Wide tables" in errors[0][1]


def test_a_section_cannot_be_inside_another_section_even_through_an_if():
    inner = {"type": "section", "page_style": "Wide tables", "children": []}
    tree = tree_with(page_styles=[WIDE], nodes=[{"type": "section", "page_style": "Wide tables", "children": [
        {"type": "if", "when": {"module": "requirements"}, "children": [inner]},
    ]}])
    assert any("inside another section" in message for _, message in errors_of(tree))
    # Two sections one after the other are fine.
    assert_valid(tree_with(page_styles=[WIDE], nodes=[inner, inner]))


def test_a_section_is_a_document_block_only():
    tree = {**empty_tree("table"), "nodes": [{"type": "section", "page_style": "x", "children": []}]}
    assert "cannot be used in a table layout" in errors_of(tree)[0][1]


# ---- the text view ----------------------------------------------------------------------------------------------

TEXT = """@layout document landscape A3 mirror margins top=2 gutter=1
@pagestyle "Wide tables" portrait A4 margins left=1.5 right=1.5
@pagestyle "Plain"

Intro

@section "Wide tables"
  Inside

  @page_break
@end
"""


def test_the_text_view_round_trips_margins_page_styles_and_sections():
    parsed = parse_text(TEXT)
    assert parsed.ok, parsed.errors
    assert parsed.tree["mirror_margins"] is True and parsed.tree["margins"] == {"top": 2, "gutter": 1}
    assert parsed.tree["page_styles"][0] == {
        "name": "Wide tables", "orientation": "portrait", "page_size": "A4", "margins": {"left": 1.5, "right": 1.5},
    }
    assert parsed.tree["nodes"][1]["type"] == "section"
    written = format_tree(parsed.tree)
    assert written.splitlines()[:3] == [
        "@layout document landscape A3 mirror margins top=2 gutter=1",
        '@pagestyle "Wide tables" portrait A4 margins left=1.5 right=1.5',
        '@pagestyle "Plain"',
    ]
    assert parse_text(written).tree == parsed.tree


@pytest.mark.parametrize(("text", "fragment"), [
    ("@layout document margins\n", "Write the margins as side=cm"),
    ("@layout document margins top=2 top=3\n", '"top" is given twice'),
    ("@layout document margins side=2\n", "Unknown margin"),
    ("@layout document margins top=wide\n", "number of cm"),
    ("@layout document mirror mirror\n", "written once"),
    ("@layout document\n\nText\n\n@pagestyle \"A\"\n", "before the first block"),
    ("@layout document\n@pagestyle\n", "name in quotes"),
    ("@layout document\n@section\n@end\n", "name in quotes"),
])
def test_a_mistake_in_the_new_lines_is_a_positioned_error(text, fragment):
    parsed = parse_text(text)
    assert not parsed.ok and any(fragment in e.message for e in parsed.errors)
    assert all(e.line >= 1 for e in parsed.errors)


def test_a_section_naming_no_page_style_is_reported_on_its_line():
    parsed = parse_text('@layout document\n\n@section "Nope"\n  Hi\n@end\n')
    assert parsed.tree is not None and [(e.line, "not a page style" in e.message) for e in parsed.errors] == [(3, True)]


def test_a_margin_error_in_a_page_style_is_reported_on_its_line():
    parsed = parse_text('@layout document\n@pagestyle "A" margins top=0.2\n')
    assert [e.line for e in parsed.errors] == [2]


# ---- the renderers ----------------------------------------------------------------------------------------------


def document(*blocks, **options: Any) -> ReportDocument:
    return ReportDocument(title="T", blocks=list(blocks), **options)


A3_LANDSCAPE = PageStyle("A3", True, Margins(1, 1, 1, 1))


def media_boxes(pdf: bytes) -> list[tuple[int, int]]:
    return [(round(float(w)), round(float(h))) for w, h in re.findall(rb"/MediaBox \[ ?0 0 ([\d.]+) ([\d.]+) ?\]", pdf)]


def first_text_x(pdf: bytes) -> list[float]:
    """The x position of the first text on each page, from the page content streams."""
    found: list[float] = []
    for raw in re.findall(rb"stream\r?\n(.*?)endstream", pdf, re.S):
        try:
            content = zlib.decompress(base64.a85decode(raw.strip(), adobe=True)).decode("latin-1")
        except Exception:  # noqa: BLE001 - not a page content stream (a font or the document info)
            continue
        match = re.search(r"1 0 0 1 ([\d.]+) [\d.]+ cm", content[content.find("q"):])
        if match:
            found.append(float(match.group(1)))
    return found


def test_the_pdf_changes_page_size_where_a_page_style_changes_and_returns_after():
    pdf = render_pdf_document(document(Paragraph("A"), PageStyleChange(A3_LANDSCAPE), Paragraph("B"), PageStyleChange(PageStyle()), Paragraph("C")))
    assert media_boxes(pdf) == [(595, 842), (1191, 842), (595, 842)]


def test_the_pdf_text_starts_inside_the_chosen_margin():
    margin_cm = 3
    pdf = render_pdf_document(document(Paragraph("A"), margins=Margins(2, 2, margin_cm, 2)))
    # The page's frame adds 6 points of padding inside the margin.
    assert first_text_x(pdf)[0] == pytest.approx(margin_cm * 28.3465 + 6, abs=0.5)


def test_pdf_margins_mirror_on_facing_pages_and_the_gutter_is_on_the_inner_edge():
    blocks = [Paragraph("one"), PageBreak(), Paragraph("two"), PageBreak(), Paragraph("three")]
    mirrored = first_text_x(render_pdf_document(document(*blocks, margins=Margins(2, 2, 2, 1, 1), mirror=True)))
    plain = first_text_x(render_pdf_document(document(*blocks, margins=Margins(2, 2, 2, 1, 1))))
    inner, outer = (2 + 1) * 28.3465 + 6, 1 * 28.3465 + 6
    assert mirrored == pytest.approx([inner, outer, inner], abs=0.5)
    assert plain == pytest.approx([inner, inner, inner], abs=0.5)


def test_a_mirrored_section_starts_on_the_side_that_matches_its_page_number():
    # "one" is page 1 (right-hand); the section starts page 2 (left-hand); the return starts page 3 (right-hand).
    default = PageStyle("A4", False, Margins(2, 2, 2, 1, 1), mirror=True)
    section = PageStyle("A4", False, Margins(1, 1, 3, 0.5), mirror=True)
    pdf = render_pdf_document(document(
        Paragraph("one"), PageStyleChange(section), Paragraph("two"), PageStyleChange(default), Paragraph("three"),
        margins=default.margins, mirror=True,
    ))
    unit = 28.3465
    assert first_text_x(pdf) == pytest.approx([(2 + 1) * unit + 6, 0.5 * unit + 6, (2 + 1) * unit + 6], abs=0.5)


def test_the_word_file_has_a_section_per_page_style_with_its_size_margins_and_gutter():
    data = render_docx_document(document(
        Paragraph("A"), PageStyleChange(A3_LANDSCAPE), Paragraph("B"), PageStyleChange(PageStyle(margins=Margins(2, 2, 2, 2, 1))),
        Paragraph("C"), margins=Margins(2, 2, 2, 2, 1),
    ))
    sections = DocxDocument(io.BytesIO(data)).sections
    assert [(round(s.page_width.cm, 1), round(s.page_height.cm, 1)) for s in sections] == [(21.0, 29.7), (42.0, 29.7), (21.0, 29.7)]
    assert [round(s.left_margin.cm, 1) for s in sections] == [2.0, 1.0, 2.0]
    assert [round(s.gutter.cm, 1) for s in sections] == [1.0, 0.0, 1.0]


def test_word_mirrors_margins_for_the_whole_file_only_when_asked():
    def settings(mirror: bool) -> str:
        data = render_docx_document(document(Paragraph("A"), mirror=mirror))
        return zipfile.ZipFile(io.BytesIO(data)).read("word/settings.xml").decode()

    assert "<w:mirrorMargins/>" in settings(True) and "mirrorMargins" not in settings(False)


def test_the_word_cover_page_footer_rule_applies_to_the_first_section_only():
    from app.services.report_document import ReportBranding

    branding = ReportBranding(accent_color_hex="#112233", include_cover_page=True)
    data = render_docx_document(document(Paragraph("A"), PageStyleChange(A3_LANDSCAPE), Paragraph("B"), branding=branding))
    first, second = DocxDocument(io.BytesIO(data)).sections
    assert first.different_first_page_header_footer is True and second.different_first_page_header_footer is False


def test_the_opendocument_file_has_a_page_layout_and_master_page_per_page_style():
    data = render_odt_document(document(
        Paragraph("A"), PageStyleChange(PageStyle("A3", True, Margins(1, 1, 1, 1, 0.5), mirror=True)), Paragraph("B"),
        PageStyleChange(PageStyle(margins=Margins(2, 2, 2, 2, 1), mirror=True)), Paragraph("C"), margins=Margins(2, 2, 2, 2, 1), mirror=True,
    ))
    package = zipfile.ZipFile(io.BytesIO(data))
    styles, content = package.read("styles.xml").decode(), package.read("content.xml").decode()
    assert re.findall(r'<style:master-page style:name="(\w+)" style:page-layout-name="(\w+)"', styles) == [
        ("Standard", "PL0"), ("Style1", "PL1"), ("Cover", "PL0"),
    ]
    assert 'style:page-usage="mirrored"' in styles
    assert 'fo:page-width="42.000cm"' in styles and 'fo:margin-left="1.500cm"' in styles  # the gutter joins the inner margin
    assert 'fo:margin-left="3.000cm"' in styles
    assert content.count('<text:p text:style-name="P_Style1"/>') == 1 and content.count('<text:p text:style-name="P_AfterCover"/>') == 1


def test_the_html_view_shows_a_labelled_section_per_page_style_and_prints_with_named_pages():
    home = PageStyle(margins=Margins(2, 2, 2, 2, 1), mirror=True)
    page = render_html_document(document(
        Paragraph("A"), PageStyleChange(A3_LANDSCAPE), Paragraph("B"), PageStyleChange(home), Paragraph("C"),
        margins=home.margins, mirror=True,
    )).decode()
    assert '<section class="page-style ps1" style="width: 99rem">' in page and page.count("</section>") == 1
    assert "Page style: A3 landscape; margins" in page and "Page style: A4 portrait" in page
    assert "@page { size: 21cm 29.7cm; margin: 2cm 2cm 2cm 3cm; }" in page and "@page:left { margin: 2cm 3cm 2cm 2cm; }" in page
    assert "@page ps1 { size: 42cm 29.7cm; margin: 1cm 1cm 1cm 1cm; }" in page and ".ps1 { page: ps1;" in page
    assert page.index("B") < page.index("</section>") < page.rindex("C")


def test_a_table_in_a_wide_section_uses_that_sections_width():
    columns = [f"Column {i}" for i in range(12)]
    wide = Table(columns, [[f"value {i}-{j}" for i in range(12)] for j in range(3)], compact=True)
    pdf = render_pdf_document(document(Paragraph("A"), PageStyleChange(A3_LANDSCAPE), wide))
    assert media_boxes(pdf) == [(595, 842), (1191, 842)]
    odt = zipfile.ZipFile(io.BytesIO(render_odt_document(document(Paragraph("A"), PageStyleChange(A3_LANDSCAPE), wide)))).read("content.xml").decode()
    assert f'style:width="{A3_LANDSCAPE.usable_width_cm:.3f}cm"' in odt


# ---- the evaluator and the API ----------------------------------------------------------------------------------


@pytest.fixture
def world(client, admin_token) -> World:
    return World(client, admin_token, "Page Styles Org")


def preview(world: World, project: dict, text: str) -> dict:
    resp = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts/preview", json={"text": text}, headers=auth_headers(world.token),
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_a_section_runs_as_a_page_style_change_and_back(world):
    project = world.project("Alpha")
    text = (
        '@layout document\n@pagestyle "Wide" A3 landscape margins left=1 right=1\n\nBefore\n\n'
        '@section "Wide"\n  Inside\n@end\n\nAfter\n'
    )
    html = preview(world, project, text)["html"]
    assert html.index("Before") < html.index('class="page-style ps1"') < html.index("Inside") < html.index("</section>") < html.index("After")


def test_a_section_with_nothing_in_it_leaves_no_page_behind(world):
    project = world.project("Alpha")
    text = (
        '@layout document\n@pagestyle "Wide" A3 landscape\n\nBefore\n\n'
        '@section "Wide"\n  @each requirement\n    {{record.title}}\n  @end\n@end\n\nAfter\n'
    )
    html = preview(world, project, text)["html"]
    assert 'class="page-style' not in html and "Before" in html and "After" in html


def test_a_layout_with_page_styles_is_saved_read_back_and_shown_as_text(world):
    project = world.project("Alpha")
    tree = tree_with(margins={"top": 2}, mirror_margins=True, page_styles=[WIDE], nodes=[
        {"type": "section", "page_style": "Wide tables", "children": [{"type": "text", "text": "Hi"}]},
    ])
    created = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts",
        json={"name": "Styled", "description": "d", "family": "document", "tree": tree}, headers=auth_headers(world.token),
    )
    assert created.status_code == 201, created.text
    version = world.client.get(
        f"/api/v1/projects/{project['id']}/report-layouts/{created.json()['id']}/versions/1", headers=auth_headers(world.token),
    )
    assert version.status_code == 200, version.text
    saved = version.json()["tree"]
    assert saved["page_styles"] == [WIDE] and saved["margins"] == {"top": 2} and saved["mirror_margins"] is True
    text = world.client.post(
        f"/api/v1/projects/{project['id']}/report-layouts/text/format", json={"tree": saved}, headers=auth_headers(world.token),
    )
    assert text.status_code == 200, text.text
    assert text.json()["text"].splitlines()[:2] == [
        "@layout document portrait mirror margins top=2", '@pagestyle "Wide tables" landscape A3 margins left=1 right=1',
    ]


def test_the_block_catalogue_describes_the_section_block_and_the_new_layout_fields(world):
    project = world.project("Alpha")
    resp = world.client.get(f"/api/v1/projects/{project['id']}/report-layouts/blocks", headers=auth_headers(world.token))
    assert resp.status_code == 200, resp.text
    catalogue = resp.json()
    section = next(b for b in catalogue["blocks"] if b["type"] == "section")
    assert [f["kind"] for f in section["fields"]] == ["page_style"] and [c["key"] for c in section["containers"]] == ["children"]
    kinds = {f["key"]: f["kind"] for f in catalogue["layout_fields"]}
    assert kinds["margins"] == "margins" and kinds["page_styles"] == "page_styles" and kinds["mirror_margins"] == "boolean"
    assert {"page_style", "margins", "page_styles"} <= set(catalogue["field_kinds"])
