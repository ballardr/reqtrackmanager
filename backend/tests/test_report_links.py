"""Tests for links in report text: `[words](address)` for web, email and in-document (section) addresses.

The rules under test: only an absolute http / https / mailto address or `#heading` is a link (`link_problem`); bare
addresses and autolinks stay text; an address may hold a placeholder, which the layout evaluator fills and checks
on the final text (a record value can never smuggle in another kind of address); a link to a missing heading is plain
words; and each renderer writes a link only after `finalize_links`. The model and renderer tests need no database;
the evaluator tests build an organisation.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document as DocxFile

from app.services.report_docx import render_docx_document
from app.services.report_html import render_html_document
from app.services.report_model import (
    BulletList,
    Gap,
    Heading,
    MathText,
    Paragraph,
    ReportDocument,
    Styled,
    finalize_links,
    heading_anchor,
    link_problem,
    markdown_to_blocks,
    text_runs,
)
from app.services.report_pdf import _InlineImages, _markup
from app.services.report_render import RenderInput, render
from tests.test_report_layout_math import blocks_of, db, doc, make_requirement, per_requirement, run, runs, text, world  # noqa: F401

LINK = Styled("the spec", href="https://example.org/spec")


def read(markdown: str, **kwargs: Any) -> list:
    return [b for b in markdown_to_blocks(markdown, **kwargs) if not isinstance(b, Gap)]


# --- what may be a link ------------------------------------------------------------------------------------------


@pytest.mark.parametrize("href", [
    "https://example.org", "http://example.org/a/b?x=1&y=2#frag", "HTTPS://EXAMPLE.ORG", "https://user@example.org:8443/p",
    "mailto:team@example.org", "mailto:a@b.co?subject=Hi", "#overview", "#scope-goals-2026", "#überblick",
])
def test_these_addresses_may_be_links(href):
    assert link_problem(href) is None


@pytest.mark.parametrize("href", [
    "", "javascript:alert(1)", "JaVaScRiPt:alert(1)", "data:text/html,<b>x</b>", "file:///etc/passwd", "vbscript:x", "ftp://example.org",
    "//example.org", "/relative/path", "relative", "http://", "https:///path", "mailto:", "mailto:not an address", "mailto:a@b",
    "#", "# spaces", "https://example.org/a b", "https://example.org/\"onclick=x", "https://example.org/<script>", "https://e.org/\\x",
    "https://example.org/\x00", "https://" + "a" * 2000,
])
def test_these_addresses_may_not_and_the_reason_never_repeats_the_address(href):
    problem = link_problem(href)
    assert problem and (not href or href not in problem)


def test_an_anchor_is_the_headings_words_lower_cased_with_hyphens():
    assert [heading_anchor(t) for t in ("Overview", "Scope & Goals (2026)", "  ", "Überblick 1", "A_b", "Req 12.3", "ALL CAPS")] == [
        "overview", "scope-goals-2026", "section", "überblick-1", "a-b", "req-12-3", "all-caps",
    ]


def test_the_link_vectors_agree_with_the_server():
    """`tests/data/link_vectors.json` holds addresses and headings with the answers the rules must give."""
    vectors = json.loads((Path(__file__).resolve().parent / "data" / "link_vectors.json").read_text(encoding="utf-8"))

    assert all(link_problem(href) is None for href in vectors["links"]["allowed"])
    assert all(link_problem(href) is not None for href in vectors["links"]["refused"])
    assert all(heading_anchor(words) == anchor for words, anchor in vectors["anchors"])


# --- the reader --------------------------------------------------------------------------------------------------


def test_a_link_is_a_styled_run_that_reads_as_its_words():
    block = read("See [the **spec**](https://example.org/s?a=1) now")[0]

    assert list(text_runs(block.text)) == [
        "See ", Styled("the ", href="https://example.org/s?a=1"), Styled("spec", bold=True, href="https://example.org/s?a=1"), " now",
    ]
    assert str(block.text) == "See the spec now"


def test_email_and_section_links_and_links_in_lists():
    blocks = read("[mail](mailto:a@b.co) and [jump](#details)\n\n1. [first](https://example.org/1)\n2. plain")

    assert [p.href for p in blocks[0].text.parts if isinstance(p, Styled)] == ["mailto:a@b.co", "#details"]
    assert list(text_runs(blocks[1].items[0])) == [Styled("first", href="https://example.org/1")]


@pytest.mark.parametrize("markdown", [
    "[bad](javascript:alert(1)) here", "[bad](data:text/html,x) here", "[bad](/relative) here", "[bad](https://a b) here",
    "[empty]() here", "[bad](ftp://x.org) here",
])
def test_a_link_with_an_address_that_is_not_allowed_stays_as_written(markdown):
    block = read(markdown)[0]
    assert block == Paragraph(markdown) and not isinstance(block.text, MathText)


@pytest.mark.parametrize("markdown", [
    "A bare address https://example.org/x here", "An autolink <https://example.org/x> here", "www.example.org and mail me@example.org",
    "A reference [text][ref] link", "[ref]: https://example.org/x",
])
def test_bare_addresses_autolinks_and_reference_links_stay_text(markdown):
    assert read(markdown)[0] == Paragraph(markdown)


def test_content_placeholders_in_an_address_are_filled_and_then_checked():
    values = {"project_name": "Atlas", "organisation": "javascript:alert(1)"}

    good = read("[docs](https://example.org/{{project_name}}/x)", placeholders=values)[0]
    bad = read("[docs]({{organisation}})", placeholders=values)[0]

    assert list(text_runs(good.text)) == [Styled("docs", href="https://example.org/Atlas/x")]
    assert bad == Paragraph("[docs](javascript:alert(1))")


def test_in_a_layout_an_address_with_a_placeholder_is_kept_for_the_evaluator():
    block = read("[x]({{record.url}}) and [y](https://h.org/{{record.code}})")[0]

    assert [p.href for p in block.text.parts if isinstance(p, Styled)] == ["{{record.url}}", "https://h.org/{{record.code}}"]


def test_equations_inside_link_text_are_put_back_unread():
    block = read(r"[load \(F_1 * m\) here](https://example.org)")[0]

    assert list(text_runs(block.text)) == [Styled(r"load \(F_1 * m\) here", href="https://example.org")]


# --- finalize_links ----------------------------------------------------------------------------------------------


def test_headings_get_unique_anchors_in_reading_order_and_the_input_is_unchanged():
    original = ReportDocument("T", [Heading("Overview"), Heading("Details"), Heading("Overview"), Heading("Overview 2"), Heading("overview")])

    final = finalize_links(original)

    assert [b.anchor for b in final.blocks] == ["overview", "details", "overview-2", "overview-2-2", "overview-3"]
    assert all(b.anchor is None for b in original.blocks)


def test_a_section_link_follows_its_heading_and_a_missing_one_becomes_plain_words():
    text = MathText([
        Styled("here", href="#details"), " and ", Styled("gone", bold=True, href="#nowhere"), " and ", Styled("plain", href="#nowhere"),
    ])
    final = finalize_links(ReportDocument("T", [Heading("Details"), Paragraph(text), BulletList([text])]))

    parts = final.blocks[1].text.parts
    assert parts == (Styled("here", href="#details"), " and ", Styled("gone", bold=True), " and ", "plain")
    assert str(final.blocks[1].text) == "here and gone and plain"


def test_an_address_that_slipped_in_without_the_reader_is_dropped_again_before_rendering():
    hostile = MathText([Styled("x", href='https://example.org/"><script>alert(1)</script>'), Styled("y", href="javascript:alert(1)")])

    final = finalize_links(ReportDocument("T", [Paragraph(hostile)]))

    assert final.blocks[0].text == "xy" and not isinstance(final.blocks[0].text, MathText)


# --- the layout evaluator ----------------------------------------------------------------------------------------


def test_an_address_filled_in_from_a_record_is_checked_on_the_final_text(world, db):  # noqa: F811
    project = world.project("Link P")
    make_requirement(world, project, "Spec doc", description="https://example.org/spec")
    make_requirement(world, project, "Evil doc", description="javascript:alert(1)")
    make_requirement(world, project, "Spacey doc", description="https://example.org/a b")
    tree = per_requirement(text("[{{record.title}}]({{record.description}})"))

    result = run(db, world, project, tree)
    found = {str(b.text): list(text_runs(b.text)) for b in blocks_of(result)}

    assert found["Spec doc"] == [Styled("Spec doc", href="https://example.org/spec")]
    assert found["Evil doc"] == ["Evil doc"] and found["Spacey doc"] == ["Spacey doc"]
    assert any("link address was not allowed" in w for w in result.warnings)
    assert not any("javascript" in w for w in result.warnings)


def test_a_record_value_in_link_text_or_a_section_link_works_and_a_title_never_becomes_a_link(world, db):  # noqa: F811
    project = world.project("Text P")
    make_requirement(world, project, "[sneaky](https://evil.example)")
    tree = doc(
        {"type": "heading", "text": "Overview"},
        {"type": "each", "source": {"record_type": "requirement"}, "children": [text("See [{{record.title}}](#overview)")]},
    )

    found = blocks_of(run(db, world, project, tree))

    assert runs(found[1]) == ["See ", Styled("[sneaky](https://evil.example)", href="#overview")]


# --- the renderers -----------------------------------------------------------------------------------------------


def linked_document() -> ReportDocument:
    paragraph = MathText([
        "See ", Styled("the spec", bold=True, href="https://example.org/spec?a=1&b=2"), ", ", Styled("mail", href="mailto:team@example.org"), ", ",
        Styled("details", href="#details"), " and ", Styled("lost", href="#nowhere"),
    ])
    back = BulletList([MathText([Styled("back", href="#overview")])])
    return ReportDocument("T", [Heading("Overview", 1), Paragraph(paragraph), Heading("Details", 2), back])


def test_html_links_are_anchors_with_ids_and_safe_attributes():
    html = render_html_document(linked_document()).decode()

    assert '<h2 id="overview">Overview</h2>' in html and '<h3 id="details">Details</h3>' in html
    assert '<a href="https://example.org/spec?a=1&amp;b=2" rel="noopener noreferrer nofollow" target="_blank"><strong>the spec</strong></a>' in html
    assert '<a href="mailto:team@example.org" rel="noopener noreferrer nofollow" target="_blank">mail</a>' in html
    assert '<a href="#details">details</a>' in html and "lost" in html and 'href="#nowhere"' not in html
    assert "<script" not in html


def test_docx_has_external_relationships_internal_anchors_and_bookmarks():
    data = render_docx_document(linked_document())
    archive = zipfile.ZipFile(io.BytesIO(data))
    body = archive.read("word/document.xml").decode()
    relations = archive.read("word/_rels/document.xml.rels").decode()

    assert relations.count('TargetMode="External"') == 2
    assert "https://example.org/spec?a=1&amp;b=2" in relations and "mailto:team@example.org" in relations
    names = re.findall(r'<w:bookmarkStart [^>]*w:name="(sec_[0-9a-f]+)"', body)
    anchors = re.findall(r'<w:hyperlink w:anchor="(sec_[0-9a-f]+)"', body)
    assert len(names) == 2 and set(anchors) <= set(names) and len(anchors) == 2
    paragraph = next(p for p in DocxFile(io.BytesIO(data)).paragraphs if p.text.startswith("See"))
    assert "lost" in paragraph.text and "details" in paragraph.text


def test_odt_links_are_text_a_elements_and_headings_carry_bookmarks():
    out = render("odt", "document", RenderInput("t", document=linked_document()))
    content = zipfile.ZipFile(io.BytesIO(out.content)).read("content.xml").decode()

    assert 'xlink:href="https://example.org/spec?a=1&amp;b=2"' in content and 'xlink:href="mailto:team@example.org"' in content
    assert 'xlink:href="#details"' in content and 'xlink:href="#nowhere"' not in content
    assert '<text:bookmark text:name="overview"/>' in content and '<text:bookmark text:name="details"/>' in content


def test_pdf_markup_makes_anchors_and_the_file_has_link_annotations():
    document = finalize_links(linked_document())
    images = _InlineImages()
    try:
        markup = _markup(document.blocks[1].text, images)
    finally:
        images.close()

    assert '<a href="https://example.org/spec?a=1&amp;b=2" color="#0563C1"><u><b>the spec</b></u></a>' in markup
    assert '<a href="#details"' in markup and "#nowhere" not in markup
    pdf = render("pdf", "document", RenderInput("t", document=linked_document())).content
    assert pdf.startswith(b"%PDF") and b"/URI" in pdf


def test_a_hostile_address_never_reaches_any_renderer():
    hostile = ReportDocument("T", [Paragraph(MathText([Styled("x", href='https://e.org/"><b>'), Styled("y", href="javascript:alert(1)")]))])

    for fmt in ("html", "docx", "odt", "pdf"):
        content = render(fmt, "document", RenderInput("t", document=hostile)).content
        if fmt in ("docx", "odt"):
            archive = zipfile.ZipFile(io.BytesIO(content))
            content = b"".join(archive.read(name) for name in archive.namelist() if name.endswith((".xml", ".rels")))
        assert b"javascript" not in content, fmt
        assert b'"><b>' not in content and b"e.org" not in content, fmt
