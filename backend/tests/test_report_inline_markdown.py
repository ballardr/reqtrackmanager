"""Tests for bold, italic, code and numbered lists in report text: the Markdown reader (`report_model.markdown_to_blocks`),
the layout evaluator's handling of styled runs, and the four document renderers (HTML, DOCX, ODT, PDF markup).

The rules under test: only `**bold**`, `*italic*` / `_italic_` and `` `code` `` are styles; links, raw HTML, entities
and images inside a sentence stay as written; equations are set aside and never read for styles; a record value or
placeholder substituted later is never read for styles. The reader and renderer tests need no database; the evaluator
tests build an organisation.
"""

from __future__ import annotations

import io
import zipfile
from typing import Any

import pytest
from docx import Document as DocxFile
from lxml import etree

from app.services.report_docx import render_docx_document
from app.services.report_html import render_html_document
from app.services.report_layout_text import format_tree, parse_text
from app.services.report_model import (
    BulletList,
    Equation,
    Gap,
    Heading,
    MathText,
    NumberedList,
    Paragraph,
    ReportDocument,
    Styled,
    markdown_to_blocks,
    text_runs,
)
from app.services.report_pdf import _InlineImages, _markup
from app.services.report_render import RenderInput, render
from tests.test_report_layout_math import blocks_of, db, doc, per_requirement, run, runs, text, world  # noqa: F401  (fixtures)

BOLD = Styled("x", bold=True)


def read(markdown: str, **kwargs: Any) -> list:
    return [b for b in markdown_to_blocks(markdown, **kwargs) if not isinstance(b, Gap)]


def runs_of(markdown: str) -> list:
    block = read(markdown)[0]
    return list(text_runs(block.text))


# --- the reader --------------------------------------------------------------------------------------------------


def test_bold_italic_and_code_become_styled_runs_and_the_words_read_without_markers():
    block = read("Plain **bold** and *italic* and `code` and ***both***.")[0]

    assert list(text_runs(block.text)) == [
        "Plain ", Styled("bold", bold=True), " and ", Styled("italic", italic=True), " and ",
        Styled("code", code=True), " and ", Styled("both", bold=True, italic=True), ".",
    ]
    assert str(block.text) == "Plain bold and italic and code and both."


def test_text_without_styles_is_the_plain_string_it_always_was():
    assert read("Just words, 5 * 3 and snake_case_name.")[0] == Paragraph("Just words, 5 * 3 and snake_case_name.")


def test_an_underscore_inside_a_word_is_not_emphasis_but_a_pair_of_them_around_words_is():
    assert runs_of("snake_case_name and _em_ here") == ["snake_case_name and ", Styled("em", italic=True), " here"]


def test_code_keeps_markers_inside_it_literally():
    assert runs_of("Use `**not bold**` here") == ["Use ", Styled("**not bold**", code=True), " here"]


def test_equations_are_set_aside_and_never_read_for_styles_or_escapes():
    assert read(r"Math \(a_1 * b_2\) and **bold** and $$x*y*z$$ end")[0].text.parts == (
        r"Math \(a_1 * b_2\) and ", Styled("bold", bold=True), " and $$x*y*z$$ end",
    )
    assert runs_of(r"**see \(x^2\) now**") == [Styled(r"see \(x^2\) now", bold=True)]


@pytest.mark.parametrize("markdown", [
    "An <http://example.com> autolink",
    "Raw <b>html</b> and <img src=x>",
    "An &amp; entity",
    "An ![image](attachment:1) inside a sentence",
    "~~struck~~ text",
    r"Path C:\temp\new",
])
def test_everything_else_stays_as_the_author_wrote_it(markdown):
    block = read(markdown)[0]
    assert block == Paragraph(markdown) and not isinstance(block.text, MathText)


def test_a_backslash_escapes_a_bracket_so_it_is_not_taken_for_a_link():
    assert read(r"Regex \d+\. and \[x\]")[0] == Paragraph(r"Regex \d+\. and [x]")


def test_a_backslash_escape_makes_a_marker_literal():
    assert read(r"a \*not italic\* b")[0] == Paragraph("a *not italic* b")
    assert read(r"snake \_case\_ and back\\slash and tick \` here")[0] == Paragraph("snake _case_ and back\\slash and tick ` here")


@pytest.mark.parametrize("markdown", [r"An unclosed \( opener", r"A stray \) closer", r"Path C:\temp\new\(x)", r"Regex \d+\. and \d"])
def test_a_backslash_anywhere_else_prints_as_written(markdown):
    assert read(markdown)[0] == Paragraph(markdown)


def test_a_paragraph_that_is_only_an_image_is_still_an_image():
    from app.services.report_model import Image

    assert read("![alt](attachment:1)", images={"attachment:1": b"png"}) == [Image(b"png")]
    assert read("![alt](attachment:2)", images={"attachment:1": b"png"}) == []


def test_lists_keep_their_kind_and_their_styles():
    blocks = read("- one\n- **two**\n\n1. first\n2. *second*\n3. third")

    assert isinstance(blocks[0], BulletList) and isinstance(blocks[1], NumberedList)
    assert blocks[0].items == ["one", "two"] and isinstance(blocks[0].items[1], MathText)
    assert blocks[1].items == ["first", "second", "third"]
    assert list(text_runs(blocks[1].items[1])) == [Styled("second", italic=True)]


def test_a_nested_list_is_flattened_not_cut_short():
    blocks = read("- a\n  - b\n  - c\n- d\n\nafter")

    assert blocks[0].items == ["a", "b", "c", "d"] and blocks[1] == Paragraph("after")


def test_a_heading_is_plain_words_without_markers():
    assert read("## A **bold** `title`")[0] == Heading("A bold title", 2)


def test_placeholders_in_content_are_filled_run_by_run_and_a_value_is_never_read_for_styles():
    block = read("**{{project_name}}** and {{organisation}}", placeholders={"project_name": "Atlas", "organisation": "*Acme* **Co**"})[0]

    assert list(text_runs(block.text)) == [Styled("Atlas", bold=True), " and *Acme* **Co**"]


# --- the layout evaluator ----------------------------------------------------------------------------------------


def test_the_evaluator_keeps_styles_and_a_substituted_value_takes_the_runs_style_but_is_not_read(world, db):  # noqa: F811
    project = world.project("Styled P")
    from tests.test_report_layout_math import make_requirement

    make_requirement(world, project, "Weird **title** `x`")
    tree = per_requirement(text("Req **{{record.title}}** is *done*"))

    found = blocks_of(run(db, world, project, tree))

    assert runs(found[0]) == ["Req ", Styled("Weird **title** `x`", bold=True), " is ", Styled("done", italic=True)]


def test_an_equation_inside_a_styled_run_is_drawn_and_the_words_around_it_keep_the_style(world, db):  # noqa: F811
    project = world.project("Mixed P")

    found = blocks_of(run(db, world, project, doc(text(r"**load \(F = m a\) applies** now"))))

    assert runs(found[0]) == [Styled("load ", bold=True), Equation("F = m a", False), Styled(" applies", bold=True), " now"]


def test_a_display_equation_still_ends_the_paragraph_and_styles_continue_after_it(world, db):  # noqa: F811
    project = world.project("Display P")

    found = blocks_of(run(db, world, project, doc(text("**before $$y$$ after**"))))

    assert found[1] == Equation("y", True)
    assert runs(found[0]) == [Styled("before ", bold=True)] and runs(found[2]) == [Styled(" after", bold=True)]


def test_lists_in_a_layout_keep_their_styles(world, db):  # noqa: F811
    project = world.project("List P")

    found = blocks_of(run(db, world, project, doc(text("1. one\n2. **two**"))))

    assert isinstance(found[0], NumberedList) and list(text_runs(found[0].items[1])) == [Styled("two", bold=True)]


def test_a_text_block_with_styles_reads_back_from_its_text_view_unchanged():
    tree = doc(text("A **bold** line.\n\n1. one\n2. two"))
    assert parse_text(format_tree(tree)).tree == tree


# --- the renderers -----------------------------------------------------------------------------------------------


def styled_document() -> ReportDocument:
    mixed = MathText([
        "Plain ", Styled("bold", bold=True), " ", Styled("it", italic=True), " ", Styled("co<de>", code=True), " ",
        Styled("both", bold=True, italic=True),
    ])
    return ReportDocument("T", [
        Heading("Head", 1), Paragraph(mixed), BulletList([MathText([Styled("bullet", bold=True), " b"])]),
        NumberedList(["first", MathText(["second ", Styled("it", italic=True)])]),
    ])


def test_html_draws_styles_with_semantic_tags_and_escapes_the_words():
    html = render_html_document(styled_document()).decode()

    assert "<p>Plain <strong>bold</strong> <em>it</em> <code>co&lt;de&gt;</code> <strong><em>both</em></strong></p>" in html
    assert "<ul><li><strong>bullet</strong> b</li></ul>" in html
    assert "<ol><li>first</li><li>second <em>it</em></li></ol>" in html


def test_docx_sets_bold_italic_and_a_code_font_on_the_runs_and_uses_numbered_list_style():
    out = DocxFile(io.BytesIO(render_docx_document(styled_document())))
    paragraph = next(p for p in out.paragraphs if p.text.startswith("Plain"))

    flags = [(r.text, bool(r.bold), bool(r.italic), r.font.name) for r in paragraph.runs]
    assert flags == [
        ("Plain ", False, False, None), ("bold", True, False, None), (" ", False, False, None), ("it", False, True, None),
        (" ", False, False, None), ("co<de>", False, False, "Courier New"), (" ", False, False, None), ("both", True, True, None),
    ]
    assert [p.text for p in out.paragraphs if p.style.name == "List Number"] == ["first", "second it"]


def test_odt_wraps_runs_in_character_styles_that_are_defined_and_numbers_the_list():
    out = render("odt", "document", RenderInput("t", document=styled_document()))
    archive = zipfile.ZipFile(io.BytesIO(out.content))
    ns = {"text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0", "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0"}
    content = etree.fromstring(archive.read("content.xml"))
    styles = etree.fromstring(archive.read("styles.xml"))

    used = {span.get(f"{{{ns['text']}}}style-name"): span.text for span in content.iter(f"{{{ns['text']}}}span")}
    assert used["T_Code"] == "co<de>" and used["T_BoldItalic"] == "both"
    assert {"T_Bold", "T_Italic", "T_Code", "T_BoldItalic"} <= set(used)
    defined = {s.get(f"{{{ns['style']}}}name") for s in styles.iter(f"{{{ns['style']}}}style")}
    assert set(used) <= defined
    numbered = [item for lst in content.iter(f"{{{ns['text']}}}list") if lst.get(f"{{{ns['text']}}}style-name") == "L_Number" for item in lst]
    assert len(numbered) == 2
    assert "L_Number" in {s.get(f"{{{ns['style']}}}name") for s in styles.iter(f"{{{ns['text']}}}list-style")}


def test_pdf_markup_nests_the_tags_and_escapes_the_words():
    images = _InlineImages()
    try:
        markup = _markup(styled_document().blocks[1].text, images)
    finally:
        images.close()

    assert "<b>bold</b>" in markup and "<i>it</i>" in markup
    assert '<font face="Courier">co&lt;de&gt;</font>' in markup
    assert "<b><i>both</i></b>" in markup


def test_pdf_renders_a_document_with_styles_and_a_numbered_list():
    out = render("pdf", "document", RenderInput("t", document=styled_document()))

    assert out.content.startswith(b"%PDF")


def test_a_bold_run_in_a_heading_is_not_wrapped_in_a_second_bold_tag_in_the_pdf():
    images = _InlineImages()
    try:
        markup = _markup(MathText([Styled("x", bold=True), Styled("y", italic=True)]), images, bold=True)
    finally:
        images.close()

    assert "<b>" not in markup and "<i>" in markup
